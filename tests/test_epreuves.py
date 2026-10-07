# -*- coding: utf-8 -*-
"""Tests du module Comité d'évaluation, épreuves et grilles (m_epreuves).
Lancer : python3 -m unittest tests.test_epreuves   (depuis /home/claude/gescol)

L'édition testée est choisie par GESCOL_EDITION_TEST (privee par défaut) ; la suite est rejouée dans un second
processus pour l'édition « publique » (TestEditionPublique)."""
import html as html_mod
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "privee")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_epreuves_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import m_epreuves as E  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def txt(r):
    return html_mod.unescape(r.data.decode())


def q(sql, params=(), un=False):
    with app.app_context():
        from app import q as _q
        return _q(sql, params, un=un)


class TestCalculs(unittest.TestCase):
    def test_items(self):
        f = {"nb_items": "4", "it_0_intitule": "Q1", "it_0_bareme": "2", "it_1_intitule": "", "it_1_bareme": "9",
             "it_2_intitule": "Q2", "it_2_bareme": "3,5", "it_2_critere": "Cohérence", "it_3_intitule": "Q3", "it_3_bareme": "abc"}
        items, erreurs = E.lire_items(f)
        self.assertEqual([i["ordre"] for i in items], [1, 2, 3])
        self.assertEqual(E.total_items(items), 5.5)
        self.assertEqual(len(erreurs), 1)
        self.assertEqual(E.num(5.5), "5,5")
        self.assertEqual(E.num(3.0), "3")

    def test_grille_type(self):
        gm = E.grille_modele("production_ecrite")
        self.assertEqual([l["bareme"] for l in gm["lignes"]], [3, 3, 2, 2])
        self.assertEqual(E.total_grille(gm["lignes"]), 10)
        self.assertEqual(E.grille_modele("maternelle")["type"], "observation")
        self.assertEqual(E.total_grille(E.grille_modele("maternelle")["lignes"]), 0)


class TestEpreuves(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        cls.c = app.test_client()
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex
            ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
            ex("UPDATE ecoles SET nom=?, departement=?, cs=?, directeur_nom=? WHERE id=1", ("EPP TEST EPREUVES", "Zou", "CS Abomey", "HOUNSOU Paul"))
        r = cls.c.post("/login", data={"identifiant": "admin", "mdp": "admin123"})
        assert r.status_code == 302, r.status_code
        cls.annee = config.ANNEE_ACTIVE

    def test_01_pages(self):
        for url in ("/epreuves", "/epreuves/comite", "/epreuves/grilles", "/epreuves/0", "/epreuves/comite/0", "/epreuves/grilles/0",
                    "/epreuves/grilles/0?modele=mathematiques", "/epreuves/comite/import"):
            r = self.c.get(url)
            self.assertEqual(r.status_code, 200, f"{url} -> {r.status_code}")
        self.assertIn("/epreuves", txt(self.c.get("/")))     # lien de menu
        self.assertEqual(self.c.get("/epreuves/999").status_code, 404)
        self.assertEqual(self.c.get("/epreuves/999.pdf").status_code, 404)
        self.assertTrue(self.c.get("/epreuves/comite.pdf").data.startswith(b"%PDF"))
        html = txt(self.c.get("/epreuves/grilles/0?modele=mathematiques"))
        self.assertIn("Compréhension de la situation", html)
        self.assertIn("Démarche / raisonnement", html)

    def test_02_comite(self):
        r = self.c.post("/epreuves/comite/0", data={"nom_prenoms": "AKPO Jean", "qualite": "Directeur / directrice", "qualification": "CAP, A1-5",
                                                   "fonction_comite": "Président(e)", "classe_ou_niveau": "Toute l'école", "contact": "97 00 00 00"},
                        follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("AKPO Jean", txt(r))
        self.assertIn("ajouté au comité", txt(r))
        r = self.c.post("/epreuves/comite/0", data={"nom_prenoms": ""})
        self.assertEqual(r.status_code, 400)
        self.c.post("/epreuves/comite/0", data={"nom_prenoms": "DOSSOU Marie", "qualite": "Enseignant(e)", "fonction_comite": "Rapporteur", "classe_ou_niveau": "CM2"})
        m = q("SELECT * FROM eval_comite WHERE nom_prenoms='AKPO Jean'", un=True)
        self.assertEqual(m["annee"], self.annee)
        self.assertEqual(m["ecole_id"], 1)
        r = self.c.post(f"/epreuves/comite/{m['id']}", data={"nom_prenoms": "AKPO Jean-Pierre", "fonction_comite": "Président(e)", "qualification": "CAP"}, follow_redirects=True)
        self.assertIn("AKPO Jean-Pierre", txt(r))
        # PDF et Excel
        r = self.c.get("/epreuves/comite.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        if openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/epreuves/comite.xlsx").data))
            vals = [c.value for row in wb.active.iter_rows() for c in row]
            self.assertIn("AKPO Jean-Pierre", vals)
            self.assertIn("Rapporteur", vals)
            # import : mise à jour + ajout
            wb2 = openpyxl.Workbook()
            ws = wb2.active
            ws.append(E.ENT_COMITE)
            ws.append(["DOSSOU Marie", "Enseignant(e)", "CEAP, B1-12", "Rapporteur", "CM2", "96 00 00 00", "oui"])
            ws.append(["KOTO Alice", "Conseiller pédagogique (CP)", "", "Personne ressource", "", "", "oui"])
            ws.append(["", "", "", "", "", "", ""])
            buf = io.BytesIO()
            wb2.save(buf)
            r = self.c.post("/epreuves/comite/import", data={"fichier": (io.BytesIO(buf.getvalue()), "comite.xlsx")},
                            content_type="multipart/form-data", follow_redirects=True)
            self.assertIn("1 membre(s) ajouté(s), 1 mis à jour", txt(r))
            self.assertEqual(q("SELECT qualification FROM eval_comite WHERE nom_prenoms='DOSSOU Marie'", un=True)["qualification"], "CEAP, B1-12")
            self.assertTrue(q("SELECT id FROM eval_comite WHERE nom_prenoms='KOTO Alice'", un=True))
        r = self.c.get("/epreuves/comite/modele.xlsx")
        self.assertIn(r.status_code, (200, 302))
        # suppression
        r = self.c.post(f"/epreuves/comite/{m['id']}/supprimer", follow_redirects=True)
        self.assertIn("retiré du comité", txt(r))
        self.assertEqual(self.c.post(f"/epreuves/comite/{m['id']}/supprimer").status_code, 404)

    def test_03_epreuve(self):
        mat = q("SELECT * FROM matieres WHERE cycle='primaire' ORDER BY ordre LIMIT 1", un=True)
        data = {"titre": "Évaluation sommative n° 2 — Français CM2", "type_evaluation": "sommative", "numero": "2", "niveau": "CM2",
                "matiere": f"{mat['code']}|{mat['nom']}", "duree_minutes": "60", "bareme_total": "10", "date_prevue": f"{self.annee[:4]}-12-10",
                "situation": "Texte d'amorce : Le fermier et ses enfants…\nDeuxième ligne.", "consignes": "Lis attentivement puis réponds.",
                "auteur": "DOSSOU Marie", "nb_items": "3",
                "it_0_intitule": "Qui est le personnage principal ?", "it_0_critere": "Pertinence de la production", "it_0_bareme": "3", "it_0_reponse": "Le fermier",
                "it_1_intitule": "Raconte la fin en trois phrases.", "it_1_critere": "Cohérence de la production", "it_1_bareme": "4,5", "it_1_reponse": "",
                "it_2_intitule": "Souligne les verbes.", "it_2_critere": "Correction de la langue", "it_2_bareme": "2,5", "it_2_reponse": "mange, court"}
        r = self.c.post("/epreuves/0", data=data, follow_redirects=True)
        self.assertEqual(r.status_code, 200, txt(r)[:300])
        html = txt(r)
        self.assertIn("créée (3 item(s), barème 10)", html)
        self.assertNotIn("Attention : barème déclaré", html)
        e = q("SELECT * FROM epreuves WHERE titre LIKE 'Évaluation sommative n° 2%'", un=True)
        self.assertEqual(e["annee"], self.annee)
        self.assertEqual(e["matiere_code"], mat["code"])
        self.assertEqual(e["statut"], "brouillon")
        items = q("SELECT * FROM epreuves_items WHERE epreuve_id=? ORDER BY ordre", (e["id"],))
        self.assertEqual(len(items), 3)
        self.assertEqual(sum(i["bareme"] for i in items), 10)
        self.assertEqual(items[1]["bareme"], 4.5)
        TestEpreuves.eid = e["id"]
        # liste et filtres
        html = txt(self.c.get("/epreuves"))
        self.assertIn("Évaluation sommative n° 2", html)
        self.assertIn("DOSSOU Marie", html)
        self.assertNotIn("Évaluation sommative n° 2", txt(self.c.get("/epreuves?type=formative")))
        self.assertIn("Évaluation sommative n° 2", txt(self.c.get(f"/epreuves?niveau=CM2&matiere={mat['code']}&statut=brouillon")))
        # modification : barème déclaré ≠ somme -> avertissement
        r = self.c.post(f"/epreuves/{e['id']}", data=dict(data, bareme_total="20", statut="brouillon"), follow_redirects=True)
        self.assertIn("Attention : barème déclaré 20 ≠ somme des items 10", txt(r))
        self.assertIn("≠ 20", txt(self.c.get("/epreuves")))
        self.c.post(f"/epreuves/{e['id']}", data=dict(data, statut="brouillon"), follow_redirects=True)
        # erreurs
        r = self.c.post("/epreuves/0", data=dict(data, titre=""))
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/epreuves/0", data=dict(data, it_0_bareme="xx"))
        self.assertEqual(r.status_code, 400)
        self.assertIn("barème invalide", txt(r))

    def test_04_pdf_sujet_corrige_bareme(self):
        eid = self.eid
        r = self.c.get(f"/epreuves/{eid}.pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"%PDF"))
        r2 = self.c.get(f"/epreuves/{eid}.pdf?corrige=1")
        self.assertTrue(r2.data.startswith(b"%PDF"))
        r3 = self.c.get(f"/epreuves/{eid}/bareme.pdf")
        self.assertTrue(r3.data.startswith(b"%PDF"))
        try:
            import pdfplumber
        except ImportError:
            pdfplumber = None
        if pdfplumber:
            with pdfplumber.open(io.BytesIO(r.data)) as pdf:
                t = "\n".join(p.extract_text() or "" for p in pdf.pages)
            self.assertIn("SITUATION", t)
            self.assertIn("Le fermier et ses enfants", t)
            self.assertIn("Qui est le personnage principal", t)
            self.assertIn("EPP TEST EPREUVES", t)
            self.assertNotIn("mange, court", t)
            with pdfplumber.open(io.BytesIO(r2.data)) as pdf:
                t2 = "\n".join(p.extract_text() or "" for p in pdf.pages)
            self.assertIn("CORRIG", t2)
            self.assertIn("mange, court", t2)

    def test_05_statut_duplication_suppression(self):
        eid = self.eid
        r = self.c.post(f"/epreuves/{eid}/statut", data={"statut": "validee"}, follow_redirects=True)
        self.assertIn("validée", txt(r))
        self.assertEqual(q("SELECT statut FROM epreuves WHERE id=?", (eid,), un=True)["statut"], "validee")
        r = self.c.post(f"/epreuves/{eid}/statut", data={"statut": "administree"}, follow_redirects=True)
        self.assertEqual(q("SELECT statut FROM epreuves WHERE id=?", (eid,), un=True)["statut"], "administree")
        self.assertEqual(self.c.post(f"/epreuves/{eid}/statut", data={"statut": "xx"}).status_code, 400)
        # duplication : copie en brouillon avec ses items
        r = self.c.post(f"/epreuves/{eid}/dupliquer", follow_redirects=True)
        self.assertIn("dupliquée", txt(r))
        copie = q("SELECT * FROM epreuves WHERE titre LIKE '%(copie)'", un=True)
        self.assertEqual(copie["statut"], "brouillon")
        self.assertEqual(q("SELECT COUNT(*) n FROM epreuves_items WHERE epreuve_id=?", (copie["id"],), un=True)["n"], 3)
        # validation refusée sans item
        r = self.c.post("/epreuves/0", data={"titre": "Vide", "type_evaluation": "formative", "nb_items": "0"}, follow_redirects=True)
        vide = q("SELECT * FROM epreuves WHERE titre='Vide'", un=True)
        r = self.c.post(f"/epreuves/{vide['id']}/statut", data={"statut": "validee"}, follow_redirects=True)
        self.assertIn("sans item", txt(r))
        # export Excel de la liste
        if openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/epreuves/liste.xlsx").data))
            self.assertEqual(wb.sheetnames, ["Épreuves", "Items"])
            vals = [c.value for row in wb["Items"].iter_rows() for c in row]
            self.assertIn("Souligne les verbes.", vals)
        # suppression de la copie
        r = self.c.post(f"/epreuves/{copie['id']}/supprimer", follow_redirects=True)
        self.assertIn("supprimée", txt(r))
        self.assertEqual(q("SELECT COUNT(*) n FROM epreuves_items WHERE epreuve_id=?", (copie["id"],), un=True)["n"], 0)
        self.assertEqual(self.c.post(f"/epreuves/{copie['id']}/supprimer").status_code, 404)
        # interface exposée
        import notif
        with app.app_context():
            liste = notif.H.epreuves_liste(1, self.annee)
        self.assertTrue(any(e["id"] == eid and e["bareme_items"] == 10 and len(e["items"]) == 3 for e in liste))

    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_06_import_items(self):
        eid = self.eid
        r = self.c.get("/epreuves/items-modele.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(E.ENT_ITEMS)
        ws.append([1, "Item importé A", "Pertinence de la production", "", 4, "réponse A"])
        ws.append([2, "Item importé B", "Cohérence de la production", "ind", "6", ""])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post(f"/epreuves/{eid}/items/import", data={"fichier": (io.BytesIO(buf.getvalue()), "items.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("2 item(s) importé(s)", txt(r))
        items = q("SELECT * FROM epreuves_items WHERE epreuve_id=? ORDER BY ordre", (eid,))
        self.assertEqual([i["intitule"] for i in items], ["Item importé A", "Item importé B"])
        self.assertEqual(sum(i["bareme"] for i in items), 10)
        r = self.c.post(f"/epreuves/{eid}/items/import", data={}, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Choisissez un fichier", txt(r))

    def test_07_grilles(self):
        # grille type production écrite chargée puis enregistrée
        html = txt(self.c.get("/epreuves/grilles/0?modele=production_ecrite"))
        self.assertIn("Pertinence de la production", html)
        self.assertIn("Correction de la langue", html)
        data = {"titre": "Grille production écrite CM2", "type": "correction", "niveau": "CM2", "nb_lignes": "4",
                "lg_0_critere": "Pertinence de la production", "lg_0_indicateurs": "Respect du type de texte\nIdées en rapport", "lg_0_bareme": "3",
                "lg_1_critere": "Cohérence de la production", "lg_1_indicateurs": "Enchaînement", "lg_1_bareme": "3",
                "lg_2_critere": "Correction de la langue", "lg_2_indicateurs": "Orthographe", "lg_2_bareme": "2",
                "lg_3_critere": "Perfectionnement", "lg_3_indicateurs": "Originalité", "lg_3_bareme": "2", "epreuve_id": str(self.eid)}
        r = self.c.post("/epreuves/grilles/0", data=data, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("créée (4 critère(s))", txt(r))
        gr = q("SELECT * FROM grilles WHERE titre='Grille production écrite CM2'", un=True)
        self.assertEqual(gr["epreuve_id"], self.eid)
        contenu = json.loads(gr["contenu"])
        self.assertEqual([l["bareme"] for l in contenu["lignes"]], [3, 3, 2, 2])
        self.assertEqual(contenu["lignes"][0]["indicateurs"], ["Respect du type de texte", "Idées en rapport"])
        self.assertIn("10", txt(self.c.get("/epreuves/grilles")))
        # erreurs
        self.assertEqual(self.c.post("/epreuves/grilles/0", data={"titre": "", "nb_lignes": "0"}).status_code, 400)
        # modification
        r = self.c.post(f"/epreuves/grilles/{gr['id']}", data=dict(data, titre="Grille PE CM2 v2", lg_3_bareme="4"), follow_redirects=True)
        self.assertIn("Grille PE CM2 v2", txt(r))
        self.assertEqual(E.total_grille(json.loads(q("SELECT contenu FROM grilles WHERE id=?", (gr["id"],), un=True)["contenu"])["lignes"]), 12)
        # PDF
        r = self.c.get(f"/epreuves/grilles/{gr['id']}.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        # grille maternelle (observation) : PDF A / ECA / NA
        r = self.c.post("/epreuves/grilles/0", data={"titre": "Observation maternelle", "type": "observation", "niveau": "Maternelle 1", "nb_lignes": "1",
                                                      "lg_0_critere": "Langage", "lg_0_indicateurs": "Parle en phrases\nComprend une consigne"}, follow_redirects=True)
        gm = q("SELECT * FROM grilles WHERE titre='Observation maternelle'", un=True)
        r = self.c.get(f"/epreuves/grilles/{gm['id']}.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        self.assertIn("A / ECA / NA", txt(self.c.get("/epreuves/grilles")))
        # duplication, Excel, suppression
        r = self.c.post(f"/epreuves/grilles/{gr['id']}/dupliquer", follow_redirects=True)
        self.assertIn("dupliquée", txt(r))
        copie = q("SELECT * FROM grilles WHERE titre LIKE '%(copie)'", un=True)
        self.assertIsNone(copie["epreuve_id"])
        self.assertEqual(copie["contenu"], q("SELECT contenu FROM grilles WHERE id=?", (gr["id"],), un=True)["contenu"])
        if openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/epreuves/grilles.xlsx").data))
            vals = [c.value for row in wb.active.iter_rows() for c in row]
            self.assertIn("Grille PE CM2 v2", vals)
            self.assertIn("Observation maternelle", vals)
        r = self.c.post(f"/epreuves/grilles/{copie['id']}/supprimer", follow_redirects=True)
        self.assertIn("supprimée", txt(r))
        self.assertEqual(self.c.get(f"/epreuves/grilles/{copie['id']}.pdf").status_code, 404)
        # la suppression de l'épreuve détache la grille
        self.c.post(f"/epreuves/{self.eid}/supprimer", follow_redirects=True)
        self.assertIsNone(q("SELECT epreuve_id FROM grilles WHERE id=?", (gr["id"],), un=True)["epreuve_id"])

    def test_08_droits(self):
        """Un comptable lit mais n'écrit pas ; un enseignant écrit."""
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex
            for ident, role in (("compta_ep", "comptable"), ("ens_ep", "enseignant")):
                ex("INSERT OR IGNORE INTO utilisateurs (identifiant, nom, role, mdp_hash, doit_changer_mdp, ecole_id) VALUES (?,?,?,?,0,1)",
                   (ident, ident, role, generate_password_hash("x")))
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "compta_ep", "mdp": "x"}).status_code, 302)
        self.assertEqual(c2.get("/epreuves").status_code, 200)
        self.assertEqual(c2.get("/epreuves/comite").status_code, 200)
        self.assertEqual(c2.get("/epreuves/0").status_code, 403)
        self.assertEqual(c2.post("/epreuves/comite/0", data={"nom_prenoms": "X"}).status_code, 403)
        c3 = app.test_client()
        self.assertEqual(c3.post("/login", data={"identifiant": "ens_ep", "mdp": "x"}).status_code, 302)
        self.assertEqual(c3.get("/epreuves/0").status_code, 200)


class TestEditionPublique(unittest.TestCase):
    """Rejoue toute la suite dans un processus séparé avec l'édition « publique »."""

    def test_publique(self):
        if EDITION != "privee":
            self.skipTest("déjà dans l'édition secondaire")
        env = dict(os.environ, GESCOL_EDITION_TEST="publique")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "tests.test_epreuves.TestEpreuves", "tests.test_epreuves.TestCalculs"],
                           cwd=RACINE, env=env, capture_output=True, text=True, timeout=600)
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
