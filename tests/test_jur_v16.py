# -*- coding: utf-8 -*-
"""Tests v1.6 des applications Circonscription (CS) et DDEMP : classes par cours (tableau de bord), divisions / modules /
attributions, unités pédagogiques (UP) et module des conseillers pédagogiques (CP).
Lancer : python3 -m unittest tests.test_jur_v16   (depuis /home/claude/gescol)

L'édition testée est choisie par GESCOL_EDITION_TEST (cs par défaut) ; la suite est rejouée dans un second processus
pour l'édition « ddemp » (TestEditionsSeparees). Les données viennent de demo_juridiction.main() (base de démonstration
dans un dossier temporaire) ; connexion avec le compte « responsable » créé par la démo (mot de passe demo1234)."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "cs")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_jur16_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import a_tableau as AT  # noqa: E402
import a_up as AU  # noqa: E402
import demo_juridiction  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def txt(r):
    return html_mod.unescape(r.data.decode())


def xlsx_feuilles(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


# Avec « unittest discover », plusieurs modules de test partagent un processus : l'édition chargée par le premier module
# l'emporte. Si ce n'est pas la nôtre, la suite est rejouée dans des sous-processus (TestEditionsSeparees).
EN_PLACE = config.EDITION == EDITION


def _rejouer(edition):
    env = dict(os.environ, GESCOL_EDITION_TEST=edition)
    env.pop("GESCOL_EDITION", None)
    env.pop("GESCOL_DATA", None)
    return subprocess.run([sys.executable, "-m", "unittest", "tests.test_jur_v16.TestJurV16"], cwd=RACINE, env=env,
                          capture_output=True, text=True, timeout=900)


class TestJurV16(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not EN_PLACE:
            raise unittest.SkipTest(f"édition « {config.EDITION} » déjà chargée : suite rejouée en sous-processus")
        app.config["TESTING"] = True
        demo_juridiction.main()
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code
        with app.app_context():
            from app import q
            cls.cs_id = q("SELECT id FROM jur_cs ORDER BY numero, nom", un=True)["id"]
            cls.zone_id = q("SELECT id FROM jur_zones WHERE cs_id=? ORDER BY ordre", (cls.cs_id,), un=True)["id"]
        cls.annee = config.ANNEE_ACTIVE
        cls.ddemp = EDITION == "ddemp"

    def get(self, url, code=200):
        r = self.c.get(url)
        self.assertEqual(r.status_code, code, f"{url} → {r.status_code}\n{r.data[:400]}")
        return r

    def post(self, url, data=None, code=302):
        r = self.c.post(url, data=data or {})
        self.assertEqual(r.status_code, code, f"POST {url} → {r.status_code}\n{txt(r)[-1500:] if r.status_code != 302 else ''}")
        return r

    # ------------------------------------------------------------ 1. tableau de bord : classes par cours
    def test_01_tableau_classes(self):
        t = txt(self.get("/juridiction/tableau-de-bord"))
        self.assertIn("Classes par cours", t)
        self.assertIn("Total classes", t)
        for n in ("CI", "CP", "CE1", "CE2", "CM1", "CM2", "Maternelle 1"):
            self.assertIn(f"<th class=\"num\">{n}</th>", t)
        t2 = txt(self.get("/juridiction/classes"))
        self.assertIn("Par école", t2)
        if self.ddemp:
            self.get(f"/juridiction/classes?cs={self.cs_id}")

    def test_02_classes_calcul(self):
        """Le nombre de classes du tableau = nombre d'entrées « classes » des sections effectifs."""
        with app.app_context():
            import json
            from app import q
            tc = AT.jur_classes_par_niveau(self.annee)
            attendu = {}
            for r in q("SELECT s.donnees FROM jur_sections s JOIN jur_ecoles e ON e.id=s.ecole_id AND e.actif=1 "
                       "WHERE s.annee=? AND s.section='effectifs'", (self.annee,)):
                for c in json.loads(r["donnees"]).get("classes", []):
                    attendu[c["niveau"]] = attendu.get(c["niveau"], 0) + 1
            self.assertEqual(tc["total"]["total"], sum(attendu.values()))
            for n, v in attendu.items():
                self.assertEqual(tc["total"]["par_niveau"][n], v)
            # total des lignes (zones, ou CS en DDEMP) = total général
            lignes = [l for l in tc["lignes"] if l["niveau"] == ("cs" if self.ddemp else "zone")]
            self.assertEqual(sum(l["total"] for l in lignes), tc["total"]["total"])
            self.assertEqual(sum(l["ecoles"] for l in lignes), tc["total"]["ecoles"])
            # replis : effectifs par niveau, puis aucune donnée
            d, src = AT.classes_ecole({"par_niveau": {"CI": [12, 10], "CP": [0, 0], "CM2": {"garcons": 5, "filles": 4}}})
            self.assertEqual((d, src), ({"CI": 1, "CM2": 1}, "niveaux"))
            self.assertEqual(AT.classes_ecole(None, {"code": "XX-000", "nom": "École inexistante"}), ({}, "aucune"))
            self.assertEqual(AT.classes_ecole({"classes": [{"niveau": "ce1"}, {"niveau": "CE1"}, {"niveau": "GS"}]})[0],
                             {"CE1": 2, "Maternelle 2": 1})

    def test_03_classes_exports(self):
        r = self.get("/juridiction/classes.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.get("/juridiction/classes.xlsx")
        self.assertTrue(r.data.startswith(b"PK"))
        if openpyxl:
            f = xlsx_feuilles(r.data)
            self.assertIn("Classes par cours", f)
            self.assertIn("Par école", f)
            lignes = [l for l in f["Classes par cours"] if l and l[0] == "Total"]
            self.assertEqual(len(lignes), 1)
            with app.app_context():
                self.assertEqual(lignes[0][-3], AT.jur_classes_par_niveau(self.annee)["total"]["total"])

    # ------------------------------------------------------------ 2. divisions
    def test_10_divisions(self):
        t = txt(self.get("/juridiction/divisions"))
        self.assertIn("Divisions, modules et attributions", t)
        code = "organisation" if not self.ddemp else "spaf"
        sigle = "CDOSP" if not self.ddemp else "SPAF"
        self.assertIn(sigle, t)
        self.assertIn("Tableau de bord", t)                 # module existant → lien affiché
        self.assertIn("/juridiction/up", t)                 # module UP (ce chantier)
        t2 = txt(self.get(f"/juridiction/divisions/{code}"))
        self.assertIn("Modules de l'application", t2)
        self.assertIn("Agents affectés", t2)
        self.get("/juridiction/divisions/inconnue", 404)
        # modification des attributions
        self.post(f"/juridiction/divisions/{code}", {"attributions": "- nouvelle attribution de test ;\n- seconde ligne."})
        with app.app_context():
            from app import q
            self.assertEqual(q("SELECT attributions FROM jur_services WHERE code=? AND niveau=?", (code, EDITION), un=True)["attributions"],
                             "- nouvelle attribution de test ;\n- seconde ligne.")
        self.assertIn("nouvelle attribution de test", txt(self.get(f"/juridiction/divisions/{code}")))
        r = self.get("/juridiction/divisions.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        # seuls les endpoints existants sont proposés
        import a_divisions as AD
        with app.test_request_context("/"):
            from flask import g
            g.user = {"id": 1, "role": "responsable"}
            for m in AD.modules_division(app, code):
                self.assertIn(m["endpoint"], app.view_functions)
            self.assertTrue(AD.modules_division(app, "division_inconnue"))   # modules communs « * »

    # ------------------------------------------------------------ 3. unités pédagogiques
    def test_20_up(self):
        t = txt(self.get("/juridiction/up"))
        self.assertIn("Unités pédagogiques", t)
        self.assertIn("RUP", t)
        # ajout
        d = {"nom": "UP TEST", "zone_id": str(self.zone_id), "rup_nom": "KOUTON Basile", "rup_tel": "0190000000",
             "corup_nom": "OLOU Rosine", "nb_ecoles": "4", "ecole_siege": "EPP Test"}
        if self.ddemp:
            d["cs_id"] = str(self.cs_id)
        r = self.post("/juridiction/up/nouvelle", d)
        uid = int(r.headers["Location"].rstrip("/").split("/")[-1])
        self.post("/juridiction/up/nouvelle", d, 400)        # doublon refusé
        t = txt(self.get(f"/juridiction/up/{uid}"))
        self.assertIn("UP TEST", t)
        self.assertIn("KOUTON Basile", t)
        # modification
        d2 = dict(d, nom="UP TEST 2", corup_nom="ADJAHO Afiavi")
        self.post(f"/juridiction/up/{uid}/modifier", d2)
        self.assertIn("ADJAHO Afiavi", txt(self.get(f"/juridiction/up/{uid}")))
        # activités : 4 prévues, 3 exécutées → 75 %
        a0 = int(self.annee[:4])
        for i, (m, exe) in enumerate(((10, 1), (11, 1), (12, 0), (1, 1))):
            y = a0 if m >= 9 else a0 + 1
            self.post(f"/juridiction/up/{uid}/activites/nouvelle",
                      {"date": f"{y}-{m:02d}-1{i}", "theme": f"Thème {i}", "type": "Atelier", "prevue": "1",
                       **({"executee": "1"} if exe else {}), "participants": "12", "taux_participation": "80"})
        self.post(f"/juridiction/up/{uid}/activites/nouvelle", {"date": "", "theme": "", "prevue": "1"}, 400)   # vide refusé
        with app.app_context():
            from app import q
            acts = q("SELECT * FROM jur_up_activites WHERE up_id=? ORDER BY date", (uid,))
            self.assertEqual(len(acts), 4)
            self.assertEqual({a["annee"] for a in acts}, {self.annee})
            self.assertEqual(AU.taux_execution(4, 3), 75.0)
            self.assertIsNone(AU.taux_execution(0, 2))
            data = _H().jur_up_liste(self.annee, self.cs_id if self.ddemp else None)
            u = next(x for x in data["ups"] if x["id"] == uid)
            self.assertEqual((u["prevues"], u["executees"], u["taux"]), (4, 3, 75.0))
            self.assertEqual(len(u["dates"]), 4)
            self.assertEqual(u["themes"], ["Thème 0", "Thème 1", "Thème 2", "Thème 3"])
            z = next(z for z in data["zones"] if z["zone_id"] == self.zone_id)
            self.assertEqual(z["tot"]["prevues"], sum(x["prevues"] for x in z["ups"]))
            self.assertEqual(data["tot"]["up"], len(data["ups"]))
            self.assertEqual(data["tot"]["executees"], sum(x["executees"] for x in data["ups"]))
            aid = acts[2]["id"]
        t = txt(self.get(f"/juridiction/up/{uid}?annee={self.annee}"))
        self.assertIn("75 %", t)
        self.assertIn("Thème 2", t)
        t = txt(self.get("/juridiction/up?annee=" + self.annee + (f"&cs={self.cs_id}" if self.ddemp else "")))
        self.assertIn("UP TEST 2", t)
        self.assertIn("75 %", t)
        # modification puis suppression d'une activité → 4 prévues, 4 exécutées (100 %) puis 3 prévues
        self.post(f"/juridiction/up/activites/{aid}/modifier", {"date": f"{a0}-12-12", "theme": "Thème 2 bis", "prevue": "1", "executee": "1"})
        with app.app_context():
            u = next(x for x in _H().jur_up_liste(self.annee)["ups"] if x["id"] == uid)
            self.assertEqual((u["prevues"], u["executees"], u["taux"]), (4, 4, 100.0))
        self.post(f"/juridiction/up/activites/{aid}/supprimer")
        with app.app_context():
            u = next(x for x in _H().jur_up_liste(self.annee)["ups"] if x["id"] == uid)
            self.assertEqual((u["prevues"], u["executees"], u["taux"]), (3, 3, 100.0))
        # exports
        r = self.get(f"/juridiction/up.pdf?annee={self.annee}")
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.get(f"/juridiction/up.xlsx?annee={self.annee}")
        self.assertTrue(r.data.startswith(b"PK"))
        if openpyxl:
            f = xlsx_feuilles(r.data)
            self.assertEqual(set(f) >= {"Unités pédagogiques", "Activités", "Par zone"}, True)
            ligne = next(l for l in f["Unités pédagogiques"] if l and l[2] == "UP TEST 2")
            self.assertEqual((ligne[9], ligne[10], ligne[11]), (3, 3, 100))
            self.assertEqual(sum(1 for l in f["Activités"] if l and l[2] == "UP TEST 2"), 3)
        self.get("/juridiction/up/import")
        r = self.get("/juridiction/up/import?modele=1")
        self.assertTrue(r.data.startswith(b"PK"))
        # import : nouvelle UP + 2 activités (1 exécutée) ; réimport sans doublon
        if openpyxl:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Circonscription", "Zone pédagogique", "UP", "École siège", "RUP (nom et prénoms)", "Téléphone RUP",
                       "Co-RUP (nom et prénoms)", "Téléphone Co-RUP", "Nombre d'écoles", "Date", "Thème", "Type",
                       "Prévue (oui/non)", "Exécutée (oui/non)", "Participants"])
            cs_nom = ""
            with app.app_context():
                from app import q
                cs_nom = q("SELECT nom FROM jur_cs WHERE id=?", (self.cs_id,), un=True)["nom"]
            ws.append([cs_nom, "Zone pédagogique 1", "UP IMPORT", "EPP Import", "SALAMI Bio", "0191111111", "OKE Sika", "",
                       3, f"{a0}-10-20", "Thème import A", "Atelier", "oui", "oui", 10])
            ws.append([cs_nom, "Zone pédagogique 1", "UP IMPORT", "", "", "", "", "", "", f"{a0}-11-20", "Thème import B",
                       "Leçon modèle", "oui", "non", ""])
            buf = io.BytesIO()
            wb.save(buf)
            for _ in range(2):
                r = self.c.post("/juridiction/up/import", data={"fichier": (io.BytesIO(buf.getvalue()), "up.xlsx")},
                                content_type="multipart/form-data")
                self.assertEqual(r.status_code, 302)
            with app.app_context():
                from app import q
                ups = q("SELECT * FROM jur_up WHERE nom='UP IMPORT'")
                self.assertEqual(len(ups), 1)
                self.assertEqual(ups[0]["rup_nom"], "SALAMI Bio")
                self.assertEqual(ups[0]["cs_id"], self.cs_id)
                acts = q("SELECT * FROM jur_up_activites WHERE up_id=?", (ups[0]["id"],))
                self.assertEqual(len(acts), 2)
                u = next(x for x in _H().jur_up_liste(self.annee)["ups"] if x["nom"] == "UP IMPORT")
                self.assertEqual((u["prevues"], u["executees"], u["taux"]), (2, 1, 50.0))
            # le fichier exporté se réimporte sans créer de doublon
            x = self.get(f"/juridiction/up.xlsx?annee={self.annee}").data
            r = self.c.post("/juridiction/up/import", data={"fichier": (io.BytesIO(x), "export.xlsx")}, content_type="multipart/form-data")
            self.assertEqual(r.status_code, 302)
            with app.app_context():
                from app import q
                self.assertEqual(q("SELECT COUNT(*) n FROM jur_up WHERE nom='UP IMPORT'", un=True)["n"], 1)
                self.assertEqual(q("SELECT COUNT(*) n FROM jur_up_activites WHERE up_id=?", (ups[0]["id"],), un=True)["n"], 2)
        # suppression de l'UP (et de ses activités)
        self.post(f"/juridiction/up/{uid}/supprimer")
        with app.app_context():
            from app import q
            self.assertIsNone(q("SELECT 1 FROM jur_up WHERE id=?", (uid,), un=True))
            self.assertEqual(q("SELECT COUNT(*) n FROM jur_up_activites WHERE up_id=?", (uid,), un=True)["n"], 0)
        self.get(f"/juridiction/up/{uid}", 404)

    def test_21_up_droits(self):
        """Un agent (sans droit jur_up) lit la page mais ne saisit pas."""
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "agent", "mdp": "demo1234"}).status_code, 302)
        self.assertEqual(c2.get("/juridiction/up").status_code, 200)
        self.assertEqual(c2.get("/juridiction/up/nouvelle").status_code, 403)
        self.assertEqual(c2.post("/juridiction/divisions/" + ("organisation" if not self.ddemp else "spaf"),
                                 data={"attributions": "x"}).status_code, 403)

    # ------------------------------------------------------------ 4. conseillers pédagogiques
    def test_30_cp(self):
        t = txt(self.get("/juridiction/cp"))
        self.assertIn("Conseillers pédagogiques", t)
        self.assertIn("Synthèse mensuelle", t)
        self.assertIn("KOUTON Basile", t)       # CP de zone (démo) et visiteur
        if self.ddemp:
            self.get(f"/juridiction/cp?cs={self.cs_id}")
        with app.app_context():
            s = _H().jur_synthese_cp(self.annee, self.cs_id if self.ddemp else None)
            self.assertTrue(s["cps"])
            for x in s["cps"]:
                self.assertEqual(x["total"], x["visites_classe"] + x["visites_ecole"])
                self.assertEqual(sum(x["mois"]), x["total"])
                self.assertLessEqual(x["enseignants_perimetre_visites"], x["enseignants_zone"] or x["enseignants_perimetre_visites"])
                if x["total"]:
                    self.assertTrue(x["derniere"])
            self.assertEqual(s["tot"]["total"], sum(x["total"] for x in s["cps"]))
            if self.ddemp:
                s2 = _H().jur_synthese_cp(self.annee)
                self.assertEqual(len(s2["cs"]), len(_H().jur_cs_liste()))
                self.assertEqual(s2["tot"]["total"], sum(c["tot"]["total"] for c in s2["cs"]))
        r = self.get("/juridiction/cp?export=xlsx")
        self.assertTrue(r.data.startswith(b"PK"))
        if openpyxl:
            f = xlsx_feuilles(r.data)
            self.assertIn("Par CP", f)
            self.assertIn("Par mois", f)
        r = self.get("/juridiction/cp?export=pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        # l'existant n'est pas cassé
        self.get("/juridiction/bilan-activites")
        self.get("/juridiction/visites")

    # ------------------------------------------------------------ 5. menu
    def test_40_menu(self):
        t = txt(self.get("/juridiction/tableau-de-bord"))
        for lib in ("Divisions, modules et attributions", "Unités pédagogiques (UP)", "Conseillers pédagogiques (visites, synthèses)",
                    "Classes par cours et par zone", "Recensement, rentrée, fin d’année"):
            self.assertIn(lib, t)


def _H():
    """Objet H (helpers) de l'application : récupéré via la fonction exposée par a_tableau."""
    import a_tableau
    return a_tableau._H


class TestEditionsSeparees(unittest.TestCase):
    """Rejoue la suite dans un processus séparé pour l'édition « cs » (si elle n'a pas tourné ici) et « ddemp »."""

    def test_cs(self):
        if EN_PLACE and EDITION == "cs" or EDITION == "ddemp":
            self.skipTest("édition cs déjà testée dans ce processus, ou processus secondaire")
        r = _rejouer("cs")
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])

    def test_ddemp(self):
        if EDITION == "ddemp":
            self.skipTest("déjà dans l'édition secondaire")
        r = _rejouer("ddemp")
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
