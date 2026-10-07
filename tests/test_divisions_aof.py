# -*- coding: utf-8 -*-
"""Tests v1.6 — divisions : rubriques, lignes d'action conformes à l'AOF (a_divisions).
Lancer : python3 -m unittest tests.test_divisions_aof   (depuis /home/claude/gescol)

Base de démonstration (demo_juridiction.main()) dans un dossier temporaire, édition choisie par GESCOL_EDITION_TEST.
Sans cette variable (« python3 -m unittest discover » ou lancement direct), ce module n'impose PAS d'édition au
processus (il est le premier chargé par « discover » et les modules suivants attendent l'édition privée) : la suite
est rejouée dans deux sous-processus, « cs » puis « ddemp » (TestEditionsSeparees), comme tests/test_jur_v16.py.
Connexion avec le compte « responsable » (mot de passe demo1234)."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST")
EN_PLACE = EDITION is not None
if EN_PLACE:
    os.environ["GESCOL_EDITION"] = EDITION
    os.environ["GESCOL_DATA"] = tempfile.mkdtemp(prefix=f"gescol_divaof_{EDITION}_")
else:
    EDITION = "cs"
    os.environ.setdefault("GESCOL_DATA", tempfile.mkdtemp(prefix="gescol_tests_"))

from app import app  # noqa: E402
import config  # noqa: E402
import a_divisions as AD  # noqa: E402
import a_structure as S  # noqa: E402
import demo_juridiction  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

EN_PLACE = EN_PLACE and config.EDITION == EDITION
ANNEE_2 = "2031-2032"      # année sans aucune ligne d'action : génération automatique au premier affichage


def txt(r):
    return html_mod.unescape(r.data.decode())


def xlsx_feuilles(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def _rejouer(edition):
    env = dict(os.environ, GESCOL_EDITION_TEST=edition)
    env.pop("GESCOL_EDITION", None)
    env.pop("GESCOL_DATA", None)
    return subprocess.run([sys.executable, "-m", "unittest", "tests.test_divisions_aof.TestDivisionsAOF"], cwd=RACINE,
                          env=env, capture_output=True, text=True, timeout=900)


class TestDivisionsAOF(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not EN_PLACE:
            raise unittest.SkipTest("suite rejouée en sous-processus (éditions cs et ddemp)")
        app.config["TESTING"] = True
        demo_juridiction.main()
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code
        cls.annee = config.ANNEE_ACTIVE
        cls.ddemp = EDITION == "ddemp"
        cls.code = "spaf" if cls.ddemp else "animation"           # division de travail
        cls.redige = "spaf" if cls.ddemp else "personnel"         # division dont le texte est rédigé par GESCOL

    def get(self, url, code=200):
        r = self.c.get(url)
        self.assertEqual(r.status_code, code, f"{url} → {r.status_code}\n{r.data[:600]}")
        return r

    def post(self, url, data=None, code=302):
        r = self.c.post(url, data=data or {})
        self.assertEqual(r.status_code, code, f"POST {url} → {r.status_code}\n{txt(r)[-1500:] if r.status_code != 302 else ''}")
        return r

    def q(self, sql, params=()):
        with app.app_context():
            from app import q
            return q(sql, params)

    # ------------------------------------------------------------ 1. fonctions pures
    def test_01_fonctions(self):
        lignes = AD.lignes_depuis_attributions(S.ART_11)
        self.assertEqual(len(lignes), S.ART_11.count("\n- "))
        self.assertTrue(all(l[0].isupper() and not l.endswith(";") for l in lignes))
        self.assertEqual(len(AD.lignes_depuis_attributions(S.ART_10)), S.ART_10.count("• "))
        self.assertEqual(AD.lignes_depuis_attributions(S.ART_9), [])
        self.assertEqual(AD.lignes_depuis_attributions("1. premier ;\n2) second.\nsans puce"), ["Premier", "Second"])
        self.assertEqual(AD.numero_article("Article 12"), "12")
        self.assertEqual(AD.numero_article("AOF – article 7"), "7")
        self.assertEqual(AD.numero_article("13"), "13")
        self.assertEqual(AD.numero_article(""), "")
        self.assertEqual(AD.source_division({"reference": "Article 11"}), ("AOF – article 11", True))
        self.assertEqual(AD.source_division({"reference": "", "note": S.REDIGE}), (S.SOURCE_GESCOL, False))
        sy = AD.synthese_lignes([{"etat": "réalisée", "taux": 10}, {"etat": "en cours", "taux": 50}, {"etat": "prévue"},
                                 {"etat": "reportée", "taux": 0}])
        self.assertEqual((sy["total"], sy["realisees"], sy["en_cours"], sy["prevues"], sy["reportees"], sy["taux"]), (4, 1, 1, 1, 1, 37.5))
        self.assertIsNone(AD.synthese_lignes([])["taux"])
        self.assertEqual(AD.pct(37.5), "37.5 %")
        self.assertEqual(AD.pct(100.0), "100 %")

    # ------------------------------------------------------------ 2. pages
    def test_10_pages(self):
        t = txt(self.get("/juridiction/divisions"))
        self.assertIn("Divisions, modules et attributions", t)
        self.assertIn("Lignes d'action", t)
        self.assertIn("Article AOF", t)
        if not self.ddemp:
            for art in ("Article 10", "Article 11", "Article 12", "Article 13"):
                self.assertIn(art, t)
            self.assertIn("hors AOF", t)                      # CSPAF conservée, marquée hors AOF
            self.assertIn("CDOSP", t)
        self.assertIn(S.SOURCE_GESCOL, t)                      # textes rédigés par GESCOL signalés
        t2 = txt(self.get(f"/juridiction/divisions/{self.code}"))
        for mot in ("Rubrique", "Autorité de rattachement", "Chef", "Source", "Lignes d'action", "Modules de l'application", "Agents affectés"):
            self.assertIn(mot, t2)
        t3 = txt(self.get("/juridiction/divisions/plan"))
        self.assertIn("Plan d'action des divisions", t3)
        self.assertIn("TOTAL STRUCTURE", t3)
        self.get("/juridiction/divisions/inconnue", 404)
        self.get("/juridiction/divisions/inconnue/plan.pdf", 404)
        # colonnes ajoutées à jur_services et initialisées
        cols = {r["name"] for r in self.q("PRAGMA table_info(jur_services)")}
        self.assertTrue({"article_aof", "source_aof"} <= cols)
        if not self.ddemp:
            r = self.q("SELECT article_aof, source_aof FROM jur_services WHERE niveau=? AND code='animation'", (EDITION,))[0]
            self.assertEqual((r["article_aof"], r["source_aof"]), ("11", "AOF – article 11"))
        r = self.q("SELECT source_aof FROM jur_services WHERE niveau=? AND code=?", (EDITION, self.redige))[0]
        self.assertEqual(r["source_aof"], S.SOURCE_GESCOL)
        # menu
        self.assertIn("Plan d'action des divisions", txt(self.get("/juridiction/tableau-de-bord")))

    # ------------------------------------------------------------ 3. génération automatique des lignes depuis l'AOF
    def test_20_generation_auto(self):
        code = "animation" if not self.ddemp else "direction"
        with app.app_context():
            attr = S.service_par_code(code)["attributions"]
        attendu = len(AD.lignes_depuis_attributions(attr))
        if not self.ddemp:
            self.assertEqual(attr, S.ART_11)
            self.assertEqual(attendu, S.ART_11.count("\n- "))
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, ANNEE_2))[0]["n"], 0)
        t = txt(self.get(f"/juridiction/divisions/{code}?annee={ANNEE_2}"))
        lignes = self.q("SELECT * FROM jur_lignes_action WHERE niveau=? AND service_code=? AND annee=? ORDER BY ordre", (EDITION, code, ANNEE_2))
        self.assertEqual(len(lignes), attendu)
        self.assertEqual([l["ordre"] for l in lignes], list(range(1, attendu + 1)))
        self.assertTrue(all(l["etat"] == "prévue" for l in lignes))
        for l in lignes:
            self.assertIn(l["libelle"], t)
        # suppression d'une ligne, puis réaffichage : pas de re-génération silencieuse
        self.post(f"/juridiction/divisions/lignes/{lignes[0]['id']}/supprimer")
        self.get(f"/juridiction/divisions/{code}?annee={ANNEE_2}")
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, ANNEE_2))[0]["n"], attendu - 1)
        # « Régénérer depuis l'AOF » : la ligne manquante revient, les autres sont conservées
        self.post(f"/juridiction/divisions/{code}/lignes/generer", {"annee": ANNEE_2})
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, ANNEE_2))[0]["n"], attendu)
        self.post(f"/juridiction/divisions/{code}/lignes/generer", {"annee": ANNEE_2})
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, ANNEE_2))[0]["n"], attendu)
        # la démo a fourni des lignes de l'année active en état varié
        etats = {r["etat"] for r in self.q("SELECT DISTINCT etat FROM jur_lignes_action WHERE annee=?", (self.annee,))}
        self.assertGreaterEqual(len(etats), 3)

    # ------------------------------------------------------------ 4. CRUD des lignes d'action
    def test_30_crud(self):
        code = self.code
        base = f"/juridiction/divisions/{code}"
        avant = self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, self.annee))[0]["n"]
        self.get(f"{base}/lignes/nouvelle?annee={self.annee}")
        self.post(f"{base}/lignes/nouvelle", {"annee": self.annee, "libelle": ""}, 400)       # libellé obligatoire
        mods = AD.MODULES_PAR_DIVISION.get(code, [])
        ep = next((m[1] for m in mods if m[1] in app.view_functions), None)
        d = {"annee": self.annee, "libelle": "Organiser la journée pédagogique de test", "indicateur": "Journée tenue",
             "cible": "1", "echeance": "2027-02-15", "responsable": "KOUTON Basile", "etat": "en cours", "taux": "40",
             "observations": "obs test"}
        if ep:
            d["modules"] = ep
        self.post(f"{base}/lignes/nouvelle", d)
        l = self.q("SELECT * FROM jur_lignes_action WHERE libelle=?", ("Organiser la journée pédagogique de test",))
        self.assertEqual(len(l), 1)
        l = l[0]
        self.assertEqual((l["etat"], l["taux"], l["ordre"], l["service_code"], l["annee"]), ("en cours", 40, avant + 1, code, self.annee))
        if ep:
            self.assertEqual(l["modules"], f'["{ep}"]')
        t = txt(self.get(f"{base}?annee={self.annee}"))
        self.assertIn("Organiser la journée pédagogique de test", t)
        if ep:
            self.assertIn(AD.LIBELLES_MODULES[ep], t)
        # modification (fiche complète)
        self.get(f"/juridiction/divisions/lignes/{l['id']}/modifier")
        self.post(f"/juridiction/divisions/lignes/{l['id']}/modifier", dict(d, libelle="Organiser la journée pédagogique (modifiée)",
                                                                           etat="réalisée", taux="100", ordre="1"))
        l2 = self.q("SELECT * FROM jur_lignes_action WHERE id=?", (l["id"],))[0]
        self.assertEqual((l2["libelle"], l2["etat"], l2["ordre"]), ("Organiser la journée pédagogique (modifiée)", "réalisée", 1))
        # suivi rapide (état, taux, réordonnancement)
        self.post(f"/juridiction/divisions/lignes/{l['id']}/suivi", {"etat": "reportée", "taux": "150", "ordre": "99",
                                                                     "echeance": "2027-06-30", "responsable": "X", "observations": "o"})
        l3 = self.q("SELECT * FROM jur_lignes_action WHERE id=?", (l["id"],))[0]
        self.assertEqual((l3["etat"], l3["taux"], l3["ordre"], l3["echeance"], l3["responsable"]), ("reportée", 100, 99, "2027-06-30", "X"))
        self.post(f"/juridiction/divisions/lignes/{l['id']}/suivi", {"etat": "invalide", "taux": "abc"})
        l4 = self.q("SELECT * FROM jur_lignes_action WHERE id=?", (l["id"],))[0]
        self.assertEqual((l4["etat"], l4["taux"]), ("reportée", 100))
        # taux de la division cohérent avec la synthèse
        with app.app_context():
            from flask import g
            with app.test_request_context("/"):
                g.user = {"id": 1, "role": "responsable"}
                import a_tableau
                H = a_tableau._H
                lignes = H.jur_lignes_action(code, self.annee)
        sy = AD.synthese_lignes(lignes)
        self.assertEqual(sy["total"], avant + 1)
        self.assertIn(AD.pct(sy["taux"]), txt(self.get(f"{base}?annee={self.annee}")))
        # suppression
        self.post(f"/juridiction/divisions/lignes/{l['id']}/supprimer")
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE id=?", (l["id"],))[0]["n"], 0)
        self.get(f"/juridiction/divisions/lignes/{l['id']}/modifier", 404)
        self.post(f"/juridiction/divisions/lignes/{l['id']}/supprimer", {}, 404)

    # ------------------------------------------------------------ 5. saisie du texte officiel de l'AOF
    def test_40_saisie_aof(self):
        code = self.redige
        t = txt(self.get(f"/juridiction/divisions/{code}"))
        self.assertIn("Saisir le texte de l'AOF", t)
        self.assertIn(S.SOURCE_GESCOL, t)
        self.post(f"/juridiction/divisions/{code}/aof", {"annee": self.annee, "article": "", "attributions": "x"})     # refusé
        self.assertEqual(self.q("SELECT source_aof FROM jur_services WHERE niveau=? AND code=?", (EDITION, code))[0]["source_aof"], S.SOURCE_GESCOL)
        texte = "Le Service est chargé de :\n- première attribution officielle ;\n- deuxième attribution officielle ;\n- troisième attribution officielle."
        self.post(f"/juridiction/divisions/{code}/aof", {"annee": ANNEE_2, "article": "Article 17", "attributions": texte})
        r = self.q("SELECT * FROM jur_services WHERE niveau=? AND code=?", (EDITION, code))[0]
        self.assertEqual((r["article_aof"], r["source_aof"], r["attributions"], r["reference"]), ("17", "AOF – article 17", texte, "Article 17"))
        t = txt(self.get(f"/juridiction/divisions/{code}?annee={ANNEE_2}"))
        self.assertIn("AOF – article 17", t)
        self.assertNotIn(S.SOURCE_GESCOL, t)
        self.assertNotIn("hors AOF", t)
        self.assertIn("Première attribution officielle", t)        # lignes d'action générées depuis le texte officiel
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, ANNEE_2))[0]["n"], 3)
        t = txt(self.get("/juridiction/divisions"))
        self.assertIn("AOF – article 17", t)
        # modification simple des attributions (route existante) conservée
        self.post(f"/juridiction/divisions/{code}", {"annee": self.annee, "attributions": "- attribution modifiée."})
        self.assertEqual(self.q("SELECT attributions FROM jur_services WHERE niveau=? AND code=?", (EDITION, code))[0]["attributions"],
                         "- attribution modifiée.")

    # ------------------------------------------------------------ 6. exports et import
    def test_50_exports_import(self):
        code = self.code
        base = f"/juridiction/divisions/{code}"
        for u in (f"{base}/plan.pdf?annee={self.annee}", f"/juridiction/divisions/plan.pdf?annee={self.annee}", "/juridiction/divisions.pdf"):
            self.assertTrue(self.get(u).data.startswith(b"%PDF"), u)
        r = self.get(f"{base}/plan.xlsx?annee={self.annee}")
        self.assertTrue(r.data.startswith(b"PK"))
        r2 = self.get(f"/juridiction/divisions/plan.xlsx?annee={self.annee}")
        self.assertTrue(r2.data.startswith(b"PK"))
        self.assertTrue(self.get(f"{base}/lignes/import?modele=1").data.startswith(b"PK"))
        self.get(f"{base}/lignes/import")
        if not openpyxl:
            return
        f = xlsx_feuilles(r.data)
        self.assertIn("Lignes d'action", f)
        nb = self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, self.annee))[0]["n"]
        self.assertEqual(sum(1 for l in f["Lignes d'action"] if l and l[2] and l[1] not in (None, "Ordre")), nb)
        f2 = xlsx_feuilles(r2.data)
        self.assertIn("Synthèse", f2)
        tot = next(l for l in f2["Synthèse"] if l and l[0] == "TOTAL STRUCTURE")
        self.assertEqual(tot[4], self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE niveau=? AND annee=?", (EDITION, self.annee))[0]["n"])
        # import : 1 nouvelle ligne + 1 mise à jour d'une ligne existante (même libellé)
        exist = self.q("SELECT * FROM jur_lignes_action WHERE service_code=? AND annee=? ORDER BY ordre LIMIT 1", (code, self.annee))[0]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(AD.ENT_LIGNES)
        ws.append(["", 50, "Ligne importée de test", "Indic", "Cible", "2027-05-31", "SALAMI Bio", "Archives ; Courrier (arrivée, départ)",
                   "réalisée", 100, "import"])
        ws.append(["", "", exist["libelle"], "Indicateur importé", "", "", "", "", "en cours", "55", ""])
        ws.append(["", "", "Ligne avec état faux", "", "", "", "", "", "inconnu", "", ""])
        buf = io.BytesIO()
        wb.save(buf)
        for _ in range(2):      # réimport sans doublon
            r = self.c.post(f"{base}/lignes/import", data={"annee": self.annee, "fichier": (io.BytesIO(buf.getvalue()), "plan.xlsx")},
                            content_type="multipart/form-data")
            self.assertEqual(r.status_code, 302)
        li = self.q("SELECT * FROM jur_lignes_action WHERE libelle='Ligne importée de test'")
        self.assertEqual(len(li), 1)
        self.assertEqual((li[0]["etat"], li[0]["taux"], li[0]["ordre"], li[0]["service_code"], li[0]["annee"]), ("réalisée", 100, 50, code, self.annee))
        self.assertEqual(sorted(AD._modules_liste(li[0]["modules"])), ["archives", "courrier"])
        e2 = self.q("SELECT * FROM jur_lignes_action WHERE id=?", (exist["id"],))[0]
        self.assertEqual((e2["indicateur"], e2["etat"], e2["taux"]), ("Indicateur importé", "en cours", 55))
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE libelle='Ligne avec état faux'")[0]["n"], 0)
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, self.annee))[0]["n"], nb + 1)
        # l'export se réimporte tel quel sans créer de ligne
        x = self.get(f"{base}/plan.xlsx?annee={self.annee}").data
        r = self.c.post(f"{base}/lignes/import", data={"annee": self.annee, "fichier": (io.BytesIO(x), "export.xlsx")},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.q("SELECT COUNT(*) n FROM jur_lignes_action WHERE service_code=? AND annee=?", (code, self.annee))[0]["n"], nb + 1)
        # fichier absent
        r = self.c.post(f"{base}/lignes/import", data={"annee": self.annee}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302)

    # ------------------------------------------------------------ 7. droits
    def test_60_droits(self):
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "agent", "mdp": "demo1234"}).status_code, 302)
        self.assertEqual(c2.get(f"/juridiction/divisions/{self.code}").status_code, 200)
        self.assertEqual(c2.get("/juridiction/divisions/plan").status_code, 200)
        self.assertEqual(c2.get(f"/juridiction/divisions/{self.code}/lignes/nouvelle").status_code, 403)
        self.assertEqual(c2.post(f"/juridiction/divisions/{self.code}/aof", data={"article": "1", "attributions": "x"}).status_code, 403)
        self.assertEqual(c2.post(f"/juridiction/divisions/{self.code}/lignes/generer", data={}).status_code, 403)
        self.assertEqual(c2.get(f"/juridiction/divisions/{self.code}/lignes/import").status_code, 403)
        self.assertNotIn("Régénérer depuis l'AOF", txt(c2.get(f"/juridiction/divisions/{self.code}")))


class TestEditionsSeparees(unittest.TestCase):
    """Rejoue la suite dans un processus séparé pour l'édition « cs » (si elle n'a pas tourné ici) et « ddemp »."""

    def test_cs(self):
        if EN_PLACE or EDITION == "ddemp":
            self.skipTest("édition déjà testée dans ce processus, ou processus secondaire")
        r = _rejouer("cs")
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])

    def test_ddemp(self):
        if EN_PLACE:
            self.skipTest("processus secondaire")
        r = _rejouer("ddemp")
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
