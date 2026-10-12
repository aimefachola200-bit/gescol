# -*- coding: utf-8 -*-
"""Tests des thèmes d'UP, de conférence pédagogique et de formation groupée (m_themes_up) — v1.6.
Lancer : python3 -m unittest tests.test_themes_up   (depuis /home/claude/gescol)

Édition « privee » (données demo.py, connexion directeur / demo1234) : préremplissage officiel 2026-2027
(12 + 12 + 5 + 5), dates, CRUD, recharger, réinitialiser, année suivante, Excel, import, PDF, ?retour=synoptique,
rôle enseignant en lecture. L'édition « cs » (demo_juridiction, responsable / demo1234, ?retour=jur_synoptique, select
des thèmes officiels dans le formulaire d'activité d'UP) est rejouée dans un sous-processus (TestEditionCS)."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "privee")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_themes_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import m_themes_up as T  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

EN_PLACE = config.EDITION == EDITION


def txt(r):
    return html_mod.unescape(r.data.decode())


def q(sql, params=(), un=False):
    with app.app_context():
        from app import q as _q
        return _q(sql, params, un=un)


def xlsx_feuilles(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def classeur(feuilles):
    """{nom: [entêtes] + lignes} -> octets xlsx."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nom, lignes in feuilles.items():
        ws = wb.create_sheet(nom)
        for l in lignes:
            ws.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================ fonctions pures
class TestCalculs(unittest.TestCase):
    def test_officiels(self):
        offi = T.themes_officiels("2026-2027")
        self.assertEqual(len(offi), 34)
        par = {c: [t for t in offi if t["categorie"] == c] for c, _l in T.CATEGORIES}
        self.assertEqual([len(par[c]) for c, _l in T.CATEGORIES], [12, 12, 5, 5])
        self.assertEqual(par["up_maternelle"][1]["date_administration"], "2026-10-23")
        self.assertEqual(par["up_primaire"][11]["date_administration"], "2027-05-07")
        self.assertEqual(par["up_primaire"][11]["theme"], "Bilan")
        self.assertEqual(par["up_primaire"][0]["date_administration"], "2026-10-09")
        self.assertEqual(par["up_maternelle"][4]["observations"], "Animation mixte")
        self.assertTrue(par["conference"][4]["theme"].startswith("L'utilisation pédagogique de l'intelligence artificielle"))
        self.assertIn("premier trimestre", par["formation_groupee"][0]["observations"])
        self.assertTrue(all(t["date_deroulement"] == "PM" for t in par["conference"] + par["formation_groupee"]))
        self.assertEqual(T.themes_officiels("2031-2032"), [])
        self.assertEqual(len(T.themes_officiels("2026-2027", "conference")), 5)

    def test_dates(self):
        self.assertEqual(T.date_affichee("2026-10-23"), "Vendredi 23/10/2026")
        self.assertEqual(T.date_affichee("PM"), "PM")
        self.assertEqual(T.normaliser_date("23/10/2026"), "2026-10-23")
        self.assertEqual(T.normaliser_date("Vendredi 23/10/2026"), "2026-10-23")
        self.assertEqual(T.normaliser_date(" pm "), "pm")
        self.assertEqual(T.decaler_un_an("2026-10-23"), "2027-10-23")
        self.assertEqual(T.decaler_un_an("2028-02-29"), "2029-02-28")
        self.assertEqual(T.decaler_un_an("PM"), "PM")
        self.assertEqual(T.trimestre_de(date(2026, 10, 9)), "I")
        self.assertEqual(T.trimestre_de(date(2027, 1, 15)), "II")
        self.assertEqual(T.trimestre_de(date(2027, 5, 7)), "III")
        self.assertEqual(T.annee_suivante("2026-2027"), "2027-2028")

    def test_suivi(self):
        lst = T.themes_officiels("2026-2027")
        lst[1]["execute"] = 1                      # UP maternelle n° 2 (23/10/2026)
        s = T.suivi(lst, aujourdhui=date(2026, 11, 1))
        m = s["up_maternelle"]
        self.assertEqual((m["total"], m["executes"], m["taux"]), (12, 1, 8))
        self.assertEqual(m["retard"], 1)           # n° 1 du 09/10/2026 non exécuté
        self.assertEqual(m["a_venir"], 10)
        self.assertEqual([t["programmes"] for t in m["trimestres"]], [5, 5, 2])
        self.assertEqual(m["trimestres"][0]["executes"], 1)
        self.assertEqual(m["prochain"]["numero"], 3)
        c = s["conference"]
        self.assertEqual((c["sans_date"], c["taux"]), (5, 0))

    def test_lire_theme(self):
        v, err = T.lire_theme({"categorie": "up_primaire", "theme": " Dictée ", "axes": "- a\n\n• b", "date_administration": "05/03/2027",
                               "date_preparation": "PM", "execute": "1", "date_execution": "2027-03-05", "participants": "12"})
        self.assertEqual(err, [])
        self.assertEqual((v["theme"], v["axes"], v["date_administration"], v["participants"]), ("Dictée", "a\nb", "2027-03-05", 12))
        v, err = T.lire_theme({"categorie": "conference", "theme": "", "date_execution": "hier", "participants": "x"})
        self.assertEqual(len(err), 3)
        self.assertEqual(T.categorie_depuis_texte("Formations groupées"), "formation_groupee")
        self.assertEqual(T.categorie_depuis_texte("UP Maternelle"), "up_maternelle")
        self.assertIsNone(T.categorie_depuis_texte("Feuil1"))


# ============================================================ édition privée (école)
class TestThemesUP(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not EN_PLACE or EDITION != "privee":
            raise unittest.SkipTest(f"édition « {config.EDITION} » : suite école rejouée seulement en édition privée")
        app.config["TESTING"] = True
        import demo
        demo.charger()
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "directeur", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code
        cls.annee = config.ANNEE_ACTIVE

    def get(self, url, code=200):
        r = self.c.get(url)
        self.assertEqual(r.status_code, code, f"{url} → {r.status_code}\n{r.data[:400]}")
        return r

    def post(self, url, data=None, code=302):
        r = self.c.post(url, data=data or {})
        self.assertEqual(r.status_code, code, f"{url} → {r.status_code}\n{r.data[:600]}")
        return r

    def test_01_preremplissage_et_pages(self):
        h = txt(self.get("/themes-up"))
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["n"], 34)
        for c, n in (("up_maternelle", 12), ("up_primaire", 12), ("conference", 5), ("formation_groupee", 5)):
            self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=? AND categorie=? AND officiel=1",
                               (self.annee, c), un=True)["n"], n)
        self.assertIn("Chasser les facteurs chronophages", h)
        self.assertIn("Vendredi 23/10/2026", h)
        self.assertIn("Animation libre", h)
        hp = txt(self.get("/themes-up?cat=up_primaire"))
        self.assertIn("23/10/2026", hp)
        self.assertIn("07/05/2027", hp)
        self.assertIn("Construction des figures symétriques", hp)
        hc = txt(self.get("/themes-up?cat=conference"))
        self.assertIn("ressources numériques", hc)
        self.assertIn("003/MEMP/DC/SGM/DPAF/DEM/DEP/DIIP/SP 041 SGG25", hc)
        hf = txt(self.get("/themes-up?cat=formation_groupee"))
        self.assertIn("Pédagogie Explicite", hf)
        hs = txt(self.get("/themes-up?cat=suivi"))
        self.assertIn("1er trimestre", hs)
        self.assertIn("0 / 12", hs)
        self.assertIn("Référence du texte", h)
        self.assertIn("DIIP", h)
        # menu
        self.assertIn("/themes-up", txt(self.get("/")))
        # année inconnue : vide
        self.assertIn("Aucun thème pour 2031-2032", txt(self.get("/themes-up?annee=2031-2032")))
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE annee='2031-2032'", un=True)["n"], 0)

    def test_02_retour_synoptique(self):
        h = txt(self.get("/themes-up?retour=synoptique"))
        self.assertIn("Tableau synoptique de l'école", h)
        self.assertIn("Thèmes d'UP", h)
        self.assertIn("/synoptique/reglement", h)
        self.assertIn("retour=synoptique", h)
        self.assertIn("Chasser les facteurs chronophages", h)
        h = txt(self.get("/themes-up?retour=synoptique&cat=suivi"))
        self.assertIn("Tableau synoptique de l'école", h)
        # l'onglet du synoptique pointe vers la page
        self.assertIn("/themes-up?retour=synoptique", txt(self.get("/synoptique")))

    def test_03_crud(self):
        self.get("/themes-up/theme/0?cat=up_primaire")
        self.get("/themes-up/theme/0?cat=conference")
        r = self.post("/themes-up/theme/0", {"categorie": "up_primaire", "theme": "", "axes": ""}, code=400)
        self.assertIn("obligatoire", txt(r))
        self.post("/themes-up/theme/0?cat=up_primaire", {"categorie": "up_primaire", "theme": "Thème local de test", "axes": "- Axe A\n- Axe B",
                                                          "date_preparation": "PM", "date_interpretation": "PM",
                                                          "date_administration": "21/05/2027", "observations": "Animation mixte"})
        t = q("SELECT * FROM themes_up WHERE theme='Thème local de test'", un=True)
        self.assertIsNotNone(t)
        self.assertEqual((t["numero"], t["officiel"], t["date_administration"], t["axes"]), (13, 0, "2027-05-21", "Axe A\nAxe B"))
        h = txt(self.get("/themes-up?cat=up_primaire"))
        self.assertIn("Thème local de test", h)
        self.assertIn("Vendredi 21/05/2027", h)
        self.assertIn("ajouté", h)
        # modification
        self.get(f"/themes-up/theme/{t['id']}")
        self.post(f"/themes-up/theme/{t['id']}", {"categorie": "up_primaire", "numero": "13", "theme": "Thème local modifié", "axes": "Axe C",
                                                   "date_administration": "2027-05-21", "execute": "1", "date_execution": "21/05/2027",
                                                   "lieu": "EPP Centre", "participants": "18", "compte_rendu": "RAS"})
        t2 = q("SELECT * FROM themes_up WHERE id=?", (t["id"],), un=True)
        self.assertEqual((t2["theme"], t2["execute"], t2["date_execution"], t2["participants"], t2["lieu"]),
                         ("Thème local modifié", 1, "2027-05-21", 18, "EPP Centre"))
        h = txt(self.get("/themes-up?cat=up_primaire"))
        self.assertIn("Exécuté le 21/05/2027", h)
        self.assertIn("18 participant(s)", h)
        # bascule exécuté
        self.post(f"/themes-up/theme/{t['id']}/executer")
        self.assertEqual(q("SELECT execute FROM themes_up WHERE id=?", (t["id"],), un=True)["execute"], 0)
        self.post(f"/themes-up/theme/{t['id']}/executer")
        self.assertEqual(q("SELECT execute FROM themes_up WHERE id=?", (t["id"],), un=True)["execute"], 1)
        hs = txt(self.get("/themes-up?cat=suivi"))
        self.assertIn("1 / 13", hs)
        # suppression
        self.post(f"/themes-up/theme/{t['id']}/supprimer")
        self.assertIsNone(q("SELECT 1 FROM themes_up WHERE id=?", (t["id"],), un=True))
        self.get(f"/themes-up/theme/{t['id']}", 404)
        self.assertTrue(q("SELECT 1 FROM journal WHERE action='themes_up' AND details LIKE '%supprimé%'", un=True))
        # référence
        self.post("/themes-up/meta", {"reference": "Réf. test", "source": "DIIP"})
        self.assertEqual(q("SELECT reference FROM themes_up_meta WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["reference"], "Réf. test")

    def test_04_recharger_reinitialiser(self):
        offi = q("SELECT id FROM themes_up WHERE ecole_id=1 AND annee=? AND categorie='up_maternelle' AND numero=2", (self.annee,), un=True)
        self.post(f"/themes-up/theme/{offi['id']}", {"categorie": "up_maternelle", "numero": "2", "theme": "Thème officiel altéré",
                                                      "date_administration": "23/10/2026"})
        self.post("/themes-up/theme/0", {"categorie": "conference", "theme": "Conférence locale", "date_deroulement": "12/12/2026"})
        self.post("/themes-up/recharger")
        self.assertIsNone(q("SELECT 1 FROM themes_up WHERE theme='Thème officiel altéré'", un=True))
        self.assertTrue(q("SELECT 1 FROM themes_up WHERE theme='Conférence locale'", un=True))
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=? AND officiel=1", (self.annee,), un=True)["n"], 34)
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["n"], 35)
        self.post("/themes-up/reinitialiser")
        self.assertIsNone(q("SELECT 1 FROM themes_up WHERE theme='Conférence locale'", un=True))
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["n"], 34)
        # année sans officiel : recharger refusé, réinitialiser vide
        self.post("/themes-up/theme/0?annee=2031-2032", {"categorie": "up_primaire", "theme": "X", "annee": "2031-2032"})
        r = self.c.post("/themes-up/recharger?annee=2031-2032", follow_redirects=True)
        self.assertIn("pas connus", txt(r))
        self.post("/themes-up/reinitialiser?annee=2031-2032")
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE annee='2031-2032'", un=True)["n"], 0)

    def test_05_annee_suivante(self):
        suiv = T.annee_suivante(self.annee)
        self.post("/themes-up/preparer")
        lst = q("SELECT * FROM themes_up WHERE ecole_id=1 AND annee=? ORDER BY categorie, numero", (suiv,))
        self.assertEqual(len(lst), 34)
        up2 = [t for t in lst if t["categorie"] == "up_maternelle" and t["numero"] == 2][0]
        self.assertEqual((up2["date_administration"], up2["officiel"], up2["execute"], up2["date_preparation"]), ("2027-10-23", 0, 0, "PM"))
        fg1 = [t for t in lst if t["categorie"] == "formation_groupee" and t["numero"] == 1][0]
        self.assertIn("2027-2028", fg1["observations"])
        h = txt(self.get(f"/themes-up?annee={suiv}"))
        self.assertIn("23/10/2027", h)
        r = self.c.post("/themes-up/preparer", follow_redirects=True)
        self.assertIn("contient déjà", txt(r))
        self.post(f"/themes-up/reinitialiser?annee={suiv}")
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE annee=?", (suiv,), un=True)["n"], 0)

    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_06_excel(self):
        r = self.get("/themes-up/export.xlsx")
        self.assertTrue(r.data.startswith(b"PK"))
        f = xlsx_feuilles(r.data)
        self.assertEqual(set(f), {"UP maternelle", "UP primaire", "Conférences pédagogiques", "Formations groupées"})
        lignes = [l for l in f["UP primaire"] if l and l[2]]
        self.assertTrue(any("Bilan" in str(l[2]) and "07/05/2027" in str(l[6]) for l in lignes))
        m = self.get("/themes-up/modele.xlsx")
        self.assertTrue(m.data.startswith(b"PK"))
        self.assertIn("Catégorie", str(xlsx_feuilles(m.data)))

    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_07_import(self):
        self.get("/themes-up/import")
        ent = T.ENTETES_XLSX
        data = classeur({"Conférences pédagogiques": [ent, ["", 1, "CP importée 1", "a\nb", "", "", "10/12/2026", "", "Oui"],
                                                      ["", 2, "CP importée 2", "c", "", "", "PM", "obs", "Non"]],
                         "Feuil1": [ent, ["Formations groupées", 1, "FG importée", "x", "", "", "15/01/2027", "", "", "Oui", "15/01/2027", "Salle", "20", "ok"]]})
        r = self.c.post("/themes-up/import", data={"fichier": (io.BytesIO(data), "themes.xlsx"), "reference": "Réf. import"},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("3 thème(s) importé(s)", txt(r))
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=? AND categorie='conference'", (self.annee,), un=True)["n"], 2)
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=? AND categorie='formation_groupee'", (self.annee,), un=True)["n"], 1)
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=1 AND annee=? AND categorie='up_primaire'", (self.annee,), un=True)["n"], 12)
        cp1 = q("SELECT * FROM themes_up WHERE theme='CP importée 1'", un=True)
        self.assertEqual((cp1["date_deroulement"], cp1["officiel"], cp1["axes"]), ("2026-12-10", 1, "a\nb"))
        fg = q("SELECT * FROM themes_up WHERE theme='FG importée'", un=True)
        self.assertEqual((fg["execute"], fg["date_execution"], fg["participants"], fg["lieu"]), (1, "2027-01-15", 20, "Salle"))
        self.assertEqual(q("SELECT reference FROM themes_up_meta WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["reference"], "Réf. import")
        # fichier sans colonne Thème / ligne invalide
        r = self.c.post("/themes-up/import", data={"fichier": (io.BytesIO(classeur({"X": [["A", "B"], [1, 2]]})), "x.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Aucune feuille exploitable", txt(r))
        r = self.c.post("/themes-up/import", data={"fichier": (io.BytesIO(classeur({"Feuil1": [ent, ["Inconnue", 1, "T"]]})), "x.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("catégorie inconnue", txt(r))
        self.post("/themes-up/reinitialiser")

    def test_08_pdf(self):
        for u in ("/themes-up/themes.pdf", "/themes-up/themes.pdf?partie=cp_fg", "/themes-up/themes.pdf?partie=complet",
                  "/themes-up/themes.pdf?retour=synoptique"):
            r = self.get(u)
            self.assertTrue(r.data.startswith(b"%PDF"), u)
        try:
            import pdfplumber
        except ImportError:  # pragma: no cover
            return
        with pdfplumber.open(io.BytesIO(self.get("/themes-up/themes.pdf?partie=complet").data)) as pdf:
            texte = "\n".join(p.extract_text() or "" for p in pdf.pages)
        self.assertIn("PLANNING D'EXÉCUTION DES THÈMES D'UNITÉ PÉDAGOGIQUE", texte)
        self.assertIn("CONFÉRENCE PÉDAGOGIQUE ET DE FORMATION GROUPÉE", texte)
        self.assertIn("Vendredi 23/10/2026", texte)
        self.assertIn("FORMATIONS GROUPÉES", texte)
        self.assertIn("041 SGG25", texte)

    def test_09_api_et_droits(self):
        with app.test_request_context("/themes-up"):
            lst = T._H_themes(self.annee)
            self.assertEqual(len(lst), 34)
            self.assertEqual(len(T._H_themes(self.annee, "up_primaire")), 12)
            self.assertEqual(len(T._H_themes("2031-2032")), 0)
        # enseignant : lecture seule
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "enseignant", "mdp": "demo1234"}).status_code, 302)
        h = txt(c2.get("/themes-up"))
        self.assertIn("Chasser les facteurs chronophages", h)
        self.assertNotIn("Ajouter un thème", h)
        self.assertIn("Consultation seule", h)
        self.assertEqual(c2.get("/themes-up/theme/0").status_code, 403)
        self.assertEqual(c2.post("/themes-up/recharger").status_code, 403)
        self.assertEqual(c2.post("/themes-up/reinitialiser").status_code, 403)
        self.assertEqual(c2.get("/themes-up/export.xlsx").status_code, 200)
        self.assertEqual(c2.get("/themes-up/themes.pdf").status_code, 200)
        # déconnecté
        self.assertEqual(app.test_client().get("/themes-up").status_code, 302)


# ============================================================ édition CS (sous-processus)
class TestThemesCS(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not EN_PLACE or EDITION != "cs":
            raise unittest.SkipTest("suite CS rejouée en sous-processus")
        app.config["TESTING"] = True
        import demo_juridiction
        demo_juridiction.main()
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code
        cls.annee = config.ANNEE_ACTIVE

    def test_cs(self):
        r = self.c.get("/juridiction/themes-up?retour=jur_synoptique")
        self.assertEqual(r.status_code, 200, r.data[:500])
        h = txt(r)
        self.assertIn("Tableau synoptique de la circonscription scolaire", h)
        self.assertIn("Chasser les facteurs chronophages", h)
        self.assertIn("planning commun de la structure", h)
        self.assertEqual(q("SELECT COUNT(*) n FROM themes_up WHERE ecole_id=0 AND annee=?", (self.annee,), un=True)["n"], 34)
        self.assertIn("themes-up?retour=jur_synoptique", txt(self.c.get("/juridiction/synoptique")))
        self.assertIn("themes-up", txt(self.c.get("/juridiction")))
        for u in ("/juridiction/themes-up?cat=suivi&retour=jur_synoptique", "/juridiction/themes-up/theme/0?retour=jur_synoptique",
                  "/juridiction/themes-up/import?retour=jur_synoptique"):
            self.assertEqual(self.c.get(u).status_code, 200, u)
        r = self.c.get("/juridiction/themes-up/themes.pdf?retour=jur_synoptique&partie=complet")
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.c.get("/juridiction/themes-up/export.xlsx?retour=jur_synoptique")
        self.assertTrue(r.data.startswith(b"PK"))
        # ajout dans le planning de la structure
        r = self.c.post("/juridiction/themes-up/theme/0?retour=jur_synoptique", data={"categorie": "conference", "theme": "CP de la CS", "date_deroulement": "05/12/2026"})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(q("SELECT 1 FROM themes_up WHERE ecole_id=0 AND theme='CP de la CS'", un=True))
        # select des thèmes officiels dans le formulaire d'activité d'UP
        up = q("SELECT id FROM jur_up ORDER BY id", un=True)
        self.assertIsNotNone(up)
        r = self.c.get(f"/juridiction/up/{up['id']}/activites/nouvelle")
        self.assertEqual(r.status_code, 200, r.data[:500])
        h = txt(r)
        self.assertIn("Thème officiel", h)
        self.assertIn("theme-officiel", h)
        self.assertIn("Chasser les facteurs chronophages", h)
        self.assertIn("CP de la CS", h)
        self.assertIn('<optgroup label="UP primaire">', h)
        r = self.c.post(f"/juridiction/up/{up['id']}/activites/nouvelle", data={"date": "2026-10-23", "type": "Animation pédagogique",
                                                                                "theme": T.UP_PRIMAIRE_2026[1][1], "prevue": "1"})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(q("SELECT 1 FROM jur_up_activites WHERE theme=?", (T.UP_PRIMAIRE_2026[1][1],), un=True))
        # rôle cp : écriture autorisée
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "cp1", "mdp": "demo1234"}).status_code, 302)
        self.assertEqual(c2.get("/juridiction/themes-up/theme/0?portee=structure").status_code, 200)


class TestEditionCS(unittest.TestCase):
    def test_sous_processus(self):
        if EDITION != "privee":
            raise unittest.SkipTest("rejoué seulement depuis l'édition privée")
        env = dict(os.environ, GESCOL_EDITION_TEST="cs")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "tests.test_themes_up.TestThemesCS"], cwd=RACINE, env=env,
                           capture_output=True, text=True, timeout=900)
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
