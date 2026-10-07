# -*- coding: utf-8 -*-
"""Tests du chantier « Données » v1.6 : modèles officiels EducMaster (notes primaire, élèves par classe, agents),
cycle maternelle sur /educmaster, suppression des élèves et des cartes scolaires.
Lancer : python3 -m unittest tests.test_donnees_v16   (depuis /home/claude/gescol)

L'édition testée est choisie par GESCOL_EDITION_TEST (privee par défaut) ; un sous-ensemble est rejoué dans un second
processus pour l'édition « publique » (TestEditionPublique)."""
import html as html_mod
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "privee")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_donnees_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def txt(r):
    return html_mod.unescape(r.data.decode())


def classeur(r):
    assert r.status_code == 200, r.status_code
    assert r.data.startswith(b"PK"), r.data[:50]
    return openpyxl.load_workbook(io.BytesIO(r.data))


@unittest.skipIf(openpyxl is None, "openpyxl absent")
class TestDonneesV16(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        cls.c = app.test_client()
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex, q
            ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'",
               (generate_password_hash("admin123"),))
            ex("DELETE FROM eleves")
            ex("DELETE FROM classes")
            ex("DELETE FROM personnel")
            cls.cm2 = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)", ("CM2 A", "CM2", 0, 1))
            cls.cm2b = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)", ("CM2 B", "CM2", 0, 1))
            cls.ci = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)", ("CI", "CI", 0, 1))
            cls.mat = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)",
                         ("Maternelle 2", "Maternelle 2", 0, 1))
            annee = config.ANNEE_ACTIVE

            def eleve(nom, prenoms, sexe, cid, em=None, nip=None):
                eid = ex("""INSERT INTO eleves (nom, prenoms, sexe, date_naissance, lieu_naissance, classe_id, statut,
                            date_inscription, annee, matricule_educmaster, nip, parent_telephone)
                            VALUES (?,?,?,?,?,?,'inscrit','2026-09-15',?,?,?,?)""",
                         (nom, prenoms, sexe, "2015-05-03", "Cotonou", cid, annee, em, nip, "01 97 00 00 00"))
                ex("UPDATE eleves SET matricule=? WHERE id=?", (f"T-{eid:04d}", eid))
                return eid
            cls.e1 = eleve("DOSSOU", "Grâce", "F", cls.cm2, "EM0001", "NIP0001")
            cls.e2 = eleve("HOUNGBO", "Marc", "M", cls.cm2, "EM0002", "NIP0002")
            cls.e3 = eleve("AKPO", "Sans Matricule", "M", cls.cm2, None, None)
            cls.e_ci = eleve("ZINSOU", "Bébé", "F", cls.ci, None, "NIP0010")
            cls.e_mat = eleve("KOTO", "Petit", "M", cls.mat, "EM0100", None)
            cls.e_paye = eleve("PAYEUR", "Jean", "M", cls.cm2b, None, None)
            ex("""INSERT INTO paiements (numero_recu, eleve_id, annee, montant, mode, date_paiement, annule)
                  VALUES ('R-0001',?,?,5000,'especes','2026-10-01',0)""", (cls.e_paye, annee))
            # évaluations : formative 1 et sommative 1 (CM2 A), sommative 1 (maternelle)
            mats = {m["code"]: m["id"] for m in q("SELECT id, code FROM matieres WHERE cycle='primaire'")}
            cls.mats = mats
            cls.ev_f1 = ex("""INSERT INTO evaluations (classe_id, annee, type, numero, libelle, date, matieres, cree_le)
                              VALUES (?,?,'formative',1,'Formative 1','2026-10-20',?,?)""",
                           (cls.cm2, annee, json.dumps([mats["DIC"], mats["MATH"], mats["EXP.EC"]]), "2026-10-20"))
            cls.ev_s1 = ex("""INSERT INTO evaluations (classe_id, annee, type, numero, libelle, date, matieres, cree_le)
                              VALUES (?,?,'sommative',1,'Sommative 1','2026-12-10',?,?)""",
                           (cls.cm2, annee, json.dumps([mats["DIC"], mats["MATH"], mats["COMP.EC"], mats["ED.SCI"],
                                                        mats["ED.SOC"], mats["ART.V"], mats["ART.P"], mats["EPS"]]),
                            "2026-12-10"))

            def note(eid, code, periode, valeur):
                ex("INSERT INTO notes (eleve_id, matiere_id, periode, annee, valeur) VALUES (?,?,?,?,?)",
                   (eid, mats[code], periode, annee, valeur))
            note(cls.e1, "DIC", "Formative 1", 14)
            note(cls.e1, "MATH", "Formative 1", 12.5)
            note(cls.e2, "DIC", "Formative 1", 9)
            note(cls.e1, "DIC", "Sommative 1", 15)
            note(cls.e1, "EPS", "Sommative 1", 18)
            note(cls.e1, "ART.V", "Sommative 1", 11)
            note(cls.e2, "COMP.EC", "Sommative 1", 13)
            # maternelle : appréciations
            mm = {m["code"]: m["id"] for m in q("SELECT id, code FROM matieres WHERE cycle='maternelle'")}
            cls.ev_mat = ex("""INSERT INTO evaluations (classe_id, annee, type, numero, libelle, date, matieres, cree_le)
                               VALUES (?,?,'sommative',1,'Sommative 1','2026-12-10',?,?)""",
                            (cls.mat, annee, json.dumps([mm["LANG"], mm["EV.MATH"]]), "2026-12-10"))
            ex("INSERT INTO notes (eleve_id, matiere_id, periode, annee, valeur, appreciation) VALUES (?,?,?,?,NULL,'A')",
               (cls.e_mat, mm["LANG"], "Sommative 1", annee))
            ex("INSERT INTO notes (eleve_id, matiere_id, periode, annee, valeur, appreciation) VALUES (?,?,?,?,NULL,'ECA')",
               (cls.e_mat, mm["EV.MATH"], "Sommative 1", annee))
            # personnel
            cls.p1 = ex("""INSERT INTO personnel (matricule, nom, prenoms, sexe, fonction, telephone, email, ecole_id, actif)
                           VALUES ('RH-0001','AGBO','Pierre','M','Enseignant(e)','01 96 00 00 00','agbo@ex.bj',1,1)""")
            ex("INSERT INTO personnel_carriere (personnel_id, date_prise_service) VALUES (?,?)", (cls.p1, "2010-10-01"))
        r = cls.c.post("/login", data={"identifiant": "admin", "mdp": "admin123"})
        assert r.status_code == 302, r.status_code

    @classmethod
    def tearDownClass(cls):
        """Remet la base à vide (les autres modules de tests partagent la même base lors d'un « discover »)."""
        with app.app_context():
            from app import ex
            for sql in ("DELETE FROM notes", "DELETE FROM evaluations", "DELETE FROM paiements", "DELETE FROM documents",
                        "DELETE FROM documents_modules WHERE cle LIKE '%educmaster%' OR cle LIKE '%:carte:%'",
                        "DELETE FROM eleves", "DELETE FROM classes", "DELETE FROM personnel_carriere", "DELETE FROM personnel"):
                ex(sql)

    def q(self, sql, params=()):
        with app.app_context():
            from app import q
            return q(sql, params)

    def q1(self, sql, params=()):
        with app.app_context():
            from app import q
            return q(sql, params, un=True)

    # ------------------------------------------------------------ pages et modèles vierges
    def test_01_pages(self):
        for cycle in ("maternelle", "primaire", "secondaire"):
            r = self.c.get(f"/educmaster?cycle={cycle}")
            self.assertEqual(r.status_code, 200)
            h = txt(r)
            self.assertIn("Modèles officiels EducMaster", h)
            self.assertIn("Maternelle", h)
        h = txt(self.c.get("/educmaster?cycle=primaire"))
        for mot in ("Modèle de notes (primaire)", "Modèle des élèves par classe", "Modèle des agents", "CM2 A",
                    "/educmaster/modeles/notes-primaire.xlsx", "/educmaster/modeles/classe.xlsx",
                    "/educmaster/modeles/agents.xlsx", "/educmaster/modeles/classe/rempli.xlsx",
                    "/educmaster/modeles/agents/rempli.xlsx", "Feuille « CM2 » → classe"):
            self.assertIn(mot, h)
        h = txt(self.c.get("/donnees"))
        self.assertIn("Modèles officiels EducMaster", h)
        self.assertIn("/educmaster/modeles/classe/rempli.xlsx", h)
        r = self.c.get("/educmaster/modeles")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.c.get("/educmaster/modeles/inconnu.xlsx").status_code, 404)

    def test_02_modeles_vierges(self):
        attendus = {"notes-primaire": ["Évaluation formative 1", "Évaluation sommative 1", "Évaluation sommative 2",
                                       "Évaluation sommative 3"],
                    "classe": ["CM2", "CM1", "CE2", "CE1", "CP", "CI", "Marternelle 2", "Marternelle 1"],
                    "agents": ["Mes agents"]}
        for cle, feuilles in attendus.items():
            r = self.c.get(f"/educmaster/modeles/{cle}.xlsx")
            wb = classeur(r)
            self.assertEqual(wb.sheetnames, feuilles, cle)
            self.assertIn("attachment", r.headers.get("Content-Disposition", ""))
            # identique au fichier du dossier modeles/
            with open(os.path.join(RACINE, "modeles", f"educmaster_modele_{cle.replace('-', '_')}.xlsx"), "rb") as fh:
                self.assertEqual(fh.read(), r.data)

    # ------------------------------------------------------------ a) modèle de notes rempli
    def test_03_modele_notes_rempli(self):
        r = self.c.get(f"/educmaster/modeles/notes-primaire/{self.cm2}.xlsx")
        self.assertIn(f"EducMaster_notes_cm2_a_{config.ANNEE_ACTIVE}.xlsx", r.headers.get("Content-Disposition", ""))
        wb = classeur(r)
        ref = openpyxl.load_workbook(os.path.join(RACINE, "modeles", "educmaster_modele_notes_primaire.xlsx"))
        self.assertEqual(wb.sheetnames, ref.sheetnames)
        for nom in ref.sheetnames:
            ws, wr = wb[nom], ref[nom]
            self.assertEqual(len(ws.merged_cells.ranges), 12, nom)
            self.assertEqual({str(x) for x in ws.merged_cells.ranges}, {str(x) for x in wr.merged_cells.ranges})
            for ligne in (1, 2):
                self.assertEqual([c.value for c in ws[ligne]], [c.value for c in wr[ligne]], f"{nom} ligne {ligne}")
            # en-têtes : police conservée
            self.assertEqual(ws["A1"].font.b, wr["A1"].font.b)
            # élèves (ordre nom, prénoms) : AKPO, DOSSOU, HOUNGBO
            self.assertEqual([ws.cell(row=i, column=2).value for i in (3, 4, 5)], ["AKPO", "DOSSOU", "HOUNGBO"])
            self.assertEqual([ws.cell(row=i, column=1).value for i in (3, 4, 5)], [None, "EM0001", "EM0002"])
            self.assertEqual(ws["C4"].value, "Grâce")
            self.assertIsNone(ws.cell(row=6, column=2).value)
        f1 = wb["Évaluation formative 1"]
        # Dictée = D/E, Math = F/G, Expression écrite = H/I
        self.assertEqual(f1["D4"].value, 14)
        self.assertEqual(f1["E4"].value, 20)
        self.assertEqual(f1["F4"].value, 12.5)
        self.assertEqual(f1["G4"].value, 20)
        self.assertIsNone(f1["H4"].value)
        self.assertEqual(f1["D5"].value, 9)
        self.assertIsNone(f1["D3"].value)
        s1 = wb["Évaluation sommative 1"]
        self.assertEqual(s1["D4"].value, 15)       # Dictée
        self.assertEqual(s1["T4"].value, 18)       # EPS
        self.assertEqual(s1["P4"].value, 11)       # EA (Oral) = ART.V
        self.assertIsNone(s1["R4"].value)          # EA (Dessin) = ART.P : pas de note
        self.assertEqual(s1["J5"].value, 13)       # Compréhension de l'écrit
        self.assertIsNone(s1["F4"].value)          # pas de note de maths en sommative
        s2 = wb["Évaluation sommative 2"]          # pas d'évaluation : élèves seuls
        self.assertEqual(s2["B4"].value, "DOSSOU")
        self.assertIsNone(s2["D4"].value)
        # classe non primaire refusée, classe inconnue 404
        self.assertEqual(self.c.get(f"/educmaster/modeles/notes-primaire/{self.mat}.xlsx").status_code, 302)
        self.assertIn(self.c.get("/educmaster/modeles/notes-primaire/99999.xlsx").status_code, (403, 404))

    # ------------------------------------------------------------ b) modèle classe : aller-retour
    def test_04_modele_classe_aller_retour(self):
        r = self.c.get("/educmaster/modeles/classe/rempli.xlsx")
        wb = classeur(r)
        self.assertEqual(wb.sheetnames, ["CM2", "CM1", "CE2", "CE1", "CP", "CI", "Marternelle 2", "Marternelle 1"])
        cm2 = wb["CM2"]
        self.assertEqual([c.value for c in cm2[1]][:8], ["Nom", "Prénom ( s )", "Sexe", "Date de naissance",
                                                          "Lieu de naissance", "Téléphone de parents", "Nationalité", "Npi"])
        noms = [cm2.cell(row=i, column=1).value for i in range(2, 7)]
        self.assertEqual(noms[:4], ["AKPO", "DOSSOU", "HOUNGBO", "PAYEUR"])   # CM2 A puis CM2 B
        self.assertEqual(cm2["B3"].value, "Grâce")
        self.assertEqual(cm2["D3"].value, "03/05/2015")
        self.assertEqual(cm2["G3"].value, "Béninoise")
        self.assertEqual(cm2["H3"].value, "NIP0001")
        self.assertEqual(cm2["F3"].value, "01 97 00 00 00")
        self.assertEqual(wb["CI"]["A2"].value, "ZINSOU")
        self.assertEqual(wb["Marternelle 2"]["A2"].value, "KOTO")
        self.assertIsNone(wb["CM1"]["A2"].value)
        # modification : nouveau lieu + nationalité pour DOSSOU (retrouvée par NIP), nouvel élève en CI, nouvel en CM2 B
        cm2["E3"].value = "Porto-Novo"
        cm2["G3"].value = "Togolaise"
        ci = wb["CI"]
        ci["A3"].value, ci["B3"].value, ci["C3"].value, ci["D3"].value = "NOUVEAU", "Élève", "F", "12/01/2020"
        ci["H3"].value = "NIP0099"
        cm2["A7"].value, cm2["B7"].value, cm2["C7"].value = "AJOUTE", "En CM2 B", "M"
        buf = io.BytesIO()
        wb.save(buf)
        # simulation : rien n'est enregistré
        r = self.c.post("/educmaster/modeles/classe/importer",
                        data={"fichier": (io.BytesIO(buf.getvalue()), "classe.xlsx"), "simuler": "1", "classe_cm2": str(self.cm2b)},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Simulation : 1 élève(s) mis à jour, 2 créé(s)", txt(r))
        self.assertIsNone(self.q1("SELECT 1 FROM eleves WHERE nom='NOUVEAU'"))
        # import réel
        r = self.c.post("/educmaster/modeles/classe/importer",
                        data={"fichier": (io.BytesIO(buf.getvalue()), "classe.xlsx"), "classe_cm2": str(self.cm2b)},
                        content_type="multipart/form-data", follow_redirects=True)
        h = txt(r)
        self.assertIn("1 élève(s) mis à jour, 2 créé(s)", h)
        self.assertIn("Dernier import de modèle", h)
        e1 = self.q1("SELECT * FROM eleves WHERE id=?", (self.e1,))
        self.assertEqual(e1["lieu_naissance"], "Porto-Novo")
        self.assertEqual(e1["nationalite"], "Togolaise")
        n = self.q1("SELECT * FROM eleves WHERE nom='NOUVEAU'")
        self.assertEqual(n["classe_id"], self.ci)
        self.assertEqual(n["date_naissance"], "2020-01-12")
        self.assertEqual(n["nip"], "NIP0099")
        self.assertTrue(n["matricule"].endswith(f"P-{n['id']:04d}"))
        a = self.q1("SELECT * FROM eleves WHERE nom='AJOUTE'")
        self.assertEqual(a["classe_id"], self.cm2b)
        # réimport : tout inchangé
        r = self.c.post("/educmaster/modeles/classe/importer",
                        data={"fichier": (io.BytesIO(buf.getvalue()), "classe.xlsx"), "classe_cm2": str(self.cm2b)},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("0 élève(s) mis à jour, 0 créé(s)", txt(r))
        # mauvais fichier
        r = self.c.post("/educmaster/modeles/classe/importer", data={"fichier": (io.BytesIO(b"bonjour"), "x.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("n'est pas un classeur Excel", txt(r))
        with app.app_context():
            from app import ex
            ex("DELETE FROM eleves WHERE nom IN ('NOUVEAU','AJOUTE')")

    # ------------------------------------------------------------ c) agents : aller-retour
    def test_05_modele_agents_aller_retour(self):
        wb = classeur(self.c.get("/educmaster/modeles/agents/rempli.xlsx"))
        ws = wb["Mes agents"]
        self.assertEqual(ws["A1"].value, "Matricule")
        self.assertEqual([ws.cell(row=2, column=j).value for j in (1, 3, 4, 5, 8, 9, 10, 12)],
                         ["RH-0001", "AGBO", "Pierre", "M", "Bénin", "01 96 00 00 00", "agbo@ex.bj", "01/10/2010"])
        ws["B2"].value = "ANPE123"
        ws["F2"].value = "15/06/1980"
        ws["G2"].value = "Abomey"
        ws["K2"].value = "IFU001"
        ws["A3"].value, ws["C3"].value, ws["D3"].value, ws["E3"].value = "RH-0002", "TOSSOU", "Clarisse", "F"
        ws["L3"].value = "01/10/2015"
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post("/educmaster/modeles/agents/importer", data={"fichier": (io.BytesIO(buf.getvalue()), "agents.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("1 agent(s) mis à jour, 1 créé(s)", txt(r))
        p1 = self.q1("SELECT * FROM personnel WHERE id=?", (self.p1,))
        self.assertEqual((p1["id_anpe"], p1["date_naissance"], p1["lieu_naissance"], p1["ifu"]),
                         ("ANPE123", "1980-06-15", "Abomey", "IFU001"))
        p2 = self.q1("SELECT p.*, c.date_prise_service FROM personnel p LEFT JOIN personnel_carriere c ON c.personnel_id=p.id "
                     "WHERE p.nom='TOSSOU'")
        self.assertEqual((p2["matricule"], p2["sexe"], p2["ecole_id"], p2["fonction"], p2["date_prise_service"]),
                         ("RH-0002", "F", 1, "Enseignant(e)", "2015-10-01"))
        r = self.c.post("/educmaster/modeles/agents/importer", data={"fichier": (io.BytesIO(buf.getvalue()), "agents.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("0 agent(s) mis à jour, 0 créé(s), 2 inchangé(s)", txt(r))
        # un modèle de classe déposé à la place du modèle des agents est refusé
        wb2 = classeur(self.c.get("/educmaster/modeles/classe/rempli.xlsx"))
        b2 = io.BytesIO()
        wb2.save(b2)
        r = self.c.post("/educmaster/modeles/agents/importer", data={"fichier": (io.BytesIO(b2.getvalue()), "classe.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Aucune feuille de ce classeur ne correspond", txt(r))

    # ------------------------------------------------------------ cycle maternelle : dépôt + élèves + appréciations
    def test_06_educmaster_maternelle(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Maternelle 2"
        ws.append(["Matricule", "Nom", "Prénoms", "Sexe", "Date de naissance", "Classe", "Langage et communication",
                   "Éveil mathématique", "Motricité"])
        ws.append(["EM0100", "KOTO", "Petit", "M", "2021-03-02", "Maternelle 2", None, None, None])
        ws.append(["EM0101", "NOUVELLE", "Fille", "F", "2021-07-09", "Maternelle 2", None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post("/educmaster/deposer", data={"cycle": "maternelle", "fichier": (io.BytesIO(buf.getvalue()), "mat.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        h = txt(r)
        self.assertEqual(r.status_code, 200)
        self.assertIn("3 colonne(s) de matières", h)
        self.assertIn("Langage et communication", h)
        self.assertIn("A = Acquis, ECA = En cours", h)
        r = self.c.post("/educmaster/eleves", data={"cycle": "maternelle", "creer": "1"}, follow_redirects=True)
        self.assertIn("1 créé(s)", txt(r))
        n = self.q1("SELECT * FROM eleves WHERE nom='NOUVELLE'")
        self.assertEqual(n["classe_id"], self.mat)
        self.assertTrue(n["matricule"].endswith(f"M-{n['id']:04d}"))
        r = self.c.post("/educmaster/notes", data={"cycle": "maternelle", "evaluation": str(self.ev_mat)})
        wb2 = classeur(r)
        ws2 = wb2["Maternelle 2"]
        self.assertEqual(ws2["G2"].value, "A")
        self.assertEqual(ws2["H2"].value, "ECA")
        self.assertIsNone(ws2["I2"].value)
        self.assertIsNone(ws2["G3"].value)
        self.c.post("/educmaster/retirer", data={"cycle": "maternelle"})
        with app.app_context():
            from app import ex
            ex("DELETE FROM eleves WHERE nom='NOUVELLE'")

    # ------------------------------------------------------------ suppression des élèves
    def test_07_suppression_eleve(self):
        h = txt(self.c.get("/eleves"))
        self.assertIn("Supprimer la sélection", h)
        self.assertIn(f"/eleves/{self.e3}/supprimer", h)
        h = txt(self.c.get(f"/eleves/{self.e3}/modifier"))
        self.assertIn("Supprimer définitivement", h)
        self.assertIn('name="nationalite"', h)
        # avec paiement : refus, puis « retirer »
        r = self.c.post(f"/eleves/{self.e_paye}/supprimer", follow_redirects=True)
        self.assertIn("Suppression impossible", txt(r))
        self.assertIn("1 paiement(s) non annulé(s)", txt(r))
        self.assertIsNotNone(self.q1("SELECT 1 FROM eleves WHERE id=?", (self.e_paye,)))
        r = self.c.post(f"/eleves/{self.e_paye}/supprimer", data={"mode": "retirer"}, follow_redirects=True)
        self.assertIn("retiré(e)", txt(r))
        self.assertEqual(self.q1("SELECT statut FROM eleves WHERE id=?", (self.e_paye,))["statut"], "sorti")
        # sans paiement : suppression physique, notes comprises
        self.assertTrue(self.q("SELECT 1 FROM notes WHERE eleve_id=?", (self.e2,)))
        r = self.c.post(f"/eleves/{self.e2}/supprimer", follow_redirects=True)
        self.assertIn("supprimé définitivement", txt(r))
        self.assertIsNone(self.q1("SELECT 1 FROM eleves WHERE id=?", (self.e2,)))
        self.assertFalse(self.q("SELECT 1 FROM notes WHERE eleve_id=?", (self.e2,)))
        self.assertEqual(self.c.post("/eleves/99999/supprimer").status_code, 404)
        # sélection : un supprimable + un bloqué (document officiel émis)
        with app.app_context():
            from app import ex
            ex("INSERT INTO documents (code, type, eleve_id, libelle, empreinte, cree_le) VALUES ('ABC12','bulletin',?,'B','x','2026-10-01')",
               (self.e1,))
        r = self.c.post("/eleves/supprimer-selection", data={"ids": [str(self.e1), str(self.e3)]}, follow_redirects=True)
        h = txt(r)
        self.assertIn("1 élève(s) supprimé(s)", h)
        self.assertIn("DOSSOU Grâce", h)
        self.assertIsNone(self.q1("SELECT 1 FROM eleves WHERE id=?", (self.e3,)))
        self.assertIsNotNone(self.q1("SELECT 1 FROM eleves WHERE id=?", (self.e1,)))
        r = self.c.post("/eleves/supprimer-selection", data={"ids": [str(self.e1)], "mode": "retirer"}, follow_redirects=True)
        self.assertIn("1 élève(s) retiré(s)", txt(r))
        self.assertEqual(self.q1("SELECT statut FROM eleves WHERE id=?", (self.e1,))["statut"], "sorti")
        r = self.c.post("/eleves/supprimer-selection", data={}, follow_redirects=True)
        self.assertIn("Cochez au moins un élève", txt(r))

    # ------------------------------------------------------------ suppression d'une carte
    def test_08_suppression_carte(self):
        with app.app_context():
            from app import ex
            ex("UPDATE eleves SET statut='inscrit' WHERE id=?", (self.e1,))
        r = self.c.get(f"/carte?eleve={self.e1}")
        self.assertEqual(r.status_code, 200)
        h = txt(r)
        self.assertIn("Supprimer la carte", h)
        self.assertIn(f"/carte/{self.e1}/supprimer", h)
        r = self.c.post(f"/carte/{self.e1}/enregistrer", data={"nom": "DOSSOU", "prenoms": "Grâce Corrigée", "sexe": "F",
                                                              "numero_table": "T12", "annee": config.ANNEE_ACTIVE,
                                                              "directeur_titre": "Directeur", "cachet_couleur": "bleu"},
                        follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(self.q("SELECT 1 FROM documents WHERE eleve_id=? AND type='carte_candidat'", (self.e1,)))
        self.assertTrue(self.q("SELECT 1 FROM documents_modules WHERE cle LIKE ?", (f"%:carte:eleve:{self.e1}",)))
        # (sans suivre la redirection : l'éditeur attribue un nouveau code dès qu'il affiche la carte)
        r = self.c.post(f"/carte/{self.e1}/supprimer")
        self.assertEqual(r.status_code, 302)
        self.assertFalse(self.q("SELECT 1 FROM documents WHERE eleve_id=? AND type='carte_candidat'", (self.e1,)))
        self.assertFalse(self.q("SELECT 1 FROM documents_modules WHERE cle LIKE ?", (f"%:carte:eleve:{self.e1}",)))
        self.assertIsNotNone(self.q1("SELECT 1 FROM eleves WHERE id=?", (self.e1,)))   # l'élève reste
        self.assertEqual(self.q1("SELECT numero_table FROM eleves WHERE id=?", (self.e1,))["numero_table"], "T12")
        self.assertIn("supprimée", txt(self.c.get("/carte?classe=" + str(self.cm2))))
        self.assertEqual(self.c.post("/carte/99999/supprimer").status_code, 404)


class TestEditionPublique(unittest.TestCase):
    """Rejoue un sous-ensemble (pages, modèles, suppression) dans l'édition publique, dans un processus séparé."""

    def test_publique(self):
        if EDITION != "privee":
            self.skipTest("déjà dans une autre édition")
        env = dict(os.environ, GESCOL_EDITION_TEST="publique")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "-q",
                            "tests.test_donnees_v16.TestDonneesV16.test_01_pages",
                            "tests.test_donnees_v16.TestDonneesV16.test_03_modele_notes_rempli",
                            "tests.test_donnees_v16.TestDonneesV16.test_04_modele_classe_aller_retour",
                            "tests.test_donnees_v16.TestDonneesV16.test_05_modele_agents_aller_retour",
                            "tests.test_donnees_v16.TestDonneesV16.test_07_suppression_eleve",
                            "tests.test_donnees_v16.TestDonneesV16.test_08_suppression_carte"],
                           cwd=RACINE, env=env, capture_output=True, text=True, timeout=600)
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
