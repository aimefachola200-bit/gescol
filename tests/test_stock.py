# -*- coding: utf-8 -*-
"""Tests du module Vivres et stock de la cantine (m_stock_cantine).
Lancer : python3 -m unittest tests.test_stock   (depuis /home/claude/gescol)

L'édition testée est choisie par GESCOL_EDITION_TEST (privee par défaut) ; la suite est rejouée dans un second
processus pour l'édition « publique » (TestEditionPublique)."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date, timedelta

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "privee")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_stock_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import m_stock_cantine as S  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def txt(r):
    return html_mod.unescape(r.data.decode())


def q(sql, params=(), un=False):
    with app.app_context():
        from app import q as _q
        return _q(sql, params, un=un)


class TestCalculs(unittest.TestCase):
    """Règles pures : stock, alertes, consommation."""

    def mv(self, typ, qv, d, **k):
        base = {"id": self.n, "type": typ, "quantite": qv, "date": d, "qualite": None, "prix_unitaire": None,
                "date_peremption": None, "rationnaires": None}
        self.n += 1
        base.update(k)
        return base

    def setUp(self):
        self.n = 1

    def test_stock(self):
        m = [self.mv("entree", 100, "2026-10-01"), self.mv("entree", 50, "2026-10-02"), self.mv("sortie", 30, "2026-10-03"),
             self.mv("perte", 5, "2026-10-04")]
        self.assertEqual(S.calculer_stock(m), 115)
        m.append(self.mv("inventaire", 110, "2026-10-05"))
        self.assertEqual(S.calculer_stock(m), 110)
        m.append(self.mv("sortie", 10, "2026-10-06"))
        self.assertEqual(S.calculer_stock(m), 100)
        # l'ordre d'insertion ne compte pas : seul l'ordre chronologique
        self.assertEqual(S.calculer_stock(list(reversed(m))), 100)

    def test_alertes(self):
        auj = date(2026, 10, 7)
        v = {"id": 1, "nom": "Riz", "unite": "kg", "seuil_alerte": 20, "categorie": "cereale"}
        e = S.etat_vivre(v, [self.mv("entree", 100, "2026-10-01", qualite="bonne", prix_unitaire=450)], auj)
        self.assertEqual(e["alertes"], [])
        self.assertEqual(e["valeur"], 45000)
        self.assertEqual(e["qualite"], "bonne")
        # seuil
        e = S.etat_vivre(v, [self.mv("entree", 20, "2026-10-01")], auj)
        self.assertTrue(any("seuil" in a for a in e["alertes"]))
        # qualité
        e = S.etat_vivre(v, [self.mv("entree", 100, "2026-10-01", qualite="bonne"), self.mv("inventaire", 90, "2026-10-02", qualite="avariee")], auj)
        self.assertEqual(e["qualite"], "avariee")
        self.assertTrue(any("avariée" in a for a in e["alertes"]))
        # péremption dépassée / proche
        e = S.etat_vivre(v, [self.mv("entree", 100, "2026-10-01", date_peremption="2026-09-30")], auj)
        self.assertTrue(any("dépassée" in a for a in e["alertes"]))
        e = S.etat_vivre(v, [self.mv("entree", 100, "2026-10-01", date_peremption=(auj + timedelta(days=10)).isoformat())], auj)
        self.assertTrue(any("Périme dans 10 jour" in a for a in e["alertes"]))
        e = S.etat_vivre(v, [self.mv("entree", 100, "2026-10-01", date_peremption=(auj + timedelta(days=90)).isoformat())], auj)
        self.assertEqual(e["alertes"], [])
        # stock négatif
        e = S.etat_vivre(v, [self.mv("sortie", 5, "2026-10-01")], auj)
        self.assertTrue(any("négatif" in a for a in e["alertes"]))

    def test_consommation(self):
        m = [self.mv("entree", 100, "2026-10-01"), self.mv("sortie", 10, "2026-10-02", rationnaires=200),
             self.mv("sortie", 2, "2026-10-02", rationnaires=200), self.mv("sortie", 8, "2026-10-03", rationnaires=160)]
        c = S.consommation(m, 80)
        self.assertEqual(c["jours"], 2)
        self.assertEqual(c["par_jour"], 10)
        self.assertEqual(c["rations"], 360)
        self.assertAlmostEqual(c["par_rationnaire"], 20 / 360, places=4)
        self.assertEqual(c["autonomie"], 8)


class TestStock(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        cls.c = app.test_client()
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex
            ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
            ex("UPDATE ecoles SET nom=?, departement=?, cs=? WHERE id=1", ("EPP TEST STOCK", "Zou", "CS Abomey"))
        r = cls.c.post("/login", data={"identifiant": "admin", "mdp": "admin123"})
        assert r.status_code == 302, r.status_code
        cls.annee = config.ANNEE_ACTIVE
        cls.y = int(cls.annee[:4])

    def d(self, mois, jour):
        return date(self.y + (mois < 9), mois, jour).isoformat()

    def test_01_pages_vides(self):
        for url in ("/cantine/stock", "/cantine/stock/vivres", "/cantine/stock/mouvements", "/cantine/stock/recap",
                    "/cantine/stock/import", "/cantine/stock/vivres/0", "/cantine/stock/mouvements/0?type=entree",
                    "/cantine/stock/sortie-jour"):
            r = self.c.get(url)
            self.assertEqual(r.status_code, 200, f"{url} -> {r.status_code}")
        self.assertIn("Aucun vivre", txt(self.c.get("/cantine/stock")))
        self.assertIn("/cantine/stock", txt(self.c.get("/")))      # lien de menu
        self.assertTrue(self.c.get("/cantine/stock/etat.pdf").data.startswith(b"%PDF"))

    def test_02_vivres(self):
        r = self.c.post("/cantine/stock/vivres/0", data={"nom": "Riz", "categorie": "cereale", "unite": "kg", "seuil_alerte": "20"}, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Vivre « Riz » ajouté", txt(r))
        r = self.c.post("/cantine/stock/vivres/0", data={"nom": "riz", "unite": "kg"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("existe déjà", txt(r))
        r = self.c.post("/cantine/stock/vivres/0", data={"nom": "", "unite": "kg"})
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/cantine/stock/vivres/0", data={"nom": "Huile", "categorie": "huile", "unite": "litre", "seuil_alerte": "5"}, follow_redirects=True)
        self.assertIn("Huile", txt(r))
        riz = q("SELECT * FROM cantine_vivres WHERE nom='Riz'", un=True)
        self.assertEqual(riz["ecole_id"], 1)
        self.assertEqual(riz["seuil_alerte"], 20)
        # modification
        r = self.c.post(f"/cantine/stock/vivres/{riz['id']}", data={"nom": "Riz blanc", "categorie": "cereale", "unite": "kg", "seuil_alerte": "25", "actif": "1"}, follow_redirects=True)
        self.assertIn("Riz blanc", txt(r))
        self.assertEqual(q("SELECT seuil_alerte FROM cantine_vivres WHERE id=?", (riz["id"],), un=True)["seuil_alerte"], 25)
        # suppression d'un vivre sans mouvement
        r = self.c.post("/cantine/stock/vivres/0", data={"nom": "Temporaire", "unite": "kg"}, follow_redirects=True)
        tid = q("SELECT id FROM cantine_vivres WHERE nom='Temporaire'", un=True)["id"]
        r = self.c.post(f"/cantine/stock/vivres/{tid}/supprimer", follow_redirects=True)
        self.assertIn("supprimé", txt(r))
        self.assertIsNone(q("SELECT id FROM cantine_vivres WHERE id=?", (tid,), un=True))
        self.assertEqual(self.c.post(f"/cantine/stock/vivres/{tid}/supprimer").status_code, 404)
        # liste type
        r = self.c.post("/cantine/stock/vivres/liste-type", follow_redirects=True)
        self.assertIn("Liste type chargée", txt(r))
        self.assertTrue(q("SELECT id FROM cantine_vivres WHERE nom='Maïs'", un=True))
        n1 = q("SELECT COUNT(*) n FROM cantine_vivres", un=True)["n"]
        self.c.post("/cantine/stock/vivres/liste-type", follow_redirects=True)
        self.assertEqual(q("SELECT COUNT(*) n FROM cantine_vivres", un=True)["n"], n1)   # pas de doublon

    def test_03_mouvements_et_stock(self):
        riz = q("SELECT * FROM cantine_vivres WHERE nom='Riz blanc'", un=True)
        vid = riz["id"]

        def mvt(typ, qv, d, **k):
            data = {"type": typ, "vivre_id": vid, "quantite": qv, "date": d}
            data.update(k)
            r = self.c.post("/cantine/stock/mouvements/0", data=data, follow_redirects=True)
            self.assertEqual(r.status_code, 200, txt(r)[:500])
            return txt(r)

        h = mvt("entree", "100", self.d(10, 1), qualite="bonne", origine="PNASI", prix_unitaire="450", bon_numero="BL 1",
                date_peremption=self.d(6, 30))
        self.assertIn("Entrée enregistrée", h)
        self.assertIn("stock : 100 kg", h)
        mvt("entree", "50", self.d(10, 2), qualite="bonne")
        mvt("sortie", "30", self.d(10, 3), rationnaires="200")
        h = mvt("perte", "5", self.d(10, 4))
        self.assertIn("stock : 115 kg", h)
        h = mvt("inventaire", "110", self.d(10, 5), qualite="acceptable")
        self.assertIn("stock : 110 kg", h)
        h = mvt("sortie", "10", self.d(10, 6), rationnaires="180")
        self.assertIn("stock : 100 kg", h)
        html = txt(self.c.get("/cantine/stock"))
        self.assertIn("Acceptable", html)
        self.assertIn("45 000", html)          # valeur = 100 × 450
        self.assertNotIn("⚠", html)
        # erreurs de saisie
        r = self.c.post("/cantine/stock/mouvements/0", data={"type": "entree", "vivre_id": vid, "quantite": "abc", "date": self.d(10, 7)})
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/cantine/stock/mouvements/0", data={"type": "sortie", "vivre_id": vid, "quantite": "0", "date": self.d(10, 7)})
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/cantine/stock/mouvements/0", data={"type": "entree", "vivre_id": 999999, "quantite": "5", "date": self.d(10, 7)})
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/cantine/stock/mouvements/0", data={"type": "entree", "vivre_id": vid, "quantite": "5", "date": "2099-01-01"})
        self.assertEqual(r.status_code, 400)
        # modification d'un mouvement : la perte passe de 5 à 15 -> après inventaire le stock reste 100
        p = q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=? AND type='perte'", (vid,), un=True)
        r = self.c.post(f"/cantine/stock/mouvements/{p['id']}", data={"type": "perte", "vivre_id": vid, "quantite": "15", "date": self.d(10, 4)}, follow_redirects=True)
        self.assertIn("mise à jour", txt(r))
        with app.app_context():
            s = S.calculer_stock(q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=?", (vid,)))
        self.assertEqual(s, 100)
        # suppression de l'inventaire -> 100 + 50 − 30 − 15 − 10 = 95
        inv = q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=? AND type='inventaire'", (vid,), un=True)
        r = self.c.post(f"/cantine/stock/mouvements/{inv['id']}/supprimer", follow_redirects=True)
        self.assertIn("recalculé", txt(r))
        s = S.calculer_stock(q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=?", (vid,)))
        self.assertEqual(s, 95)
        self.assertEqual(self.c.post(f"/cantine/stock/mouvements/{inv['id']}/supprimer").status_code, 404)
        # suppression du vivre refusée (mouvements) ; désactivation possible
        r = self.c.post(f"/cantine/stock/vivres/{vid}/supprimer", follow_redirects=True)
        self.assertIn("Suppression refusée", txt(r))
        self.assertTrue(q("SELECT id FROM cantine_vivres WHERE id=?", (vid,), un=True))
        r = self.c.post(f"/cantine/stock/vivres/{vid}/supprimer", data={"desactiver": "1"}, follow_redirects=True)
        self.assertIn("désactivé", txt(r))
        self.assertNotIn("<b>Riz blanc</b>", txt(self.c.get("/cantine/stock")))       # plus dans l'état du stock
        self.assertIn("Riz blanc", txt(self.c.get("/cantine/stock/mouvements")))        # mais toujours dans le journal
        r = self.c.post(f"/cantine/stock/vivres/{vid}/reactiver", follow_redirects=True)
        self.assertIn("réactivé", txt(r))
        self.assertIn("<b>Riz blanc</b>", txt(self.c.get("/cantine/stock")))

    def test_04_alertes_et_synthese(self):
        huile = q("SELECT * FROM cantine_vivres WHERE nom='Huile'", un=True)
        r = self.c.post("/cantine/stock/mouvements/0", data={"type": "entree", "vivre_id": huile["id"], "quantite": "4", "date": self.d(10, 2),
                                                           "qualite": "mauvaise", "date_peremption": (date.today() + timedelta(days=5)).isoformat()},
                        follow_redirects=True)
        html = txt(self.c.get("/cantine/stock"))
        self.assertIn("seuil d'alerte", html)
        self.assertIn("Qualité mauvaise", html)
        self.assertIn("Périme dans 5 jour", html)
        syn = self._synthese()
        self.assertEqual(syn["annee"], self.annee)
        self.assertGreaterEqual(syn["alertes"], 1)
        self.assertEqual(syn["valeur"], 95 * 450)
        huile_l = next(l for l in syn["lignes"] if l["vivre"] == "Huile")
        self.assertEqual(huile_l["stock"], 4)
        self.assertEqual(huile_l["qualite"], "mauvaise")
        self.assertEqual(len(huile_l["alertes"]), 3)

    def _synthese(self):
        """H.stock_cantine_synthese (interface exposée aux autres modules) via le namespace partagé notif.H."""
        import notif
        with app.app_context():
            return notif.H.stock_cantine_synthese(1, self.annee)

    def test_05_sortie_jour_et_recap(self):
        d = self.d(10, 8)
        r = self.c.get(f"/cantine/stock/sortie-jour?date={d}")
        self.assertEqual(r.status_code, 200)
        riz = q("SELECT * FROM cantine_vivres WHERE nom='Riz blanc'", un=True)
        huile = q("SELECT * FROM cantine_vivres WHERE nom='Huile'", un=True)
        r = self.c.post("/cantine/stock/sortie-jour", data={"date": d, "rationnaires": "210", f"q_{riz['id']}": "12,5", f"q_{huile['id']}": "1",
                                                          "observations": "riz au gras"}, follow_redirects=True)
        self.assertIn("Sortie du", txt(r))
        self.assertIn("2 vivre(s), 210 rationnaire(s)", txt(r))
        m = q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=? AND date=? AND type='sortie'", (riz["id"], d), un=True)
        self.assertEqual(m["quantite"], 12.5)
        self.assertEqual(m["rationnaires"], 210)
        self.assertEqual(m["annee"], self.annee)
        html = txt(self.c.get("/cantine/stock/recap"))
        self.assertIn("Riz blanc", html)
        self.assertIn("82,5", html)             # 95 − 12,5
        self.assertIn("Octobre", html)
        r = self.c.get("/cantine/stock/recap.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        if openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/cantine/stock/recap.xlsx").data))
            self.assertIn("Par mois", wb.sheetnames)
            vals = [c.value for row in wb["Synthèse"].iter_rows() for c in row]
            self.assertIn("Riz blanc", vals)

    def test_06_journal_pdf_excel(self):
        html = txt(self.c.get("/cantine/stock/mouvements"))
        self.assertIn("Riz blanc", html)
        self.assertIn("BL 1", html)
        html = txt(self.c.get("/cantine/stock/mouvements?type=perte"))
        self.assertIn("Perte", html)
        self.assertNotIn("BL 1", html)
        r = self.c.get("/cantine/stock/mouvements.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.c.get("/cantine/stock/etat.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        if openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/cantine/stock/etat.xlsx").data))
            self.assertIn("Vivres", wb.sheetnames)
            vals = [c.value for row in wb["État du stock"].iter_rows() for c in row]
            self.assertIn("Riz blanc", vals)
            self.assertIn(82.5, vals)          # 95 − 12,5 (sortie du jour)
            wb = openpyxl.load_workbook(io.BytesIO(self.c.get("/cantine/stock/mouvements.xlsx").data))
            vals = [c.value for row in wb.active.iter_rows() for c in row]
            self.assertIn("Entrée", vals)
            self.assertIn(100, vals)

    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_07_import_excel(self):
        r = self.c.get("/cantine/stock/modele.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["Vivres", "Mouvements"])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vivres"
        ws.append(S.ENT_VIVRES)
        ws.append(["Sardine import", "Protéine", "boîte", 10, "oui"])
        ws.append(["Huile", "Huile", "litre", 8, "oui"])          # mise à jour du seuil
        ws2 = wb.create_sheet("Mouvements")
        ws2.append(S.ENT_MOUVEMENTS)
        ws2.append([f"15/10/{self.y}", "Sardine import", "Entrée", 40, "Bonne", "APE", "Marché", "BR 7", 600, f"15/10/{self.y + 1}", "", "import"])
        ws2.append([f"16/10/{self.y}", "Sardine import", "Sortie", 4, "", "", "", "", "", "", 190, ""])
        ws2.append([f"16/10/{self.y}", "Nouveau vivre", "Entrée", 3, "", "", "", "", "", "", "", ""])
        ws2.append(["", "Sardine import", "Entrée", "x", "", "", "", "", "", "", "", ""])    # ligne invalide
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post("/cantine/stock/import", data={"fichier": (io.BytesIO(buf.getvalue()), "stock.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        html = txt(r)
        self.assertIn("Import terminé", html)
        self.assertIn("Vivres : 1 ajouté(s), 1 mis à jour", html)
        self.assertIn("Mouvements : 3 importé(s)", html)
        self.assertEqual(q("SELECT seuil_alerte FROM cantine_vivres WHERE nom='Huile'", un=True)["seuil_alerte"], 8)
        s = q("SELECT * FROM cantine_vivres WHERE nom='Sardine import'", un=True)
        self.assertEqual(s["categorie"], "proteine")
        mv = q("SELECT * FROM cantine_stock_mouvements WHERE vivre_id=? ORDER BY id", (s["id"],))
        self.assertEqual([m["type"] for m in mv], ["entree", "sortie"])
        self.assertEqual(mv[0]["prix_unitaire"], 600)
        self.assertEqual(mv[0]["date_peremption"], f"{self.y + 1}-10-15")
        self.assertEqual(mv[1]["rationnaires"], 190)
        self.assertTrue(q("SELECT id FROM cantine_vivres WHERE nom='Nouveau vivre'", un=True))
        self.assertEqual(S.calculer_stock(mv), 36)
        # fichier manquant
        r = self.c.post("/cantine/stock/import", data={}, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Choisissez un fichier", txt(r))

    def test_08_reinitialisation(self):
        n = q("SELECT COUNT(*) n FROM cantine_stock_mouvements WHERE ecole_id=1 AND annee=?", (self.annee,), un=True)["n"]
        self.assertGreater(n, 0)
        r = self.c.post("/cantine/stock/reinitialiser", data={"annee": self.annee, "confirmation": "non"}, follow_redirects=True)
        self.assertIn("annulée", txt(r))
        self.assertEqual(q("SELECT COUNT(*) n FROM cantine_stock_mouvements WHERE annee=?", (self.annee,), un=True)["n"], n)
        r = self.c.post("/cantine/stock/reinitialiser", data={"annee": self.annee, "confirmation": "REINITIALISER"}, follow_redirects=True)
        self.assertIn("réinitialisé", txt(r))
        self.assertEqual(q("SELECT COUNT(*) n FROM cantine_stock_mouvements WHERE annee=?", (self.annee,), un=True)["n"], 0)
        self.assertGreater(q("SELECT COUNT(*) n FROM cantine_vivres", un=True)["n"], 0)     # vivres conservés
        self.assertTrue(q("SELECT 1 FROM journal WHERE action='stock_reinitialisation'", un=True))

    def test_09_droits(self):
        """Un enseignant lit le stock mais ne peut pas écrire."""
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex
            ex("INSERT OR IGNORE INTO utilisateurs (identifiant, nom, role, mdp_hash, doit_changer_mdp, ecole_id) VALUES (?,?,?,?,0,1)",
               ("ens_stock", "Enseignant", "enseignant", generate_password_hash("x")))
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "ens_stock", "mdp": "x"}).status_code, 302)
        self.assertEqual(c2.get("/cantine/stock").status_code, 200)
        self.assertEqual(c2.get("/cantine/stock/vivres/0").status_code, 403)
        self.assertEqual(c2.post("/cantine/stock/reinitialiser", data={"confirmation": "REINITIALISER"}).status_code, 403)
        self.assertNotIn("Réinitialiser l'année", txt(c2.get("/cantine/stock")))


class TestEditionPublique(unittest.TestCase):
    """Rejoue toute la suite dans un processus séparé avec l'édition « publique »."""

    def test_publique(self):
        if EDITION != "privee":
            self.skipTest("déjà dans l'édition secondaire")
        env = dict(os.environ, GESCOL_EDITION_TEST="publique")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "tests.test_stock.TestStock", "tests.test_stock.TestCalculs"],
                           cwd=RACINE, env=env, capture_output=True, text=True, timeout=600)
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
