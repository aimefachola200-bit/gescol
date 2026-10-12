# -*- coding: utf-8 -*-
"""Tests du calendrier scolaire officiel (m_calendrier_officiel).
Lancer : python3 -m unittest tests.test_calendrier_officiel   (depuis /home/claude/gescol)

Édition « publique » (comptes de demo.py recréés à la main, connexion directeur / demo1234) : page, préremplissage 2026-2027, CRUD,
recharger, réinitialiser, année suivante, Excel, import, PDF, calculs, onglet du synoptique (?retour=synoptique),
passerelle avec le registre. L'édition « cs » (demo_juridiction.py, responsable / demo1234, ?retour=jur_synoptique)
est rejouée dans un sous-processus (TestEditionCS), comme tests/test_jur_synoptique.py."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "publique")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_calo_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import m_calendrier_officiel as CO  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

EN_PLACE = config.EDITION == EDITION
AN = "2026-2027"


def txt(r):
    return html_mod.unescape(r.data.decode())


def xlsx_feuilles(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(r) for r in wb[ws.title].iter_rows(values_only=True)] for ws in wb.worksheets}


def classeur(lignes, entetes=CO.ENTETES_XLSX):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calendrier"
    ws.append(entetes)
    for l in lignes:
        ws.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================ fonctions pures
class TestCalculs(unittest.TestCase):
    def setUp(self):
        self.lst = CO.lignes_officielles(AN)

    def test_officiel_2026(self):
        self.assertEqual(len(self.lst), 11)
        self.assertEqual([l["type"] for l in self.lst].count("trimestre"), 3)
        self.assertEqual(self.lst[0]["date_debut"], "2026-09-07")
        self.assertEqual(CO.lignes_officielles("2031-2032"), [])

    def test_semaines_par_trimestre(self):
        st = {s["libelle"]: s for s in CO.stats_trimestres(self.lst)}
        # 1er trimestre : 14 semaines pleines (14/09 → 18/12/2026, aucun jour férié saisi)
        self.assertEqual(st["Premier trimestre"]["jours"], 70)
        self.assertEqual(st["Premier trimestre"]["semaines"], "14 semaines")
        # 2e trimestre : 04/01 → 24/03/2027 = 58 jours ouvrés, moins détente (22–26/02 : 5) et fête des cultes (07–08/01 : 2)
        self.assertEqual(st["Deuxième trimestre"]["jours"], 51)
        self.assertEqual(st["Deuxième trimestre"]["semaines"], "10 semaines 1 jour")
        # 3e trimestre : 08/04 → 25/06/2027 = 57 jours ouvrés, moins Ascension et lundi de Pentecôte
        self.assertEqual(st["Troisième trimestre"]["jours"], 55)
        self.assertEqual(st["Troisième trimestre"]["semaines"], "11 semaines")

    def test_jours_classe(self):
        jc = CO.jours_classe(self.lst)
        self.assertEqual(len(jc), 70 + 51 + 55)
        self.assertIn("2026-12-18", jc)        # vendredi 18 décembre : congés « après les cours » → jour de classe
        self.assertNotIn("2026-12-21", jc)     # lundi 21 décembre : congés
        self.assertNotIn("2027-01-07", jc)     # fête des cultes et traditions
        self.assertNotIn("2027-05-06", jc)     # Ascension
        self.assertNotIn("2026-09-12", jc)     # samedi
        self.assertNotIn("2026-09-09", jc)     # pré-rentrée : pas un jour de classe
        self.assertIn("2027-06-25", jc)        # dernier jour du 3e trimestre
        self.assertNotIn("2027-06-28", jc)

    def test_debut_effectif_et_texte(self):
        self.assertEqual(CO.debut_effectif({"type": "conges", "date_debut": "2026-12-18", "precision": "après les cours"}), date(2026, 12, 19))
        self.assertEqual(CO.debut_effectif({"type": "ferie", "date_debut": "2027-05-06", "precision": "après les cours"}), date(2027, 5, 6))
        self.assertEqual(CO.semaines_texte(70), "14 semaines")
        self.assertEqual(CO.semaines_texte(53), "10 semaines 3 jours")
        self.assertEqual(CO.semaines_texte(1), "1 jour")
        self.assertEqual(CO.semaines_texte(0), "0 jour")
        self.assertEqual(CO.annee_suivante(AN), "2027-2028")
        self.assertEqual(CO.annee_valide("2027-2028"), "2027-2028")
        self.assertIsNone(CO.annee_valide("2027-2029"))
        self.assertEqual(CO.decaler_un_an("2028-02-29"), "2029-02-28")
        self.assertEqual(CO.type_depuis_texte("Jour férié"), "ferie")
        self.assertEqual(CO.type_depuis_texte("Congés"), "conges")
        self.assertEqual(CO.type_depuis_texte("Grandes vacances"), "vacances")
        self.assertEqual(CO.type_depuis_texte("pré-rentrée"), "pre_rentree")
        self.assertEqual(CO.type_depuis_texte("n'importe quoi"), "autre")

    def test_grille_mois(self):
        mois = CO.grille_mois(AN, self.lst)
        self.assertEqual(len(mois), 12)
        self.assertEqual(mois[0]["titre"], "Septembre 2026")
        self.assertEqual(mois[11]["titre"], "Août 2027")
        self.assertEqual(mois[0]["classe"], 13)     # 14 → 30 septembre 2026 : 13 jours ouvrés
        self.assertEqual(mois[3]["classe"], 14)     # décembre 2026 : 1 → 18 = 14 jours ouvrés
        for m in mois:
            for s in m["semaines"]:
                self.assertEqual(len(s), 7)

    def test_lire_ligne(self):
        v, err = CO.lire_ligne({"type": "ferie", "libelle": "Fête du 1er août", "date_debut": "2027-08-01"})
        self.assertEqual(err, [])
        self.assertEqual((v["date_debut"], v["date_fin"]), ("2027-08-01", "2027-08-01"))
        _v, err = CO.lire_ligne({"type": "conges", "libelle": "", "date_debut": "2027-03-10", "date_fin": "2027-03-01"})
        self.assertEqual(len(err), 2)


# ============================================================ édition école (publique)
class TestCalendrierEcole(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if config.juridiction():
            raise unittest.SkipTest("édition de juridiction : voir TestCalendrierCS")
        app.config["TESTING"] = True
        # comptes de demo.py (directeur / enseignant, mot de passe demo1234) créés directement : demo.charger() remplirait
        # la base partagée avec les autres suites (élèves, paiements…) et gênerait leurs nettoyages
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex, q
            for ident, nom, role in (("directeur", "Direction", "directeur"), ("enseignant", "Enseignant", "enseignant")):
                if not q("SELECT 1 FROM utilisateurs WHERE identifiant=?", (ident,), un=True):
                    ex("INSERT INTO utilisateurs (identifiant, nom, role, mdp_hash, doit_changer_mdp) VALUES (?,?,?,?,0)",
                       (ident, nom, role, generate_password_hash("demo1234")))
            ex("DELETE FROM calendrier_officiel")
            ex("DELETE FROM calendrier_officiel_meta")
            ex("DELETE FROM jours_classe WHERE ecole_id=1")
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "directeur", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code

    def get(self, url, code=200):
        r = self.c.get(url)
        self.assertEqual(r.status_code, code, f"{url} → {r.status_code}\n{r.data[:400]}")
        return r

    def post(self, url, data=None, code=302, **k):
        r = self.c.post(url, data=data or {}, **k)
        self.assertEqual(r.status_code, code, f"POST {url} → {r.status_code}\n{txt(r)[-1500:] if r.status_code != 302 else ''}")
        return r

    def q(self, sql, args=(), un=True):
        with app.app_context():
            from app import q
            return q(sql, args, un=un)

    def lignes(self, annee=AN):
        return [dict(r) for r in self.q("SELECT * FROM calendrier_officiel WHERE ecole_id=1 AND annee=? ORDER BY date_debut, ordre, id", (annee,), un=False)]

    def test_01_page_et_preremplissage(self):
        h = txt(self.get("/calendrier-scolaire"))
        for mot in ("Calendrier de l'année scolaire 2026-2027", "Premier trimestre", "Congés de détente", "Fête des cultes et traditions",
                    "lundi 14 septembre 2026", "14 semaines", "Réinitialiser (officiel)", "Recharger", "Appliquer au registre",
                    "017/MEMP/MES/MESRS-FT", "Préparer l'année 2027-2028"):
            self.assertIn(mot, h)
        self.assertEqual(len(self.lignes()), 11)
        self.assertTrue(all(l["officiel"] == 1 for l in self.lignes()))
        m = self.q("SELECT * FROM calendrier_officiel_meta WHERE ecole_id=1 AND annee=?", (AN,))
        self.assertIn("28 juillet 2026", m["source"])
        # vue calendrier : 12 mois
        h = txt(self.get("/calendrier-scolaire?vue=calendrier"))
        for mot in ("Septembre 2026", "Août 2027", "jour de classe", "Vue liste"):
            self.assertIn(mot, h)
        # le menu propose le calendrier officiel
        self.assertIn("/calendrier-scolaire", txt(self.get("/")))
        # année inconnue : vide, sans préremplissage
        h = txt(self.get("/calendrier-scolaire?annee=2030-2031"))
        self.assertIn("Aucune période pour 2030-2031", h)
        self.assertEqual(self.lignes("2030-2031"), [])
        self.get("/calendrier-scolaire?annee=nimporte")  # année invalide → année active

    def test_02_onglet_synoptique(self):
        h = txt(self.get("/calendrier-scolaire?retour=synoptique"))
        self.assertIn("Tableau synoptique de l'école", h)
        self.assertIn('class="actif" href="/calendrier-scolaire?retour=synoptique"', h)
        self.assertIn("retour=synoptique", h)
        # l'onglet est proposé depuis le synoptique
        self.assertIn("Calendrier scolaire", txt(self.get("/synoptique")))
        self.get("/calendrier-scolaire/ligne/0?retour=synoptique")
        self.get("/calendrier-scolaire/import?retour=synoptique")

    def test_03_crud(self):
        # ajout
        r = self.post("/calendrier-scolaire/ligne/0?annee=2026-2027", {"type": "ferie", "libelle": "Fête du 1er août",
                                                                        "date_debut": "2027-08-01", "date_fin": ""})
        self.assertIn("/calendrier-scolaire?annee=2026-2027", r.headers["Location"])
        l = [x for x in self.lignes() if x["libelle"] == "Fête du 1er août"][0]
        self.assertEqual((l["date_fin"], l["officiel"]), ("2027-08-01", 0))
        # erreur : libellé vide, fin avant début
        self.post("/calendrier-scolaire/ligne/0", {"type": "ferie", "libelle": "", "date_debut": "2027-03-10", "date_fin": "2027-03-01"}, code=400)
        # modification (journée pédagogique un mercredi du 1er trimestre → 1 jour de classe en moins)
        self.get(f"/calendrier-scolaire/ligne/{l['id']}")
        self.post(f"/calendrier-scolaire/ligne/{l['id']}", {"type": "ferie", "libelle": "Journée pédagogique", "date_debut": "2026-10-07",
                                                             "date_fin": "2026-10-07", "precision": "mercredi"})
        h = txt(self.get("/calendrier-scolaire"))
        self.assertIn("Journée pédagogique", h)
        self.assertIn("13 semaines 4 jours", h)
        self.assertEqual(self.get("/calendrier-scolaire/ligne/999999", code=404).status_code, 404)
        # suppression
        self.post(f"/calendrier-scolaire/ligne/{l['id']}/supprimer")
        self.assertEqual(len(self.lignes()), 11)
        self.assertIn("ligne supprimée", self.q("SELECT details FROM journal WHERE action='calendrier_officiel' ORDER BY id DESC")["details"])

    def test_04_recharger_reinitialiser(self):
        # modifier une ligne officielle + ajouter une ligne, puis recharger : officiel remis, ajout conservé
        t1 = [x for x in self.lignes() if x["libelle"] == "Premier trimestre"][0]
        self.post(f"/calendrier-scolaire/ligne/{t1['id']}", {"type": "trimestre", "libelle": "Premier trimestre", "date_debut": "2026-09-21",
                                                              "date_fin": "2026-12-18"})
        self.post("/calendrier-scolaire/ligne/0", {"type": "ferie", "libelle": "Tabaski", "date_debut": "2027-05-27"})
        self.post("/calendrier-scolaire/recharger?annee=2026-2027")
        lst = self.lignes()
        self.assertEqual(len(lst), 12)
        self.assertEqual([x["date_debut"] for x in lst if x["libelle"] == "Premier trimestre"], ["2026-09-14"])
        self.assertEqual(len([x for x in lst if x["libelle"] == "Tabaski"]), 1)
        # réinitialiser : tout supprimé puis officiel
        self.post("/calendrier-scolaire/reinitialiser?annee=2026-2027")
        lst = self.lignes()
        self.assertEqual(len(lst), 11)
        self.assertEqual([x for x in lst if x["libelle"] == "Tabaski"], [])
        # année sans calendrier officiel : recharger refuse, réinitialiser vide
        self.post("/calendrier-scolaire/ligne/0?annee=2030-2031", {"type": "autre", "libelle": "Essai", "date_debut": "2030-10-01"})
        self.post("/calendrier-scolaire/recharger?annee=2030-2031")
        self.assertEqual(len(self.lignes("2030-2031")), 1)
        self.post("/calendrier-scolaire/reinitialiser?annee=2030-2031")
        self.assertEqual(self.lignes("2030-2031"), [])

    def test_05_annee_suivante(self):
        self.post("/calendrier-scolaire/reinitialiser?annee=2027-2028")
        self.post("/calendrier-scolaire/preparer?annee=2026-2027")
        lst = self.lignes("2027-2028")
        self.assertEqual(len(lst), 11)
        t1 = [x for x in lst if x["libelle"] == "Premier trimestre"][0]
        self.assertEqual((t1["date_debut"], t1["date_fin"], t1["officiel"]), ("2027-09-14", "2027-12-18", 0))
        vac = [x for x in lst if x["type"] == "vacances"][0]
        self.assertIn("2028", vac["precision"])
        h = txt(self.get("/calendrier-scolaire?annee=2027-2028"))
        self.assertIn("À compléter", h)
        # une seconde préparation est refusée tant que l'année contient des lignes
        self.post("/calendrier-scolaire/preparer?annee=2026-2027")
        self.assertEqual(len(self.lignes("2027-2028")), 11)
        # sélecteur d'année : active, suivante, années présentes
        self.assertIn("2027-2028", h)
        self.post("/calendrier-scolaire/meta?annee=2027-2028", {"reference_arrete": "Arrêté n° 999 (essai)", "source": "MEMP"})
        self.assertIn("Arrêté n° 999 (essai)", txt(self.get("/calendrier-scolaire?annee=2027-2028")))

    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_06_excel_et_import(self):
        r = self.get("/calendrier-scolaire/export.xlsx?annee=2026-2027")
        f = xlsx_feuilles(r.data)
        self.assertIn("Calendrier", f)
        self.assertIn("Jours de classe", f)
        cal = [l for l in f["Calendrier"] if l and l[1] == "Premier trimestre"]
        self.assertEqual(cal[0][2], "2026-09-14")
        jc = [l for l in f["Jours de classe"] if l and l[0] == "Premier trimestre"]
        self.assertEqual(jc[0][3], 70)
        r = self.get("/calendrier-scolaire/modele.xlsx")
        self.assertIn("Type", [str(x) for l in xlsx_feuilles(r.data)["Calendrier"] for x in l])
        # import : remplace les lignes de l'année
        self.get("/calendrier-scolaire/import")
        buf = classeur([["Trimestre", "Premier trimestre", "2026-09-14", "2026-12-18", "", "14 semaines", "", "Oui"],
                        ["Congés", "Congés de Noël", "18/12/2026", "03/01/2027", "après les cours", "2 semaines", "", "Non"],
                        ["Jour férié", "Fête nationale", "01/08/2027", "", "", "1 jour", "", ""]])
        r = self.c.post("/calendrier-scolaire/import?annee=2026-2027", data={"fichier": (buf, "cal.xlsx"), "reference_arrete": "Arrêté importé"},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302, txt(r)[-800:])
        lst = self.lignes()
        self.assertEqual(len(lst), 3)
        self.assertEqual([(x["type"], x["officiel"]) for x in lst], [("trimestre", 1), ("conges", 0), ("ferie", 0)])
        self.assertEqual([x["date_fin"] for x in lst if x["type"] == "ferie"], ["2027-08-01"])
        self.assertIn("Arrêté importé", txt(self.get("/calendrier-scolaire")))
        # fichier erroné : rien n'est modifié
        buf = classeur([["Trimestre", "", "2026-09-14", "2026-12-18"]])
        self.c.post("/calendrier-scolaire/import?annee=2026-2027", data={"fichier": (buf, "cal.xlsx")}, content_type="multipart/form-data")
        self.assertEqual(len(self.lignes()), 3)
        # sans fichier
        self.post("/calendrier-scolaire/import?annee=2026-2027")
        self.assertEqual(len(self.lignes()), 3)
        # retour à l'officiel
        self.post("/calendrier-scolaire/reinitialiser?annee=2026-2027")
        self.assertEqual(len(self.lignes()), 11)

    def test_07_pdf(self):
        r = self.get("/calendrier-scolaire/calendrier.pdf?annee=2026-2027")
        self.assertTrue(r.data.startswith(b"%PDF"))
        self.assertGreater(len(r.data), 2000)
        r = self.get("/calendrier-scolaire/calendrier.pdf?annee=2030-2031")
        self.assertTrue(r.data.startswith(b"%PDF"))

    def test_08_registre(self):
        self.post("/calendrier-scolaire/reinitialiser?annee=2026-2027")
        self.post("/calendrier-scolaire/appliquer-registre?annee=2026-2027")
        with app.app_context():
            from app import q
            import m_registre
            p = m_registre.parametres_calendrier(m_calendrier_H(), 1, AN)
            self.assertEqual((p["rentree"], p["fin"]), ("2026-09-14", "2027-06-25"))
            jc = {r["date"]: dict(r) for r in q("SELECT * FROM jours_classe WHERE ecole_id=1", (), un=False)}
        self.assertEqual(jc["2026-12-21"]["demi_journees"], 0)
        self.assertEqual(jc["2027-05-06"]["motif"], "Ascension")
        self.assertNotIn("2026-12-18", jc)      # vendredi 18 décembre : encore un jour de classe
        self.assertNotIn("2026-12-19", jc)      # samedi : non marqué
        self.assertIn("2027-02-22", jc)         # congés de détente
        self.assertIn("2027-04-07", jc)
        self.assertNotIn("2027-06-28", jc)      # grandes vacances : hors de la période des cours
        # le calendrier du registre reflète ces jours
        h = txt(self.get("/calendrier?mois=2026-12"))
        self.assertIn("Congés de fin du premier trimestre", h)

    def test_09_api_et_droits(self):
        with app.app_context():
            H = m_calendrier_H()
            lst = H.calendrier_officiel(AN, 1)
            self.assertEqual(len(lst), 11)
            self.assertEqual(len(H.calendrier_jours_classe(AN, 1)), 176)
            self.assertEqual(len(H.calendrier_officiel(AN, 999)), 11)       # école sans lignes : calendrier officiel
            self.assertEqual(H.calendrier_officiel("2031-2032", 1), [])
        # enseignant : lecture seule
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "enseignant", "mdp": "demo1234"}).status_code, 302)
        h = txt(c2.get("/calendrier-scolaire"))
        self.assertIn("Premier trimestre", h)
        self.assertNotIn("Réinitialiser (officiel)", h)
        self.assertEqual(c2.get("/calendrier-scolaire/ligne/0").status_code, 403)
        self.assertEqual(c2.post("/calendrier-scolaire/reinitialiser").status_code, 403)
        self.assertEqual(c2.get("/calendrier-scolaire/calendrier.pdf").status_code, 200)
        # non connecté
        self.assertEqual(app.test_client().get("/calendrier-scolaire").status_code, 302)


def m_calendrier_H():
    """Le H des modules (référence gardée par m_synoptique pour les tests)."""
    import m_synoptique
    return m_synoptique._H


# ============================================================ édition CS (sous-processus)
class TestCalendrierCS(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not config.juridiction() or not EN_PLACE:
            raise unittest.SkipTest("édition de juridiction rejouée en sous-processus")
        app.config["TESTING"] = True
        import demo_juridiction
        demo_juridiction.main()
        cls.c = app.test_client()
        r = cls.c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"})
        assert r.status_code == 302, r.status_code

    def q(self, sql, args=(), un=True):
        with app.app_context():
            from app import q
            return q(sql, args, un=un)

    def test_structure(self):
        r = self.c.get("/juridiction/calendrier-scolaire")
        self.assertEqual(r.status_code, 200, r.data[:400])
        h = txt(r)
        self.assertIn("Premier trimestre", h)
        self.assertIn("calendrier commun de la structure", h)
        self.assertEqual(self.q("SELECT COUNT(*) n FROM calendrier_officiel WHERE ecole_id=0 AND annee=?", (AN,))["n"], 11)
        # onglet du synoptique de la structure
        h = txt(self.c.get("/juridiction/calendrier-scolaire?retour=jur_synoptique"))
        self.assertIn("Tableau synoptique de la DDEMP" if config.EDITION == "ddemp" else "Tableau synoptique de la circonscription scolaire", h)
        self.assertIn("retour=jur_synoptique", h)
        self.assertIn("Calendrier scolaire", txt(self.c.get("/juridiction/synoptique")))
        # le menu de la structure propose le calendrier
        self.assertIn("/juridiction/calendrier-scolaire", txt(self.c.get("/", follow_redirects=True)))
        # CRUD sur le calendrier de la structure
        r = self.c.post("/juridiction/calendrier-scolaire/ligne/0?annee=2026-2027&retour=jur_synoptique",
                        data={"type": "evenement", "libelle": "Journée de l'UP", "date_debut": "2026-11-05"})
        self.assertEqual(r.status_code, 302)
        self.assertIn("retour=jur_synoptique", r.headers["Location"])
        self.assertEqual(self.q("SELECT COUNT(*) n FROM calendrier_officiel WHERE ecole_id=0 AND annee=?", (AN,))["n"], 12)
        r = self.c.get("/juridiction/calendrier-scolaire/calendrier.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.c.get("/juridiction/calendrier-scolaire/export.xlsx")
        self.assertEqual(r.status_code, 200)
        # pas d'école ouverte : le registre n'est pas applicable
        r = self.c.post("/juridiction/calendrier-scolaire/appliquer-registre", follow_redirects=True)
        self.assertIn("Aucune école ouverte", txt(r))
        self.assertEqual(self.c.post("/juridiction/calendrier-scolaire/reinitialiser").status_code, 302)
        self.assertEqual(self.q("SELECT COUNT(*) n FROM calendrier_officiel WHERE ecole_id=0 AND annee=?", (AN,))["n"], 11)

    def test_ecole_ouverte(self):
        jid = self.q("SELECT id FROM jur_ecoles WHERE actif=1 ORDER BY id")["id"]
        self.assertEqual(self.c.post(f"/juridiction/ecoles-modules/{jid}/ouvrir").status_code, 302)
        eid = self.q("SELECT id FROM ecoles WHERE jur_ecole_id=?", (jid,))["id"]
        try:
            # école ouverte : calendrier propre à l'école (préremplissage), bascule vers celui de la structure
            h = txt(self.c.get("/juridiction/calendrier-scolaire"))
            self.assertNotIn("calendrier commun de la structure", h)
            self.assertIn("Calendrier de la structure", h)
            self.assertIn("Appliquer au registre", h)
            self.assertEqual(self.q("SELECT COUNT(*) n FROM calendrier_officiel WHERE ecole_id=? AND annee=?", (eid, AN))["n"], 11)
            h = txt(self.c.get("/juridiction/calendrier-scolaire?portee=structure"))
            self.assertIn("calendrier commun de la structure", h)
            # l'onglet du synoptique de la structure montre toujours le calendrier commun
            self.assertIn("calendrier commun de la structure", txt(self.c.get("/juridiction/calendrier-scolaire?retour=jur_synoptique")))
            # passerelle avec le registre de l'école ouverte
            r = self.c.post("/juridiction/calendrier-scolaire/appliquer-registre", follow_redirects=True)
            self.assertIn("Calendrier appliqué au registre", txt(r))
            self.assertEqual(self.q("SELECT demi_journees FROM jours_classe WHERE ecole_id=? AND date='2026-12-21'", (eid,))["demi_journees"], 0)
        finally:
            self.c.post("/juridiction/ecole/fermer")


class TestEditionCS(unittest.TestCase):
    """Rejoue TestCalendrierCS dans l'édition « cs » (processus séparé : config.EDITION est fixé à l'import)."""

    def test_cs(self):
        if config.juridiction():
            raise unittest.SkipTest("déjà en édition de juridiction")
        env = dict(os.environ, GESCOL_EDITION_TEST="cs")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "tests.test_calendrier_officiel.TestCalendrierCS"], cwd=RACINE,
                           env=env, capture_output=True, text=True, timeout=900)
        self.assertEqual(r.returncode, 0, r.stderr[-3000:])
        self.assertNotIn("skipped", r.stderr.split("\n")[-2], r.stderr[-500:])


if __name__ == "__main__":
    unittest.main()
