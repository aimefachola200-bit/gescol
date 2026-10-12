# -*- coding: utf-8 -*-
"""Tests du tableau synoptique de l'école (m_synoptique) et des règles de retraite / ancienneté (m_carriere).
Lancer : python3 -m unittest tests.test_synoptique   (depuis /home/claude/gescol)

L'édition testée est choisie par GESCOL_EDITION_TEST (publique par défaut) ; la suite est exécutée pour
l'édition « publique » et, dans un second processus, pour l'édition « privee » (voir TestEditionPrivee)."""
import html as html_mod
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("GESCOL_EDITION_TEST", "publique")
DOSSIER = tempfile.mkdtemp(prefix=f"gescol_syno_{EDITION}_")
os.environ["GESCOL_EDITION"] = EDITION
os.environ["GESCOL_DATA"] = DOSSIER

from app import app  # noqa: E402
import config  # noqa: E402
import m_carriere as MC  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def txt(r):
    """Texte d'une réponse HTML, entités décodées (l'apostrophe est rendue &#39; par Jinja)."""
    return html_mod.unescape(r.data.decode())


class TestSynoptique(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        cls.c = app.test_client()
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex, q
            ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
            ex("UPDATE ecoles SET nom=?, departement=?, cs=?, ccs_nom=?, up=?, zone=? WHERE id=1",
               ("EPP TEST SYNOPTIQUE", "Donga", "CS Djougou II", "AGOSSOU Marc", "UP 3", "59"))
            if not q("SELECT 1 FROM classes WHERE ecole_id=1", un=True):
                ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)", ("CM2 A", "CM2", 0, 1))
            cls.classe_id = q("SELECT id FROM classes WHERE ecole_id=1 ORDER BY id", un=True)["id"]
        r = cls.c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}, follow_redirects=False)
        assert r.status_code == 302, r.status_code
        r = cls.c.get("/synoptique")
        assert r.status_code == 200, r.data[:300]

    # ------------------------------------------------------------ pages
    def test_01_pages(self):
        for url in ("/synoptique", "/synoptique/personnel", "/synoptique/comites", "/synoptique/comites/ape",
                    "/synoptique/comites/direction", "/synoptique/comites/classe", "/synoptique/edt", "/synoptique/edt/cm2",
                    "/synoptique/edt/maternelle1", "/synoptique/planification", "/synoptique/planification/cm1",
                    "/synoptique/reglement", "/synoptique/import", "/synoptique/personnel/0", "/synoptique/edt/ci/modifier",
                    "/synoptique/planification/ce1/modifier", "/synoptique/reglement/0"):
            r = self.c.get(url)
            self.assertEqual(r.status_code, 200, f"{url} -> {r.status_code}")
        html = txt(self.c.get("/synoptique"))
        for mot in ("Nom du MEMP", "Nom du DDEMP", "Nom du CCS", "CP de zone", "Zone pédagogique", "Nom du RUP", "Co-RUP",
                    "Année de création", "EPP TEST SYNOPTIQUE", "Ministère des Enseignements Maternel et Primaire", "AGOSSOU Marc"):
            self.assertIn(mot, html)
        self.assertEqual(self.c.get("/synoptique/edt/inconnu").status_code, 404)
        self.assertEqual(self.c.get("/synoptique/comites/inconnu").status_code, 404)
        # le menu propose le tableau synoptique
        self.assertIn("/synoptique", txt(self.c.get("/")))

    def test_02_identite(self):
        r = self.c.post("/synoptique/identite", data={
            "memp": "Ministère des Enseignements Maternel et Primaire", "ddemp": "DDEMP Donga", "ccs": "KOUAGOU Pierre",
            "epp": "EPP TEST SYNOPTIQUE", "cp_zone": "YOVO Jeanne", "zone_numero": "59", "up": "UP 3 Djougou", "rup_nom": "ALI Moussa",
            "corup_nom": "SEKO Awa", "annee_creation": "1975", "dirige_par": "Directrice", "directeur_grade": "A1-5",
            "directeur_nom": "TOSSOU Clarisse"}, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        html = txt(r)
        self.assertIn("Identité de l'école enregistrée", html)
        for mot in ("KOUAGOU Pierre", "YOVO Jeanne", "ALI Moussa", "SEKO Awa", "1975", "Directrice", "TOSSOU Clarisse"):
            self.assertIn(mot, html)
        with app.app_context():
            from app import q
            e = q("SELECT * FROM ecoles WHERE id=1", un=True)
            self.assertEqual(e["ccs_nom"], "KOUAGOU Pierre")
            self.assertEqual(e["directeur_nom"], "TOSSOU Clarisse")
            self.assertEqual(e["directeur_titre"], "Directrice")
            self.assertEqual(e["up"], "UP 3 Djougou")
        r = self.c.post("/synoptique/identite", data={"annee_creation": "abcd", "dirige_par": "Directeur"}, follow_redirects=True)
        self.assertIn("Année de création invalide", txt(r))

    # ------------------------------------------------------------ personnel : ajout, calculs, modification, retrait
    def test_03_personnel(self):
        ancien = config.ANNEE_ACTIVE
        data = {"nom": "akpo", "prenoms": "Jean", "sexe": "M", "fonction": "Enseignant(e)", "telephone": "97 00 00 00",
                "matricule": "SYN-001", "nip": "1234567", "date_naissance": "28/02/1969", "promotion": "2000", "categorie": "ACDPE",
                "grade": "B1-12", "indice_reel": "520", "indice_paye": "500", "cours_tenus": "CM2 A",
                "date_prise_service": "2000-10-01", "date_service_dep": "2010-10-01", "date_service_ecole": "2019-09-16",
                "date_derniere_formation": "2024-03-15", "diplome_academique": "BAC", "diplome_academique_annee": "1998",
                "ceap": "oui", "ceap_annee": "2004", "cap": "oui", "cap_annee": "2009", "visite_aptitude": "oui",
                "visite_aptitude_note": "15", "visites_cp": "2", "visites_ccs": "1", "visites_directeur": "3",
                "participation_up": "80", "gestion_collegiale": "Responsable de la cantine", "note_evaluation": "16"}
        r = self.c.post("/synoptique/personnel/0", data=data, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        html = txt(r)
        self.assertIn("Agent enregistré", html)
        self.assertIn("AKPO Jean", html)
        with app.app_context():
            from app import q
            p = q("SELECT p.*, c.* FROM personnel p JOIN personnel_carriere c ON c.personnel_id=p.id WHERE p.matricule='SYN-001'", un=True)
            self.assertEqual(p["ecole_id"], 1)
            self.assertEqual(p["promotion"], "2000")
            self.assertEqual(p["cours_tenus"], "CM2 A")
            self.assertEqual(p["visites_directeur"], 3)
            self.assertEqual(p["date_derniere_formation"], "2024-03-15")
            self.assertEqual(p["date_naissance"], "1969-02-28")
            pid = p["id"]
            # calculs : né le 28/02/1969, grade B (58 ans) -> 01/04/2027 ; anciennetés au 31/12 de la 1re année scolaire
            f = MC.calculer_fiche(dict(p), auj=date(2026, 10, 7))
            self.assertEqual(f["retraite_txt"], "01/04/2027")
            self.assertEqual(f["reference"], date(int(config.ANNEE_ACTIVE[:4]), 12, 31))
            ref_annee = int(config.ANNEE_ACTIVE[:4])
            self.assertEqual(f["anc"]["generale"], (ref_annee - 2000, 3))
            self.assertEqual(f["anc"]["departement"], (ref_annee - 2010, 3))
            self.assertEqual(f["anc"]["ecole"], (ref_annee - 2019, 3))
        html = txt(self.c.get("/synoptique/personnel"))
        self.assertIn("01/04/2027", html)
        self.assertIn("CEAP oui (2004)", html)
        self.assertIn("2 / 1", html)
        self.assertIn("SYN-001", html)
        # modification
        data2 = dict(data, prenoms="Jean-Baptiste", cours_tenus="CM1 B", visites_directeur="4")
        r = self.c.post(f"/synoptique/personnel/{pid}", data=data2, follow_redirects=True)
        self.assertIn("AKPO Jean-Baptiste", txt(r))
        self.assertIn("CM1 B", txt(r))
        # erreurs
        r = self.c.post("/synoptique/personnel/0", data=dict(data, nom="", matricule="SYN-002"))
        self.assertEqual(r.status_code, 400)
        r = self.c.post("/synoptique/personnel/0", data=dict(data, matricule="SYN-001", nom="DOUBLON"))
        self.assertEqual(r.status_code, 400)
        self.assertIn("déjà attribué", txt(r))
        # retrait = actif 0, puis réactivation
        r = self.c.post(f"/synoptique/personnel/{pid}/retirer", follow_redirects=True)
        self.assertIn("retiré", txt(r))
        with app.app_context():
            from app import q
            self.assertEqual(q("SELECT actif FROM personnel WHERE id=?", (pid,), un=True)["actif"], 0)
        self.assertNotIn("SYN-001", txt(self.c.get("/synoptique/personnel")))
        self.assertIn("SYN-001", txt(self.c.get("/synoptique/personnel?tous=1")))
        r = self.c.post(f"/synoptique/personnel/{pid}/reactiver", follow_redirects=True)
        self.assertIn("SYN-001", txt(self.c.get("/synoptique/personnel")))
        config.ANNEE_ACTIVE = ancien

    # ------------------------------------------------------------ comités
    def test_04_comites(self):
        for comite in ("sante", "ape", "direction", "classe"):
            qs = f"?classe_id={self.classe_id}" if comite == "classe" else ""
            r = self.c.post(f"/synoptique/comites/{comite}/membre/0{qs}",
                            data={"nom_prenoms": f"MEMBRE {comite.upper()}", "qualite": "Président(e)", "origine": "Parent d'élève",
                                  "contact": "96 00 00 00", "ordre": "1"}, follow_redirects=True)
            self.assertEqual(r.status_code, 200)
            self.assertIn(f"MEMBRE {comite.upper()}", txt(r))
            with app.app_context():
                from app import q
                m = q("SELECT * FROM synoptique_membres WHERE comite=? AND nom_prenoms=?", (comite, f"MEMBRE {comite.upper()}"), un=True)
                self.assertEqual(m["classe_id"], self.classe_id if comite == "classe" else None)
                mid = m["id"]
            r = self.c.post(f"/synoptique/comites/{comite}/membre/{mid}",
                            data={"nom_prenoms": f"MEMBRE {comite.upper()} MODIFIÉ", "qualite": "Secrétaire"}, follow_redirects=True)
            self.assertIn(f"MEMBRE {comite.upper()} MODIFIÉ", txt(r))
            self.assertIn("Secrétaire", txt(r))
            # PDF du comité
            r = self.c.get(f"/synoptique/comites/{comite}.pdf{qs}")
            self.assertEqual(r.status_code, 200)
            self.assertTrue(r.data.startswith(b"%PDF"))
            # suppression
            r = self.c.post(f"/synoptique/comites/{comite}/membre/{mid}/supprimer", follow_redirects=True)
            self.assertNotIn(f"MEMBRE {comite.upper()} MODIFIÉ", txt(r))
            self.assertEqual(self.c.post(f"/synoptique/comites/{comite}/membre/{mid}/supprimer").status_code, 404)
        r = self.c.post("/synoptique/comites/sante/membre/0", data={"nom_prenoms": ""})
        self.assertEqual(r.status_code, 400)

    # ------------------------------------------------------------ emplois du temps
    def test_05_edt(self):
        html = txt(self.c.get("/synoptique/edt/cp"))
        self.assertIn("Cérémonial des couleurs 15 min", html)
        self.assertIn("ES/Morale 30 min", html)
        self.assertIn("Dictée 20 min", html)
        self.assertIn("Emploi du temps officiel", html)
        html = txt(self.c.get("/synoptique/edt/ce2"))
        self.assertIn("Dictée 45 min", html)
        self.assertIn("EA Couture 45 min", html)
        html = txt(self.c.get("/synoptique/edt/cm1"))
        self.assertIn("Dictée 60 min", html)
        self.assertIn("Français 75 min", html)
        html = txt(self.c.get("/synoptique/edt/maternelle2"))
        self.assertIn("Pré-lecture", html)
        self.assertIn("7h45 – 8h30", html)
        # modification cellule par cellule
        data = {"nb_lignes": "2", "h_0": "08h à 10h15", "c_0_lundi": "Cérémonial 15 min\nFrançais 60 min", "c_0_mardi": "Maths 60 min",
                "h_1": "10h30 à 12h", "c_1_vendredi": "EPS 45 min", "remarques": "Remarque de test\n\nDeuxième", "source": "École"}
        r = self.c.post("/synoptique/edt/cm2/modifier", data=data, follow_redirects=True)
        html = txt(r)
        self.assertIn("Emploi du temps CM2 enregistré", html)
        self.assertIn("Maths 60 min", html)
        self.assertIn("Remarque de test", html)
        self.assertNotIn("Emploi du temps officiel (non modifié)", html)
        with app.app_context():
            from app import q
            import json
            d = json.loads(q("SELECT contenu FROM synoptique_edt WHERE ecole_id=1 AND niveau='cm2'", un=True)["contenu"])
            self.assertEqual(len(d["lignes"]), 2)
            self.assertEqual(d["lignes"][1]["vendredi"], "EPS 45 min")
            self.assertEqual(d["remarques"], ["Remarque de test", "Deuxième"])
        # le CM1 n'est pas touché
        self.assertIn("Dictée 60 min", txt(self.c.get("/synoptique/edt/cm1")))
        # PDF et Excel
        r = self.c.get("/synoptique/edt/cm2.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        if openpyxl:
            r = self.c.get("/synoptique/edt/cm2.xlsx")
            wb = openpyxl.load_workbook(io.BytesIO(r.data))
            self.assertIn("EPS 45 min", [c.value for row in wb.active.iter_rows() for c in row])
        # réinitialisation
        r = self.c.post("/synoptique/edt/cm2/reinitialiser", follow_redirects=True)
        self.assertIn("Dictée 60 min", txt(r))
        self.assertNotIn("Maths 60 min", txt(r))

    # ------------------------------------------------------------ planification
    def test_06_planification(self):
        html = txt(self.c.get("/synoptique/planification/cm2"))
        self.assertIn("Septembre", html)
        self.assertIn("U1 SA5 à U3 SA7", html)                      # octobre, français CM2
        self.assertIn("Évaluation sommative d'étape 1", html)
        html = txt(self.c.get("/synoptique/planification/ci"))
        self.assertIn("Jours 1-10 : SA1", html)
        html = txt(self.c.get("/synoptique/planification/maternelle1"))
        self.assertIn("Juin", html)
        # ajout d'une ligne (formulaire : 10 lignes existantes + 1 vide)
        form = {"nb_lignes": "3", "p_0_trimestre": "I", "p_0_mois": "Septembre", "p_0_francais": "U1", "p_0_maths": "D1 SA1",
                "p_1_trimestre": "I", "p_1_mois": "Octobre", "p_1_es_morale": "SA1 Séq 2", "p_1_observations": "Visite CP",
                "p_2_mois": "", "p_2_francais": ""}
        r = self.c.post("/synoptique/planification/ce1/modifier", data=form, follow_redirects=True)
        html = txt(r)
        self.assertIn("Planification CE1 enregistrée", html)
        self.assertIn("Visite CP", html)
        self.assertIn("D1 SA1", html)
        with app.app_context():
            from app import q
            import json
            lignes = json.loads(q("SELECT contenu FROM synoptique_planification WHERE ecole_id=1 AND niveau='ce1'", un=True)["contenu"])
            self.assertEqual(len(lignes), 2)
        # « enregistrer et ajouter une ligne »
        r = self.c.post("/synoptique/planification/ce1/modifier", data=dict(form, action="ajouter"), follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        with app.app_context():
            from app import q
            import json
            lignes = json.loads(q("SELECT contenu FROM synoptique_planification WHERE ecole_id=1 AND niveau='ce1'", un=True)["contenu"])
            self.assertEqual(len(lignes), 3)
        # PDF, Excel, import Excel
        self.assertTrue(self.c.get("/synoptique/planification/ce1.pdf").data.startswith(b"%PDF"))
        if openpyxl:
            r = self.c.get("/synoptique/planification/ce1.xlsx")
            wb = openpyxl.load_workbook(io.BytesIO(r.data))
            self.assertIn("Visite CP", [c.value for row in wb.active.iter_rows() for c in row])
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Trimestre", "Mois", "Français", "Maths", "Morale", "Civisme", "Histoire", "Géographie", "EST", "EPS",
                       "Chant", "Poésie", "Conte", "Dessin", "Couture", "Observations"])
            ws.append(["I", "Septembre", "U1 importée", "D1", "", "", "", "", "SA1", "", "", "", "", "", "", "import"])
            ws.append(["II", "Janvier", "U5 importée"] + [""] * 13)
            buf = io.BytesIO()
            wb.save(buf)
            r = self.c.post("/synoptique/planification/ce2/import", data={"fichier": (io.BytesIO(buf.getvalue()), "plan.xlsx")},
                            content_type="multipart/form-data", follow_redirects=True)
            html = txt(r)
            self.assertIn("importée : 2 ligne(s)", html)
            self.assertIn("U5 importée", html)
        # réinitialisation
        r = self.c.post("/synoptique/planification/ce1/reinitialiser", follow_redirects=True)
        self.assertIn("Diphtongue", txt(r))

    # ------------------------------------------------------------ règlement intérieur
    def test_07_reglement(self):
        html = txt(self.c.get("/synoptique/reglement"))
        self.assertIn("PRÉAMBULE", html)
        self.assertIn("ENTRÉE – SORTIE – HORAIRE – DISCIPLINE", html)
        self.assertIn("Article 24", html)
        self.assertIn("Les châtiments corporels sont formellement interdits.", html)
        self.assertIn("Article 63", html)
        with app.app_context():
            from app import q
            self.assertEqual(q("SELECT COUNT(*) n FROM synoptique_reglement WHERE ecole_id=1 AND type='article'", un=True)["n"], 63)
            self.assertEqual(q("SELECT COUNT(*) n FROM synoptique_reglement WHERE ecole_id=1 AND type='titre'", un=True)["n"], 18)
        # ajout
        r = self.c.post("/synoptique/reglement/0", data={"type": "article", "numero": "64", "texte": "Article propre à l'école de test."},
                        follow_redirects=True)
        self.assertIn("Article propre à l'école de test.", txt(r))
        with app.app_context():
            from app import q
            a = q("SELECT * FROM synoptique_reglement WHERE numero='64' AND ecole_id=1", un=True)
            rid = a["id"]
            self.assertGreater(a["ordre"], 600)
        # modification
        r = self.c.post(f"/synoptique/reglement/{rid}", data={"type": "article", "numero": "64", "texte": "Article modifié.", "ordre": "5"},
                        follow_redirects=True)
        html = txt(r)
        self.assertIn("Article modifié.", html)
        self.assertLess(html.index("Article modifié."), html.index("Article 1 :"))     # ordre 5 : avant tout le reste
        # erreur : texte vide
        self.assertEqual(self.c.post(f"/synoptique/reglement/{rid}", data={"type": "article", "numero": "64", "texte": ""}).status_code, 400)
        # suppression
        r = self.c.post(f"/synoptique/reglement/{rid}/supprimer", follow_redirects=True)
        self.assertNotIn("Article modifié.", txt(r))
        # PDF
        r = self.c.get("/synoptique/reglement.pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        # rétablir l'officiel après une suppression d'article
        with app.app_context():
            from app import q
            a1 = q("SELECT id FROM synoptique_reglement WHERE numero='1' AND type='article' AND ecole_id=1", un=True)["id"]
        self.c.post(f"/synoptique/reglement/{a1}/supprimer")
        r = self.c.post("/synoptique/reglement/retablir", follow_redirects=True)
        self.assertIn("L'école ouvre ses portes", txt(r))
        with app.app_context():
            from app import q
            self.assertEqual(q("SELECT COUNT(*) n FROM synoptique_reglement WHERE ecole_id=1 AND type='article'", un=True)["n"], 63)

    # ------------------------------------------------------------ import / export Excel, PDF complet
    @unittest.skipIf(openpyxl is None, "openpyxl absent")
    def test_08_import_export(self):
        r = self.c.get("/synoptique/modele.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        self.assertIn("Personnel", wb.sheetnames)
        self.assertIn("Comité de santé", wb.sheetnames)
        # classeur d'import : personnel + comités
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Personnel"
        ws.append(["Tableau du personnel"])
        ws.append(["N° matricule", "NIP", "Nom", "Prénoms", "Sexe", "Promotion", "Catégorie", "Fonction", "Grade", "Indice réel",
                   "Indice payé", "Cours tenus", "Date de naissance", "Date de première prise de service",
                   "Date de prise de service dans le département", "Date de prise de service dans l'école", "Date de la dernière formation",
                   "Diplôme académique le plus élevé", "Année du diplôme académique", "CEAP", "Année du CEAP", "CAP", "Année du CAP",
                   "Visite d'aptitude", "Note", "Nombre de visites CP / CCS", "Nombre de visites du directeur", "Participation à l'UP (%)",
                   "Gestion collégiale", "Contact"])
        ws.append(["IMP-001", "7654321", "Hounsou", "Marie", "F", "2008", "FE", "Directeur / directrice", "A1-5", 600, 580, "CM2 A",
                   "15/02/1967", "01/10/2008", "01/10/2012", "16/09/2020", "10/01/2025", "Licence", 2006, "oui", 2010, "non", "",
                   "oui", 17, "3 / 2", 1, 100, "Gestion de la cantine", "95 00 00 00"])
        ws.append(["IMP-002", "", "Dossa", "Paul", "M", "", "ACDPCTD", "Enseignant(e)", "C1-3", "", "", "CI", "10/11/1970",
                   "2015-10-01", "", "2021-09-15", "", "BEPC", "", "", "", "", "", "non", "", "", "", "", "", ""])
        ws.append(["IMP-003", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws2 = wb.create_sheet("Comité de santé")
        ws2.append(["N°", "Nom et prénoms", "Qualité / fonction", "Structure d'origine", "Contact", "Observations"])
        ws2.append([1, "SANTÉ Importé", "Président(e)", "Agent de santé", "", ""])
        ws3 = wb.create_sheet("Comités de classe")
        ws3.append(["N°", "Nom et prénoms", "Qualité / fonction", "Structure d'origine", "Contact", "Observations", "Classe"])
        ws3.append([1, "CLASSE Importé", "Président(e)", "Parent d'élève", "", "", "CM2 A"])
        ws3.append([2, "CLASSE Inconnue", "Membre", "", "", "", "Classe inexistante"])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post("/synoptique/import", data={"fichier": (io.BytesIO(buf.getvalue()), "import.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        html = txt(r)
        self.assertIn("Personnel : 2 agent(s) ajouté(s), 0 mis à jour", html)
        self.assertIn("Comité de santé : 1 membre(s)", html)
        self.assertIn("Comité de classe : 1 membre(s)", html)
        with app.app_context():
            from app import q
            p = q("SELECT p.*, c.* FROM personnel p JOIN personnel_carriere c ON c.personnel_id=p.id WHERE p.matricule='IMP-001'", un=True)
            self.assertEqual(p["nom"], "HOUNSOU")
            self.assertEqual(p["date_naissance"], "1967-02-15")
            self.assertEqual(p["visites_cp"], 3)
            self.assertEqual(p["visites_ccs"], 2)
            self.assertEqual(p["visites_directeur"], 1)
            self.assertEqual(p["ceap_annee"], 2010)
            self.assertEqual(p["cap"], "non")
            self.assertEqual(p["date_derniere_formation"], "2025-01-10")
            self.assertEqual(p["grade"], "A1-5")
            f = MC.calculer_fiche(dict(p), auj=date(2026, 10, 7))
            self.assertEqual(f["retraite_txt"], "01/04/2027")                # né 15/02/1967, grade A
            p2 = q("SELECT p.*, c.* FROM personnel p JOIN personnel_carriere c ON c.personnel_id=p.id WHERE p.matricule='IMP-002'", un=True)
            f2 = MC.calculer_fiche(dict(p2), auj=date(2026, 10, 7))
            self.assertEqual(f2["retraite_txt"], "01/01/2026")               # né 10/11/1970, grade C
            self.assertTrue(f2["retraite_atteinte"])
            self.assertTrue(q("SELECT 1 FROM synoptique_membres WHERE comite='classe' AND nom_prenoms='CLASSE Importé' AND classe_id=?",
                              (self.classe_id,), un=True))
            self.assertFalse(q("SELECT 1 FROM synoptique_membres WHERE nom_prenoms='CLASSE Inconnue'", un=True))
        # second import : mise à jour (même matricule) et non doublon
        r = self.c.post("/synoptique/import", data={"fichier": (io.BytesIO(buf.getvalue()), "import.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Personnel : 0 agent(s) ajouté(s), 2 mis à jour", txt(r))
        with app.app_context():
            from app import q
            self.assertEqual(q("SELECT COUNT(*) n FROM personnel WHERE matricule LIKE 'IMP-%'", un=True)["n"], 2)
            self.assertEqual(q("SELECT COUNT(*) n FROM synoptique_membres WHERE nom_prenoms='SANTÉ Importé'", un=True)["n"], 1)
        html = txt(self.c.get("/synoptique/personnel"))
        self.assertIn("HOUNSOU Marie", html)
        self.assertIn("01/04/2027", html)
        # export complet
        r = self.c.get("/synoptique/export.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        for feuille in ("Identité", "Personnel", "Comité de santé", "Bureau de l'APE", "Comité de direction", "Comités de classe",
                        "Emplois du temps", "Planification", "Règlement"):
            self.assertIn(feuille, wb.sheetnames)
        cellules = [c.value for row in wb["Personnel"].iter_rows() for c in row]
        self.assertIn("HOUNSOU Marie", cellules)
        self.assertIn("01/04/2027", cellules)
        self.assertIn("Date de départ à la retraite", cellules)
        cellules = [c.value for row in wb["Règlement"].iter_rows() for c in row]
        self.assertIn("Les châtiments corporels sont formellement interdits.", cellules)
        self.assertIn("Jours 1-10 : SA1", [c.value for row in wb["Planification"].iter_rows() for c in row])
        # fichier sans en-tête exploitable
        wb = openpyxl.Workbook()
        wb.active.append(["a", "b"])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post("/synoptique/import", data={"fichier": (io.BytesIO(buf.getvalue()), "vide.xlsx")},
                        content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Aucune feuille exploitable", txt(r))

    def test_09_pdf_complet(self):
        for url in ("/synoptique/synoptique.pdf", "/synoptique/personnel.pdf", "/synoptique/edt/ci.pdf",
                    "/synoptique/planification/cm2.pdf", "/synoptique/reglement.pdf", "/synoptique/comites/sante.pdf",
                    f"/synoptique/comites/classe.pdf?classe_id={self.classe_id}"):
            r = self.c.get(url)
            self.assertEqual(r.status_code, 200, url)
            self.assertTrue(r.data.startswith(b"%PDF"), url)
        self.assertGreater(len(self.c.get("/synoptique/synoptique.pdf").data), 30000)

    def test_10_droits_et_interfaces(self):
        """Un enseignant lit le règlement et le comité de classe mais n'écrit pas ; le tableau synoptique et le personnel
        relèvent du directeur (m_perimetre_enseignant) ; H.synoptique_* est exposé aux autres modules."""
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex, q
            if not q("SELECT 1 FROM utilisateurs WHERE identifiant='ens_syno'", un=True):
                ex("INSERT INTO utilisateurs (identifiant, nom, role, mdp_hash, doit_changer_mdp, ecole_id) VALUES (?,?,?,?,0,1)",
                   ("ens_syno", "Enseignant test", "enseignant", generate_password_hash("secret123")))
        c2 = app.test_client()
        self.assertEqual(c2.post("/login", data={"identifiant": "ens_syno", "mdp": "secret123"}).status_code, 302)
        # v1.6 (m_perimetre_enseignant) : le tableau synoptique et le personnel relèvent du directeur ; l'enseignant
        # lit le règlement intérieur, les emplois du temps / planifications de SES niveaux et le comité de classe
        self.assertEqual(c2.get("/synoptique").status_code, 403)
        self.assertEqual(c2.get("/synoptique/personnel").status_code, 403)
        self.assertEqual(c2.get("/synoptique/reglement").status_code, 200)
        self.assertEqual(c2.get("/synoptique/comites/classe").status_code, 200)
        self.assertEqual(c2.post("/synoptique/identite", data={"memp": "x"}).status_code, 403)
        self.assertEqual(c2.get("/synoptique/personnel/0").status_code, 403)
        self.assertEqual(c2.post("/synoptique/reglement/retablir").status_code, 403)
        self.assertEqual(app.test_client().get("/synoptique").status_code, 302)      # non connecté -> login
        import m_synoptique
        with app.test_request_context("/"):
            from flask import g
            from app import q
            g.ecole = q("SELECT * FROM ecoles WHERE id=1", un=True)
            g.user = q("SELECT * FROM utilisateurs WHERE identifiant='admin'", un=True)
            Hn = m_synoptique._H
            self.assertTrue(hasattr(Hn, "synoptique_identite") and hasattr(Hn, "synoptique_personnel"))
            ident = Hn.synoptique_identite(1)
            self.assertEqual(next(x["valeur"] for x in ident if x["cle"] == "ccs"), "KOUAGOU Pierre")
            perso = Hn.synoptique_personnel(1)
            self.assertTrue(any(p["nom_complet"] == "HOUNSOU Marie" and p["retraite_txt"] == "01/04/2027" for p in perso))


class TestEditionPrivee(unittest.TestCase):
    """Rejoue toute la suite dans un processus séparé avec l'édition « privee » (le module est chargé une fois par processus)."""

    def test_privee(self):
        if EDITION != "publique":
            self.skipTest("déjà dans l'édition secondaire")
        env = dict(os.environ, GESCOL_EDITION_TEST="privee")
        env.pop("GESCOL_EDITION", None)
        env.pop("GESCOL_DATA", None)
        r = subprocess.run([sys.executable, "-m", "unittest", "tests.test_synoptique.TestSynoptique"], cwd=RACINE, env=env,
                           capture_output=True, text=True, timeout=600)
        self.assertEqual(r.returncode, 0, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
