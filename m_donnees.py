# -*- coding: utf-8 -*-
"""Données : import / export de GESCOL, échanges avec EducMaster et conversion PDF → Excel / CSV.

* /donnees                    : page d'accueil « Import / export » (liens vers les imports / exports de tous les modules
                                installés, selon les droits de l'utilisateur) + sauvegarde de la base
* /educmaster?cycle=primaire  : IMPORT-EXPORT EDUCMASTER (onglets Primaire / Secondaire)
    - dépôt du fichier .xlsx téléchargé depuis la plateforme EducMaster (glisser-déposer ou « Choisir un fichier ») ;
    - détection souple des colonnes (matricule EducMaster, NIP, nom, prénoms, sexe, date / lieu de naissance, classe,
      matières) avec rapport, correspondance modifiable (mémorisée : Sauvegarde / Charger / Réinitialiser) et aperçu ;
    - (a) mise à jour / création des élèves (matricule EducMaster, NIP, sexe, naissance) ;
    - (b) remplissage des notes d'une évaluation GESCOL DANS le fichier déposé (mise en forme conservée : seules les
      cellules des colonnes de matières reconnues sont écrites) puis téléchargement pour le recharger sur EducMaster.
* /conversion-pdf             : extraction des tableaux d'un PDF texte (pdfplumber) -> aperçu -> .xlsx ou .csv

Les fichiers déposés sont gardés temporairement dans <INSTANCE_DIR>/tmp_donnees (nettoyés au bout de 2 jours).
"""
import csv
import io
import json
import os
import re
import secrets
import time
from datetime import date

from flask import Response, abort, flash, g, redirect, render_template, request, send_file, url_for

import evaluations_core as EC
from calcul import cycle_du_niveau

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CYCLES = {"primaire": "Primaire", "secondaire": "Secondaire"}
RE_JETON = re.compile(r"^[0-9a-f]{16}$")
DUREE_TMP = 2 * 24 * 3600
MAX_LIGNES_EM = 3000
MAX_PAGES_PDF = 200

# Champs d'identification reconnus dans un fichier EducMaster : (clé, libellé)
CHAMPS_ID = [
    ("matricule_em", "Matricule EducMaster"), ("matricule", "Matricule GESCOL (école)"), ("nip", "NIP"),
    ("nom", "Nom"), ("prenoms", "Prénoms"), ("nom_complet", "Nom et prénoms"), ("sexe", "Sexe"),
    ("date_naissance", "Date de naissance"), ("lieu_naissance", "Lieu de naissance"),
    ("naissance", "Date et lieu de naissance"), ("classe", "Classe"),
]
LIB_ID = dict(CHAMPS_ID)
# intitulés « numéro d'ordre » à ignorer
IGNORES = {"n", "no", "num", "numero", "n ordre", "no ordre", "numero d ordre", "ordre", "rang", "n d ordre", "nr",
           "observation", "observations", "decision", "moyenne", "moy", "total", "rang", "mention", "appreciation"}
# alias de matières (intitulé normalisé -> code GESCOL), en plus du nom et du code des matières
ALIAS_MATIERES = {
    "primaire": {"communication orale": "COM.OR", "co": "COM.OR", "dictee": "DIC", "lecture": "COMP.EC",
                 "comprehension de l ecrit": "COMP.EC", "ce": "COMP.EC", "maths": "MATH", "math": "MATH",
                 "mathematique": "MATH", "mathematiques": "MATH", "es": "ED.SOC", "education sociale": "ED.SOC",
                 "est": "ED.SCI", "education scientifique": "ED.SCI", "ea": "ART.P", "education artistique": "ART.P",
                 "ea plastique": "ART.P", "ea vivant": "ART.V", "anglais": "ANG", "expression ecrite": "EXP.EC",
                 "production d ecrit": "EXP.EC", "ee": "EXP.EC", "eps": "EPS", "education physique": "EPS"},
    "secondaire": {"communication ecrite": "FR.COM", "dictee": "FR.DIC", "expression ecrite": "FR.EXP",
                   "lecture": "FR.LEC", "maths": "MATH", "math": "MATH", "mathematiques": "MATH",
                   "hist geo": "HG", "histoire geographie": "HG", "anglais": "ANG", "espagnol": "ESP",
                   "allemand": "ALL", "pct": "PCT", "physique chimie": "PCT", "svt": "SVT", "philosophie": "PHILO",
                   "eps": "EPS"},
}


def normaliser(t):
    import unicodedata
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def lettre(i):
    """0 -> A, 25 -> Z, 26 -> AA."""
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def role_identite(h):
    """Intitulé normalisé -> clé de CHAMPS_ID ou None (détection souple)."""
    if not h:
        return None
    mots = h.split()
    if "educmaster" in h.replace(" ", "") or ("matricule" in h and "educ" in h) or h in ("id em", "identifiant em"):
        return "matricule_em"
    if h in ("nip", "n i p") or mots[0] == "nip" or "identification personnel" in h:
        return "nip"
    if "prenom" in h and "nom" in mots:
        return "nom_complet"
    if h in ("nom complet", "noms et prenoms", "eleve", "apprenant", "nom de l eleve et prenoms", "identite",
             "nom et prenom de l eleve", "nom et prenoms de l eleve"):
        return "nom_complet"
    if mots[0] in ("prenom", "prenoms"):
        return "prenoms"
    if mots[0] in ("nom", "noms") and "prenom" not in h:
        return "nom"
    if mots[0] in ("sexe", "genre") or h in ("s", "sx", "g f", "m f"):
        return "sexe"
    if "naissance" in h and "lieu" in h and "date" in h:
        return "naissance"
    if h in ("ne e le a", "ne le a", "ne e le et a"):
        return "naissance"
    if "lieu" in h and ("naissance" in h or h.startswith("lieu")) or h in ("ne e a", "ne a"):
        return "lieu_naissance"
    if "naissance" in h or h in ("ne e le", "ne le", "dn", "date nais", "date de nais"):
        return "date_naissance"
    if mots[0] in ("classe", "cours", "salle", "niveau", "division"):
        return "classe"
    if mots[0] == "matricule" or h in ("mat", "n matricule", "numero matricule", "no matricule", "n mle", "mle"):
        return "matricule"
    return None


def matiere_de(h, matieres, cycle):
    """Intitulé normalisé -> matière GESCOL (Row) ou None."""
    if not h or h in IGNORES:
        return None
    par_code = {normaliser(m["code"]): m for m in matieres}
    for m in matieres:
        if h in (normaliser(m["nom"]), normaliser(m["code"])):
            return m
    code = ALIAS_MATIERES.get(cycle, {}).get(h)
    if code and normaliser(code) in par_code:
        return par_code[normaliser(code)]
    for m in matieres:  # préfixe (« Mathématiques /20 », « Anglais (note) »…)
        n = normaliser(m["nom"])
        if len(h) >= 5 and (n.startswith(h) or h.startswith(n)):
            return m
    return None


def separer_naissance(v):
    """« 12/03/2014 à Cotonou » -> ('12/03/2014', 'Cotonou')."""
    s = str(v or "").strip()
    m = re.match(r"^\s*(.+?)\s+(?:à|a|A|À)\s+(.+)$", s)
    if m:
        return m.group(1), m.group(2)
    return s, ""


def enregistrer(app, H):
    q, ex = H.q, H.ex
    X = H.xlsx

    # ============================================================ outils fichiers temporaires
    def dossier_tmp():
        d = os.path.join(H.config.INSTANCE_DIR, "tmp_donnees")
        os.makedirs(d, exist_ok=True)
        return d

    def nettoyer_tmp():
        d = dossier_tmp()
        limite = time.time() - DUREE_TMP
        for nom in os.listdir(d):
            p = os.path.join(d, nom)
            try:
                if os.path.getmtime(p) < limite:
                    os.remove(p)
            except OSError:
                pass

    def chemin_tmp(prefixe, jeton, ext):
        if not RE_JETON.match(jeton or ""):
            abort(404)
        return os.path.join(dossier_tmp(), f"{prefixe}_e{H.eco()}_{jeton}.{ext}")

    # ============================================================ accueil Import / export
    @app.route("/donnees")
    @H.connexion_requise("donnees")
    def donnees_accueil():
        role = g.user["role"]
        # (groupe, [(endpoint, libellé, description, droit ou None, rôles autorisés ou None)])
        groupes = [
            ("Élèves", "👤", [
                ("eleves_export_xlsx", "Exporter la liste des élèves (Excel)", "Tous les élèves de l'école, au format .xlsx.", "eleves", None),
                ("eleves_export", "Exporter la liste des élèves (CSV)", "Même liste au format CSV (séparateur « ; »).", "eleves", None),
                ("eleves_modele_xlsx", "Modèle Excel d'inscription", "Fichier vierge à remplir pour inscrire les élèves en nombre.", "eleves", None),
                ("eleves_import", "Importer des élèves (Excel / CSV)", "Inscription ou mise à jour en masse avec aperçu avant validation.", "eleves", None),
            ]),
            ("Notes et résultats", "✎", [
                ("eval_accueil", "Notes par évaluation : modèle, import, export Excel",
                 "Ouvrez une évaluation : boutons « Modèle Excel », « Importer » et « Excel ».", "evaluations", None),
                ("notes_accueil", "Notes par trimestre", "Saisie et bulletins trimestriels.", "notes", None),
                ("eval_statistiques", "Statistiques d'une évaluation (Excel / PDF)", "Statistiques par sexe et par matière.", "evaluations", None),
                ("bilan_annuel", "Bilan annuel (Excel / PDF)", "Moyennes annuelles et décisions.", "evaluations", None),
                ("synthese_accueil", "Synthèse École / UP / CS (Excel / PDF)", "Synthèse de l'évaluation par classe et par niveau.", "synthese", None),
                ("affichage_accueil", "Affichage des résultats (modèle Excel, import)", "Liste d'affichage des résultats.", "affichage", None),
            ]),
            ("Présences", "🗓", [
                ("registre_accueil", "Registre d'appel (import des élèves, PDF)", "Registre par demi-journée.", "registre", None),
                ("releve_mensuel", "Relevés mensuels (PDF)", "Tableau des relevés mensuels de l'année.", "registre", None),
            ]),
            ("Gestion", "₣", [
                ("paie_accueil", "Paie : livre de paie (CSV)", "Ouvrez un mois puis « Livre de paie ».", "paie", None),
                ("cantine_recap", "Cantine : récapitulatif", "Récapitulatif des paiements de cantine.", "cantine", None),
                ("fondateur_xlsx", "Tableau du fondateur (Excel)", "Indicateurs consolidés de toutes les écoles.", None, ("fondateur", "admin")),
                ("demandes_ecoles_csv", "Demandes d'adhésion des écoles (CSV)", "Liste des écoles ayant demandé à rejoindre GESCOL.", "ecoles", ("fondateur", "admin")),
            ]),
            ("Échanges et conversions", "⇅", [
                ("educmaster_accueil", "Import-export EducMaster", "Élèves et notes avec la plateforme EducMaster (primaire, secondaire).", "donnees", None),
                ("conversion_pdf", "Conversion PDF → Excel / CSV", "Récupérer les tableaux d'un PDF (liste, relevé…).", "donnees", None),
            ]),
        ]
        blocs = []
        for titre, ico, liens in groupes:
            ok = [(ep, lib, desc) for ep, lib, desc, droit, roles in liens
                  if ep in app.view_functions and (droit is None or H.peut(droit)) and (roles is None or role in roles)]
            if ok:
                blocs.append((titre, ico, ok))
        return render_template("donnees_accueil.html", blocs=blocs,
                               sauvegarde="sauvegarde" in app.view_functions and H.peut("parametres"))

    # ============================================================ EducMaster : analyse du fichier
    def cycle_choisi():
        c = request.values.get("cycle", "primaire")
        return c if c in CYCLES else "primaire"

    def cle_etat(cycle):
        return f"e{H.eco()}:educmaster:{cycle}:fichier"

    def cle_mapping(cycle):
        # même clé que la barre « Sauvegarde / Charger » (/documents-modules/educmaster:<cycle>:mapping)
        return f"e{H.eco()}:educmaster:{cycle}:mapping"

    def matieres_cycle(cycle):
        return list(EC.matieres_cycle(cycle))

    def cle_colonne(entetes, j):
        """Nom de champ stable pour la colonne j (intitulé normalisé ; suffixe si l'intitulé est répété)."""
        base = normaliser(entetes[j]).replace(" ", "_")[:40] or f"col_{lettre(j)}"
        n = sum(1 for k in range(j) if (normaliser(entetes[k]).replace(" ", "_")[:40] or f"col_{lettre(k)}") == base)
        return "c_" + base + (f"_{n + 1}" if n else "")

    def ouvrir(chemin, lecture=True):
        import openpyxl  # import paresseux
        return openpyxl.load_workbook(chemin, read_only=lecture, data_only=lecture,
                                      keep_vba=chemin.lower().endswith(".xlsm") and not lecture)

    def trouver_entete(lignes, matieres, cycle):
        """Indice (0..) de la ligne d'en-tête parmi les 30 premières lignes, et son score."""
        meilleur, score_max = 0, 0
        for i, r in enumerate(lignes[:30]):
            score = 0
            for v in r:
                h = normaliser(v)
                if not h or len(h) > 60:
                    continue
                if role_identite(h):
                    score += 2
                elif matiere_de(h, matieres, cycle):
                    score += 1
            if score > score_max:
                meilleur, score_max = i, score
        return meilleur, score_max

    def analyser(chemin, cycle, feuille=None, mapping=None):
        """Lit le fichier déposé et reconnaît les colonnes. mapping : {cle_colonne: '' | '__ignorer__' |
        'id:<champ>' | 'mat:<code>'} (choix de l'utilisateur, prioritaire sur la détection automatique)."""
        mapping = mapping or {}
        matieres = matieres_cycle(cycle)
        par_code = {m["code"]: m for m in matieres}
        wb = ouvrir(chemin)
        try:
            feuilles = wb.sheetnames
            candidats = []
            for nom in feuilles:
                ws = wb[nom]
                lignes = []
                for r in ws.iter_rows(values_only=True, max_row=MAX_LIGNES_EM + 40):
                    lignes.append(list(r))
                idx, score = trouver_entete(lignes, matieres, cycle)
                candidats.append((nom, lignes, idx, score))
        finally:
            wb.close()
        if feuille in feuilles:
            nom, lignes, idx, score = next(c for c in candidats if c[0] == feuille)
        else:
            nom, lignes, idx, score = max(candidats, key=lambda c: c[3])
        entetes_brutes = lignes[idx] if lignes else []
        largeur = max([len(r) for r in lignes[idx:idx + 200]] + [len(entetes_brutes), 0])
        entetes = [("" if v is None else str(v).strip()) for v in list(entetes_brutes) + [None] * (largeur - len(entetes_brutes))]
        donnees = []
        for k, r in enumerate(lignes[idx + 1:], start=idx + 2):  # k = numéro de ligne Excel
            r = list(r) + [None] * (largeur - len(r))
            if any(v not in (None, "") for v in r):
                donnees.append((k, r))
        colonnes, deja = [], set()
        for j, h in enumerate(entetes):
            hn = normaliser(h)
            cle = cle_colonne(entetes, j)
            auto = ""
            rid = role_identite(hn) if len(hn) <= 60 else None
            if rid and rid not in deja:
                auto = "id:" + rid
            elif not rid:
                m = matiere_de(hn, matieres, cycle)
                if m:
                    auto = "mat:" + m["code"]
            choix = mapping.get(cle, "")
            role = choix if choix and (choix == "__ignorer__" or choix.startswith(("id:", "mat:"))) else auto
            if role.startswith("id:") and role[3:] not in LIB_ID:
                role = auto
            if role.startswith("mat:") and role[4:] not in par_code:
                role = auto
            if role.startswith("id:"):
                if role[3:] in deja:
                    role = ""  # un même champ ne peut être lu que dans une colonne
                else:
                    deja.add(role[3:])
            exemples = [v for _, r in donnees[:40] for v in [r[j]] if v not in (None, "")][:3]
            colonnes.append({"j": j, "lettre": lettre(j), "entete": h, "cle": cle, "auto": auto, "role": role,
                             "force": bool(choix), "exemples": [str(x)[:30] for x in exemples]})
        roles = {c["role"][3:]: c["j"] for c in colonnes if c["role"].startswith("id:")}
        mats = {c["j"]: par_code[c["role"][4:]] for c in colonnes if c["role"].startswith("mat:")}
        # « Matricule » seul dans un fichier EducMaster = matricule EducMaster
        if "matricule" in roles and "matricule_em" not in roles and not any(
                c["force"] and c["role"] == "id:matricule" for c in colonnes):
            roles["matricule_em"] = roles.pop("matricule")
            for c in colonnes:
                if c["role"] == "id:matricule":
                    c["role"] = "id:matricule_em"
        return {"feuilles": feuilles, "feuille": nom, "ligne_entete": idx + 1, "score": score, "entetes": entetes,
                "colonnes": colonnes, "donnees": donnees, "roles": roles, "matieres": mats,
                "identifiable": bool({"matricule_em", "matricule", "nip", "nom_complet"} & set(roles))
                or {"nom", "prenoms"} <= set(roles)}

    def etat(cycle):
        e = H.doc_lire(cle_etat(cycle), None)
        if not e or not RE_JETON.match(e.get("jeton", "")):
            return None
        chemin = chemin_tmp("em", e["jeton"], e.get("ext", "xlsx"))
        if not os.path.exists(chemin):
            return None
        e["chemin"] = chemin
        return e

    def mapping_courant(cycle, e):
        """Correspondance mémorisée (Sauvegarde) complétée par celle appliquée à ce fichier."""
        m = dict(H.doc_lire(cle_mapping(cycle), {}) or {})
        m.update(e.get("mapping") or {})
        return {k: v for k, v in m.items() if isinstance(v, str)}

    def analyse_etat(cycle, e):
        return analyser(e["chemin"], cycle, e.get("feuille"), mapping_courant(cycle, e))

    # --- repérage des élèves
    def eleves_ecole(cycle=None):
        rows = q("""SELECT e.*, c.nom classe_nom, c.niveau FROM eleves e JOIN classes c ON c.id=e.classe_id
                    WHERE c.ecole_id=? AND e.statut='inscrit'""", (H.eco(),))
        return [r for r in rows if cycle is None or cycle_du_niveau(r["niveau"]) == cycle]

    def index_eleves(eleves):
        idx = {"em": {}, "mat": {}, "nip": {}, "np": {}}
        for e in eleves:
            if e["matricule_educmaster"]:
                idx["em"][str(e["matricule_educmaster"]).strip().upper()] = e
            if e["matricule"]:
                idx["mat"][str(e["matricule"]).strip().upper()] = e
            if e["nip"]:
                idx["nip"][str(e["nip"]).strip().upper()] = e
            n, p = normaliser(e["nom"]), normaliser(e["prenoms"])
            idx["np"].setdefault(f"{n} {p}".strip(), e)
            idx["np"].setdefault(f"{p} {n}".strip(), e)
        return idx

    def lire_ligne(an, r):
        """Valeurs d'identité d'une ligne du fichier."""
        roles = an["roles"]

        def v(k):
            return X.valeur(r, roles.get(k)) if k in roles else ""
        d = {k: v(k) for k in LIB_ID}
        if d["nom_complet"] and not (d["nom"] or d["prenoms"]):
            parts = d["nom_complet"].split()
            d["nom"], d["prenoms"] = (parts[0], " ".join(parts[1:])) if parts else ("", "")
        if d["naissance"]:
            dn, ln = separer_naissance(d["naissance"])
            d["date_naissance"] = d["date_naissance"] or dn
            d["lieu_naissance"] = d["lieu_naissance"] or ln
        if "date_naissance" in roles:
            brut = r[roles["date_naissance"]]
            if hasattr(brut, "strftime"):
                d["date_naissance"] = brut.strftime("%Y-%m-%d")
        d["date_iso"] = X.date_iso(d["date_naissance"]) if d["date_naissance"] else ""
        d["sexe_ok"] = X.sexe(d["sexe"]) if d["sexe"] else ""
        return d

    def retrouver(idx, d):
        for cle, champ in (("em", "matricule_em"), ("mat", "matricule_em"), ("mat", "matricule"), ("em", "matricule"),
                           ("nip", "nip")):
            val = (d.get(champ) or "").strip().upper()
            if val and val in idx[cle]:
                return idx[cle][val]
        nc = normaliser(d.get("nom_complet") or "") or normaliser(f"{d.get('nom', '')} {d.get('prenoms', '')}")
        return idx["np"].get(nc) if nc else None

    def retrouver_classe(texte, classes):
        t = normaliser(texte).replace(" ", "")
        if not t:
            return None
        for c in classes:
            if normaliser(c["nom"]).replace(" ", "") == t:
                return c
        meme = [c for c in classes if normaliser(c["niveau"]).replace(" ", "") == t]
        return meme[0] if len(meme) == 1 else None

    # ============================================================ EducMaster : pages
    @app.route("/educmaster")
    @H.connexion_requise("donnees")
    def educmaster_accueil():
        cycle = cycle_choisi()
        e = etat(cycle)
        an, erreur = None, None
        if e:
            try:
                an = analyse_etat(cycle, e)
            except Exception as exc:  # noqa: BLE001 — fichier devenu illisible
                erreur = f"Le fichier déposé ne peut plus être lu ({exc.__class__.__name__}). Déposez-le de nouveau."
        classes = [c for c in H.classes_autorisees() if cycle_du_niveau(c["niveau"]) == cycle]
        evaluations = []
        for c in classes:
            for ev in EC.evaluations_classe(c["id"]):
                evaluations.append({"id": ev["id"], "libelle": f"{c['nom']} — {ev['titre']}"})
        if an:
            an["apercu"] = an["donnees"][:15]
            an["nb_lignes"] = len(an["donnees"])
        choix = [("", "— détection automatique —"), ("__ignorer__", "— ignorer cette colonne —")] \
            + [("id:" + k, lib) for k, lib in CHAMPS_ID] \
            + [("mat:" + m["code"], f"Matière : {m['nom']}") for m in matieres_cycle(cycle)]
        return render_template("educmaster_accueil.html", cycle=cycle, CYCLES=CYCLES, e=e, an=an, erreur=erreur,
                               classes=classes, evaluations=evaluations, choix=choix, LIB_ID=LIB_ID,
                               rapport=(e or {}).get("rapport"))

    @app.route("/educmaster/deposer", methods=["POST"])
    @H.connexion_requise("donnees")
    def educmaster_deposer():
        cycle = cycle_choisi()
        f = request.files.get("fichier")
        if not f or not f.filename:
            flash("Choisissez le fichier Excel téléchargé depuis EducMaster.", "warning")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        nom = f.filename.lower()
        if nom.endswith(".xls"):
            flash("Ancien format Excel (.xls) : ouvrez le fichier dans Excel ou LibreOffice et enregistrez-le au "
                  "format .xlsx, puis déposez-le de nouveau.", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        if not nom.endswith((".xlsx", ".xlsm")):
            flash("Format non reconnu : déposez le modèle Excel (.xlsx) téléchargé depuis EducMaster.", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        contenu = f.read()
        if not contenu.startswith(b"PK"):
            flash("Ce fichier n'est pas un classeur Excel valide (.xlsx).", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            flash("Le composant Excel (openpyxl) n'est pas installé : l'import EducMaster est indisponible.", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        nettoyer_tmp()
        ancien = etat(cycle)
        if ancien:
            try:
                os.remove(ancien["chemin"])
            except OSError:
                pass
        jeton = secrets.token_hex(8)
        ext = "xlsm" if nom.endswith(".xlsm") else "xlsx"
        chemin = chemin_tmp("em", jeton, ext)
        with open(chemin, "wb") as out:
            out.write(contenu)
        try:
            an = analyser(chemin, cycle, None, H.doc_lire(cle_mapping(cycle), {}) or {})
        except Exception:  # noqa: BLE001 — classeur corrompu, protégé par mot de passe…
            os.remove(chemin)
            flash("Classeur illisible (fichier abîmé ou protégé par un mot de passe). Réenregistrez-le au format .xlsx.",
                  "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        H.doc_ecrire(cle_etat(cycle), {"jeton": jeton, "ext": ext, "nom": os.path.basename(f.filename)[:120],
                                       "depose_le": H.maintenant(), "feuille": an["feuille"], "mapping": {}})
        H.journal("educmaster_depot", f"{cycle} : {f.filename[:80]} ({len(an['donnees'])} ligne(s))")
        reconnus = [LIB_ID[k] for k in an["roles"]]
        msg = (f"Fichier « {f.filename} » analysé : feuille « {an['feuille']} », en-têtes en ligne {an['ligne_entete']}, "
               f"{len(an['donnees'])} ligne(s). Reconnu : " + (", ".join(reconnus) or "aucune colonne d'identité")
               + f" ; {len(an['matieres'])} colonne(s) de matières.")
        flash(msg, "success" if an["identifiable"] else "warning")
        if not an["identifiable"]:
            flash("Aucune colonne ne permet d'identifier les élèves (matricule, NIP ou nom + prénoms) : indiquez les "
                  "colonnes dans le tableau « Correspondance des colonnes ».", "warning")
        return redirect(url_for("educmaster_accueil", cycle=cycle))

    @app.route("/educmaster/mapping", methods=["POST"])
    @H.connexion_requise("donnees")
    def educmaster_mapping():
        """Applique la correspondance des colonnes choisie (et la feuille) au fichier déposé."""
        cycle = cycle_choisi()
        e = etat(cycle)
        if not e:
            flash("Déposez d'abord un fichier EducMaster.", "warning")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        mapping = {k: v[:40] for k, v in request.form.items() if k.startswith("c_") and len(k) <= 50}
        e.pop("chemin", None)
        e["mapping"] = mapping
        if request.form.get("feuille"):
            e["feuille"] = request.form["feuille"][:60]
        if request.form.get("memoriser"):
            H.doc_ecrire(cle_mapping(cycle), mapping)
        H.doc_ecrire(cle_etat(cycle), e)
        flash("Correspondance des colonnes appliquée" + (" et mémorisée." if request.form.get("memoriser") else "."), "success")
        return redirect(url_for("educmaster_accueil", cycle=cycle) + "#apercu")

    @app.route("/educmaster/retirer", methods=["POST"])
    @H.connexion_requise("donnees")
    def educmaster_retirer():
        cycle = cycle_choisi()
        e = etat(cycle)
        if e:
            try:
                os.remove(e["chemin"])
            except OSError:
                pass
        H.doc_ecrire(cle_etat(cycle), {})
        flash("Fichier retiré.", "success")
        return redirect(url_for("educmaster_accueil", cycle=cycle))

    # --- (a) élèves
    @app.route("/educmaster/eleves", methods=["POST"])
    @H.connexion_requise("donnees")
    def educmaster_eleves():
        cycle = cycle_choisi()
        e = etat(cycle)
        if not e:
            flash("Déposez d'abord un fichier EducMaster.", "warning")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        an = analyse_etat(cycle, e)
        if not an["identifiable"]:
            flash("Impossible d'identifier les élèves : indiquez la colonne du matricule EducMaster ou des noms.", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        simuler = bool(request.form.get("simuler"))
        creer = bool(request.form.get("creer"))
        classes = [c for c in H.classes_autorisees() if cycle_du_niveau(c["niveau"]) == cycle]
        defaut = request.form.get("classe_defaut", type=int)
        defaut = next((c for c in classes if c["id"] == defaut), None)
        idx = index_eleves(eleves_ecole())
        db = H.get_db()
        stats = {"maj": 0, "inchanges": 0, "crees": 0, "ignores": 0}
        details = []
        for k, r in an["donnees"][:MAX_LIGNES_EM]:
            d = lire_ligne(an, r)
            qui = (d["nom_complet"] or f"{d['nom']} {d['prenoms']}").strip() or d["matricule_em"] or "?"
            if not (d["nom"] or d["prenoms"] or d["matricule_em"] or d["nip"]):
                continue
            el = retrouver(idx, d)
            maj = {}
            if d["matricule_em"] and (el is None or (el["matricule_educmaster"] or "") != d["matricule_em"]):
                maj["matricule_educmaster"] = d["matricule_em"][:40]
            if d["nip"] and (el is None or (el["nip"] or "") != d["nip"]):
                maj["nip"] = d["nip"][:40]
            if d["sexe_ok"] and (el is None or el["sexe"] != d["sexe_ok"]):
                maj["sexe"] = d["sexe_ok"]
            if d["date_iso"] and (el is None or el["date_naissance"] != d["date_iso"]):
                maj["date_naissance"] = d["date_iso"]
            if d["lieu_naissance"] and (el is None or (el["lieu_naissance"] or "") != d["lieu_naissance"]):
                maj["lieu_naissance"] = d["lieu_naissance"][:80]
            if el is not None:
                if el["classe_id"] not in {c["id"] for c in H.classes_autorisees()}:
                    stats["ignores"] += 1
                    details.append([k, qui, "Ignoré", "élève d'une classe qui ne vous est pas accessible"])
                    continue
                if not maj:
                    stats["inchanges"] += 1
                    continue
                if not simuler:
                    cols = list(maj)
                    db.execute(f"UPDATE eleves SET {', '.join(c + '=?' for c in cols)} WHERE id=?",
                               [maj[c] for c in cols] + [el["id"]])
                stats["maj"] += 1
                details.append([k, qui, "Mis à jour", ", ".join(f"{c} = {v}" for c, v in maj.items())])
                continue
            # élève inconnu de GESCOL
            if not creer:
                stats["ignores"] += 1
                details.append([k, qui, "Ignoré", "élève absent de GESCOL (cochez « créer les élèves absents »)"])
                continue
            cl = retrouver_classe(d["classe"], classes) or defaut
            if not (d["nom"] and d["prenoms"]):
                stats["ignores"] += 1
                details.append([k, qui, "Ignoré", "nom ou prénoms manquant"])
                continue
            if cl is None:
                stats["ignores"] += 1
                details.append([k, qui, "Ignoré", f"classe « {d['classe'] or '?'} » introuvable : choisissez une classe par défaut"])
                continue
            if not d["sexe_ok"]:
                stats["ignores"] += 1
                details.append([k, qui, "Ignoré", "sexe manquant ou illisible (M / F)"])
                continue
            if not simuler:
                cur = db.execute("""INSERT INTO eleves (nom, prenoms, sexe, date_naissance, lieu_naissance, classe_id,
                                    statut, date_inscription, annee, matricule_educmaster, nip)
                                    VALUES (?,?,?,?,?,?,'inscrit',?,?,?,?)""",
                                 (d["nom"].upper()[:80], d["prenoms"][:120], d["sexe_ok"], d["date_iso"] or None,
                                  d["lieu_naissance"][:80] or None, cl["id"], date.today().isoformat(),
                                  H.config.ANNEE_ACTIVE, d["matricule_em"][:40] or None, d["nip"][:40] or None))
                eid = cur.lastrowid
                cyc = {"maternelle": "M", "primaire": "P"}.get(cycle_du_niveau(cl["niveau"]), "S")
                db.execute("UPDATE eleves SET matricule=? WHERE id=?",
                           (f"{H.config.ECOLE['sigle']}-{H.config.ANNEE_ACTIVE[:4]}-{cyc}-{eid:04d}", eid))
                nouvel = dict(zip(("id", "nom", "prenoms", "matricule", "matricule_educmaster", "nip"),
                                  (eid, d["nom"].upper(), d["prenoms"], "", d["matricule_em"], d["nip"])))
                nouvel["classe_id"] = cl["id"]
                idx["np"].setdefault(normaliser(f"{d['nom']} {d['prenoms']}"), nouvel)
                if d["matricule_em"]:
                    idx["em"][d["matricule_em"].upper()] = nouvel
            stats["crees"] += 1
            details.append([k, qui, "Créé", f"classe {cl['nom']}"])
        if not simuler:
            db.commit()
            H.journal("educmaster_eleves", f"{cycle} : {stats['maj']} mis à jour, {stats['crees']} créé(s)")
        e.pop("chemin", None)
        e["rapport"] = {"titre": ("SIMULATION (rien n'a été enregistré)" if simuler else "Mise à jour des élèves")
                        + f" — {H.maintenant()}", "stats": stats, "details": details[:300]}
        H.doc_ecrire(cle_etat(cycle), e)
        flash(("Simulation : " if simuler else "")
              + f"{stats['maj']} élève(s) mis à jour, {stats['crees']} créé(s), {stats['inchanges']} inchangé(s), "
              f"{stats['ignores']} ignoré(s).", "info" if simuler else "success")
        return redirect(url_for("educmaster_accueil", cycle=cycle) + "#rapport")

    # --- (b) notes : remplir le fichier EducMaster
    @app.route("/educmaster/notes", methods=["POST"])
    @H.connexion_requise("donnees")
    def educmaster_notes():
        cycle = cycle_choisi()
        e = etat(cycle)
        if not e:
            flash("Déposez d'abord un fichier EducMaster.", "warning")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        ev = EC.evaluation(request.form.get("evaluation", type=int) or 0)
        if not ev:
            flash("Choisissez l'évaluation dont les notes doivent être recopiées.", "warning")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        cl = H.classe_ou_404(ev["classe_id"])
        H.verifier_classe(cl["id"])
        if cl["cycle"] != cycle:
            abort(400)
        an = analyse_etat(cycle, e)
        if not an["matieres"]:
            flash("Aucune colonne de matière n'est reconnue dans le fichier : indiquez-les dans la correspondance des "
                  "colonnes.", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        if not an["identifiable"]:
            flash("Impossible d'identifier les élèves dans le fichier (matricule ou noms).", "danger")
            return redirect(url_for("educmaster_accueil", cycle=cycle))
        matieres, lignes = EC.resultats(ev, cl)
        ids_ev = {m["id"] for m in matieres}
        notes = {l["eleve"]["id"]: l for l in lignes}
        idx = index_eleves([l["eleve"] for l in lignes])
        wb = ouvrir(e["chemin"], lecture=False)
        ws = wb[an["feuille"]]
        ecrites, reconnues, inconnus, sans_note, hors_ev = 0, 0, [], 0, set()
        for k, r in an["donnees"]:
            d = lire_ligne(an, r)
            if not (d["nom"] or d["prenoms"] or d["matricule_em"] or d["nip"] or d["matricule"]):
                continue
            el = retrouver(idx, d)
            if el is None:
                inconnus.append(f"ligne {k} ({(d['nom_complet'] or (d['nom'] + ' ' + d['prenoms'])).strip() or d['matricule_em']})")
                continue
            reconnues += 1
            l = notes[el["id"]]
            for j, m in an["matieres"].items():
                if m["id"] not in ids_ev:
                    hors_ev.add(m["nom"])
                    continue
                v = l["notes"].get(m["id"])
                if v is None:
                    sans_note += 1
                    continue
                cell = ws.cell(row=k, column=j + 1)
                if type(cell).__name__ == "MergedCell":
                    continue
                cell.value = int(v) if float(v).is_integer() else float(v)
                ecrites += 1
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        buf.seek(0)
        H.journal("educmaster_notes", f"{cl['nom']} {ev['libelle']} : {ecrites} note(s) écrite(s)")
        msg = (f"Fichier EducMaster rempli avec « {ev['titre']} » ({cl['nom']}) : {reconnues} élève(s) reconnu(s), "
               f"{ecrites} note(s) écrite(s)")
        if sans_note:
            msg += f", {sans_note} case(s) laissée(s) telle(s) quelle(s) (pas de note dans GESCOL)"
        flash(msg + ".", "success")
        if inconnus:
            flash(f"{len(inconnus)} ligne(s) du fichier sans élève correspondant dans {cl['nom']} : "
                  + ", ".join(inconnus[:10]) + (" …" if len(inconnus) > 10 else ""), "warning")
        if hors_ev:
            flash("Matières du fichier non évaluées dans cette évaluation (non remplies) : " + ", ".join(sorted(hors_ev)),
                  "info")
        base = re.sub(r"[^A-Za-z0-9_-]+", "_", os.path.splitext(e.get("nom") or "educmaster")[0])[:60]
        suffixe = re.sub(r"[^A-Za-z0-9]+", "_", normaliser(f"{cl['nom']} {ev['libelle']}"))
        ext = e.get("ext", "xlsx")
        mime = MIME_XLSX if ext == "xlsx" else "application/vnd.ms-excel.sheet.macroEnabled.12"
        return send_file(buf, mimetype=mime, as_attachment=True, download_name=f"{base}_GESCOL_{suffixe}.{ext}")

    # ============================================================ Conversion PDF → Excel / CSV
    def cle_conversion():
        return f"e{H.eco()}:conversion:dernier"

    def nettoyer_tableau(t):
        """Supprime les lignes et colonnes entièrement vides ; cellules en texte sur une ligne."""
        rows = [[("" if c is None else re.sub(r"\s*\n\s*", " ", str(c)).strip()) for c in r] for r in t if r]
        rows = [r for r in rows if any(r)]
        if not rows:
            return []
        n = max(len(r) for r in rows)
        rows = [r + [""] * (n - len(r)) for r in rows]
        garder = [j for j in range(n) if any(r[j] for r in rows)]
        return [[r[j] for j in garder] for r in rows]

    def tableaux_texte(texte):
        """Repli : lignes de texte découpées sur les grands espaces (2 espaces ou plus / tabulations)."""
        rows = [re.split(r"\s{2,}|\t", l.strip()) for l in texte.splitlines() if l.strip()]
        return rows if sum(1 for r in rows if len(r) > 1) >= 2 else []

    def extraire(contenu):
        """PDF -> {pages, tableaux: [{page, n, lignes, methode}], caracteres}. ValueError si illisible."""
        try:
            import pdfplumber  # import paresseux
        except ImportError as exc:
            raise RuntimeError("Le composant de lecture PDF (pdfplumber) n'est pas installé : la conversion est "
                               "indisponible sur cet ordinateur.") from exc
        tableaux, caracteres, pages = [], 0, 0
        try:
            with pdfplumber.open(io.BytesIO(contenu)) as pdf:
                pages = len(pdf.pages)
                for num, page in enumerate(pdf.pages[:MAX_PAGES_PDF], 1):
                    caracteres += len(page.chars)
                    if not page.chars:
                        continue
                    trouves = [nettoyer_tableau(t) for t in (page.extract_tables() or [])]
                    trouves = [t for t in trouves if t and (len(t) > 1 or len(t[0]) > 1)]
                    methode = "lignes du tableau"
                    if not trouves:
                        par_texte = page.extract_tables({"vertical_strategy": "text",
                                                         "horizontal_strategy": "text"}) or []
                        par_texte = [nettoyer_tableau(x) for x in par_texte]
                        t = max(par_texte, key=len) if par_texte else []
                        if t and len(t) > 1 and len(t[0]) > 1:
                            trouves, methode = [t], "alignement du texte"
                        else:
                            t = tableaux_texte(page.extract_text() or "")
                            if t:
                                trouves, methode = [nettoyer_tableau(t)], "lignes de texte"
                    for n, t in enumerate(trouves, 1):
                        tableaux.append({"page": num, "n": n, "lignes": t, "methode": methode})
        except Exception as exc:  # noqa: BLE001 — PDF chiffré, abîmé…
            if isinstance(exc, RuntimeError) and "pdfplumber" in str(exc):
                raise
            raise ValueError("Ce fichier PDF ne peut pas être lu (abîmé ou protégé par un mot de passe).") from exc
        return {"pages": pages, "tableaux": tableaux, "caracteres": caracteres}

    def lire_conversion(jeton):
        d = H.doc_lire(cle_conversion(), None) or {}
        if d.get("jeton") != jeton:
            abort(404)
        chemin = chemin_tmp("conv", jeton, "json")
        if not os.path.exists(chemin):
            flash("Le résultat de la conversion a expiré : déposez de nouveau le PDF.", "warning")
            return None, d
        with open(chemin, encoding="utf-8") as fh:
            return json.load(fh), d

    @app.route("/conversion-pdf", methods=["GET", "POST"])
    @H.connexion_requise("donnees")
    def conversion_pdf():
        if request.method == "POST":
            f = request.files.get("fichier")
            if not f or not f.filename:
                flash("Choisissez un fichier PDF.", "warning")
                return redirect(url_for("conversion_pdf"))
            contenu = f.read()
            if not f.filename.lower().endswith(".pdf") or not contenu.startswith(b"%PDF"):
                flash("Ce fichier n'est pas un PDF.", "danger")
                return redirect(url_for("conversion_pdf"))
            try:
                res = extraire(contenu)
            except RuntimeError as exc:
                flash(str(exc), "danger")
                return redirect(url_for("conversion_pdf"))
            except ValueError as exc:
                flash(str(exc), "danger")
                return redirect(url_for("conversion_pdf"))
            if res["caracteres"] == 0:
                flash("Ce PDF ne contient pas de texte : c'est probablement un document scanné (une image). "
                      "La conversion n'est possible qu'avec un PDF « texte » (produit par un logiciel : Word, Excel, "
                      "GESCOL, EducMaster…). Pour un scan, il faut d'abord une reconnaissance de caractères (OCR).", "danger")
                return redirect(url_for("conversion_pdf"))
            if not res["tableaux"]:
                flash("Aucun tableau n'a été trouvé dans ce PDF (le texte ne présente pas de colonnes).", "warning")
                return redirect(url_for("conversion_pdf"))
            nettoyer_tmp()
            ancien = H.doc_lire(cle_conversion(), None) or {}
            if RE_JETON.match(ancien.get("jeton", "")):
                try:
                    os.remove(chemin_tmp("conv", ancien["jeton"], "json"))
                except OSError:
                    pass
            jeton = secrets.token_hex(8)
            with open(chemin_tmp("conv", jeton, "json"), "w", encoding="utf-8") as fh:
                json.dump(res, fh, ensure_ascii=False)
            H.doc_ecrire(cle_conversion(), {"jeton": jeton, "nom": os.path.basename(f.filename)[:120],
                                            "le": H.maintenant()})
            H.journal("conversion_pdf", f"{f.filename[:80]} : {len(res['tableaux'])} tableau(x)")
            flash(f"{len(res['tableaux'])} tableau(x) extrait(s) de {res['pages']} page(s).", "success")
            return redirect(url_for("conversion_pdf", jeton=jeton))
        jeton = request.args.get("jeton", "")
        res, info = (None, None)
        if jeton:
            res, info = lire_conversion(jeton)
        return render_template("conversion_pdf.html", res=res, info=info, jeton=jeton if res else "")

    def nom_base(info):
        return re.sub(r"[^A-Za-z0-9_-]+", "_", os.path.splitext(info.get("nom") or "conversion")[0])[:60] or "conversion"

    def nombre(v):
        """« 12,5 » -> 12.5 ; « 15 » -> 15 ; les codes à zéro initial (« 0102 ») et les textes restent du texte."""
        t = (v or "").replace("\u00a0", " ").strip()
        if re.match(r"^-?\d{1,12}([.,]\d+)?$", t) and not (re.match(r"^-?0\d", t)):
            x = float(t.replace(",", "."))
            return int(x) if x.is_integer() and "," not in t and "." not in t else x
        return v

    @app.route("/conversion-pdf/<jeton>.xlsx")
    @H.connexion_requise("donnees")
    def conversion_xlsx(jeton):
        res, info = lire_conversion(jeton)
        if res is None:
            return redirect(url_for("conversion_pdf"))
        feuilles = []
        for t in res["tableaux"]:
            lignes = t["lignes"]
            feuilles.append((f"Page {t['page']} - T{t['n']}", lignes[0], [[nombre(v) for v in l] for l in lignes[1:]],
                             {"titre_haut": [f"{info.get('nom')} — page {t['page']}, tableau {t['n']}"]}))
        H.journal("conversion_xlsx", info.get("nom", ""))
        try:
            return X.reponse_classeur(f"{nom_base(info)}.xlsx", feuilles)
        except X.ExcelIndisponible as exc:
            flash(str(exc), "danger")
            return redirect(url_for("conversion_pdf", jeton=jeton))

    @app.route("/conversion-pdf/<jeton>.csv")
    @H.connexion_requise("donnees")
    def conversion_csv(jeton):
        res, info = lire_conversion(jeton)
        if res is None:
            return redirect(url_for("conversion_pdf"))
        n = request.args.get("t", type=int)
        tabs = res["tableaux"] if n is None else [t for i, t in enumerate(res["tableaux"], 1) if i == n]
        if not tabs:
            abort(404)
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        for i, t in enumerate(tabs):
            if len(tabs) > 1:
                if i:
                    w.writerow([])
                w.writerow([f"Page {t['page']} - tableau {t['n']}"])
            w.writerows(t["lignes"])
        nom = nom_base(info) + (f"_tableau{n}" if n else "") + ".csv"
        return Response("﻿" + buf.getvalue(), mimetype="text/csv",
                        headers={"Content-Disposition": f"attachment; filename={nom}"})
